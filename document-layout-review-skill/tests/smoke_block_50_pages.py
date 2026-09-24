#!/usr/bin/env python3
"""Real 50-page scope/IO smoke test; not a full template/vision/Word F9 benchmark.

No model requests or network access. LibreOffice render is optional unless
--require-render is specified. The caller controls the temporary output folder.
"""
from __future__ import annotations
import argparse
import json
import shutil
import subprocess
import sys
import time
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Pt
from lxml import etree as E

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import block_progress as bp


def write_package(path, parts):
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)


def run(root,require_render=False):
    root=Path(root).resolve();root.mkdir(parents=True,exist_ok=True)
    if (root/'source.docx').exists():raise ValueError('Use an empty output directory; do not overwrite a prior test.')
    source=root/'source.docx';d=Document();d.styles['Normal'].font.size=Pt(12)
    for page in range(1,51):
        heading=d.add_paragraph(f'{page}. Block workflow test',style='Heading 1')
        if page>1:heading.paragraph_format.page_break_before=True
        p=d.add_paragraph(f'Page {page}: contract period 30 days; project value 500000. ' * 8)
        for r in p.runs:r.font.size=Pt(8)
    d.save(source);original_hash=bp.sha(source)
    work=root/'work';start=time.perf_counter()
    bp.initialize(source,work,target_chars=1,max_paragraphs=1)
    bp.accept_global(work,bp.load(work)['current'],'Synthetic global rules: body 12pt, no content edits.')
    blocks=list(bp.load(work)['blocks']);changes=0
    for block in blocks:
        state=bp.load(work);parts,xml,body=bp.package(state['current'])
        for node in list(body)[block['start']:block['end']]:
            for p in ([node] if node.tag==bp.Q('p') else node.iter(bp.Q('p'))):
                style=p.find('w:pPr/w:pStyle',bp.NS)
                if style is not None and style.get(bp.Q('val'))=='Heading1':continue
                for r in p.iter(bp.Q('r')):
                    rp=r.find(bp.Q('rPr'))
                    if rp is None:rp=E.Element(bp.Q('rPr'));r.insert(0,rp)
                    sz=rp.find(bp.Q('sz'))
                    if sz is None:sz=E.SubElement(rp,bp.Q('sz'))
                    if sz.get(bp.Q('val'))!='24':sz.set(bp.Q('val'),'24');changes+=1
        parts[bp.DOC]=E.tostring(xml,encoding='UTF-8',xml_declaration=True,standalone=True)
        proposal=root/'proposal.docx';write_package(proposal,parts)
        bp.checkpoint(work,block['id'],proposal,'Corrected only the selected block body font size; scope verified.')
    elapsed=time.perf_counter()-start;state=bp.load(work)
    assert len(blocks)==50 and changes==50
    old=Document(source);new=Document(state['current'])
    assert [p.text for p in old.paragraphs]==[p.text for p in new.paragraphs]
    assert bp.sha(source)==original_hash
    assert all(b['status']=='checked' for b in state['blocks'])
    assert all(r.font.size.pt==12 for p in new.paragraphs if p.style.name=='Normal' for r in p.runs)
    # Only the requested block becomes pending, not all completed pages.
    bp.reopen(work,'B0023','Synthetic targeted follow-up')
    reopened=bp.load(work);assert sum(b['status']=='checked' for b in reopened['blocks'])==49
    bp.checkpoint(work,'B0023',reopened['current'],'No further mutation; local recheck complete.')
    result={'status':'passed','kind':'real-docx-scope-and-export-smoke','blocks':50,'runs_corrected':50,
            'body_text_preserved':True,'source_unchanged':True,'reopen_kept_other_blocks':49,
            'block_processing_seconds':round(elapsed,3),'model_calls':0,'whole_document_renders_during_blocks':0,
            'native_word_wps_tested':False,'full_skill_acceptance_tested':False}
    office=shutil.which('soffice') or shutil.which('libreoffice')
    if office:
        candidate=root/'candidate.docx';shutil.copyfile(bp.load(work)['current'],candidate)
        start=time.perf_counter()
        cp=subprocess.run([office,'-env:UserInstallation='+(root/'lo-profile').as_uri(),'--headless','--convert-to','pdf','--outdir',str(root),str(candidate)],capture_output=True,text=True,timeout=120)
        pdf=root/'candidate.pdf'
        if cp.returncode or not pdf.is_file():raise RuntimeError(cp.stderr or cp.stdout)
        import fitz
        with fitz.open(pdf) as doc:
            count=len(doc)
            assert count==50
            assert 'Page 23:' in doc[22].get_text()
        result.update(renderer='LibreOffice',rendered_pages=count,render_seconds=round(time.perf_counter()-start,3),final_export_count=1)
    elif require_render:
        raise RuntimeError('No executable LibreOffice renderer for this smoke test.')
    else:result['render_status']='not_run'
    (root/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return result


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--out-dir',type=Path,required=True);ap.add_argument('--require-render',action='store_true');a=ap.parse_args()
    print(json.dumps(run(a.out_dir,a.require_render),ensure_ascii=False,indent=2))
