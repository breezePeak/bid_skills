"""Real DOCX regressions; external Office/vision boundaries are explicitly mocked.

No test response here is evidence of real model accuracy or native Word/WPS.
"""
import copy, io, json, os, shutil, sys, types, zipfile
from pathlib import Path
from unittest.mock import patch
import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree as E
from PIL import Image

SCRIPTS=Path(__file__).resolve().parents[1]/'scripts'
sys.path.insert(0,str(SCRIPTS))
import block_progress as bp
import block_workflow as bw
import block_checks as ck
import office_stability as office
import retry_guard as retry
import visual_evidence as ve
import host_visual as hv
import image_repairs as ir
import format_equivalence as eq


def write_doc(path, text='工期30天。', header=0):
    d=Document();d.styles['Normal'].font.size=Pt(12);d.add_paragraph(text);d.add_paragraph('第二块50万元。')
    if header:
        h=d.sections[0].header;h.paragraphs[0].text='页眉1'
        for i in range(1,header):h.add_paragraph('页眉'+str(i+1))
    d.save(path);return path


def rewrite(source,out,edit):
    with zipfile.ZipFile(source) as z:parts={n:z.read(n) for n in z.namelist()}
    edit(parts)
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in parts.items():z.writestr(n,data)
    return out


def xml(n):return E.tostring(n,encoding='UTF-8',xml_declaration=True,standalone=True)


def edit_body(source,out,edit):
    def change(parts):
        root=bp.parse(parts[bp.DOC]);edit(root.find(bp.Q('body')));parts[bp.DOC]=xml(root)
    return rewrite(source,out,change)


def setup(tmp_path, *, headers=0, ready=True, text='工期30天。'):
    src=write_doc(tmp_path/'source.docx',text,headers);work=tmp_path/'work'
    profile=tmp_path/'profile.json';bp.write_json(profile,{'semantic_roles':{'body':{'style_id':'Normal'}}})
    rules=tmp_path/'rules.json';bp.write_json(rules,{'body':{'preserve_fonts':['黑体','SimHei']}})
    settings={'template':str(src),'template_style_json':str(profile),'text_rules':str(rules),'renderer':'auto','field_engine':'auto'}
    bp.initialize(src,work,settings,target_chars=1,max_paragraphs=1)
    if ready:bp.accept_global(work,bp.load(work)['current'],'模板设置已核对')
    return src,work


@pytest.mark.parametrize('bad', ['工期3天。','工期300天。','工期30天。新增承诺','', '工期-30天。'])
def test_bad_business_text_does_not_advance(tmp_path,bad):
    _,work=setup(tmp_path);state=bp.load(work);d=Document(state['current']);d.paragraphs[0].text=bad
    proposal=tmp_path/'bad.docx';d.save(proposal)
    with pytest.raises(bp.BlockError,match='当前块检查失败'):bp.checkpoint(work,'B0001',proposal,'修复')
    assert bp.load(work)['revision']==state['revision']
    assert bp.next_block(work)['completed']==0


@pytest.mark.parametrize('size',[72,13,1])
def test_wrong_explicit_font_size_rejected_before_checkpoint(tmp_path,size):
    _,work=setup(tmp_path);state=bp.load(work);d=Document(state['current']);d.paragraphs[0].runs[0].font.size=Pt(size)
    proposal=tmp_path/'bad.docx';d.save(proposal)
    with pytest.raises(bp.BlockError,match='block-run-format'):bp.checkpoint(work,'B0001',proposal,'改字号')
    assert bp.load(work)['current_sha256']==state['current_sha256']


def test_default_body_heiti_exception_is_not_removed(tmp_path):
    _,work=setup(tmp_path);state=bp.load(work);d=Document(state['current']);d.paragraphs[0].runs[0].font.name='SimHei'
    proposal=tmp_path/'heiti.docx';d.save(proposal)
    bp.checkpoint(work,'B0001',proposal,'保留正文黑体')
    assert bp.next_block(work)['completed']==1


def test_explicit_body_size_override_is_used(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work);bp.write_json(s['settings']['text_rules'],{'body':{'size_pt':15}})
    d=Document(s['current']);r=d.paragraphs[0].runs[0];r.font.size=Pt(15)
    rp=r._r.get_or_add_rPr();E.SubElement(rp,bp.Q('szCs'),{bp.Q('val'):'30'})
    out=tmp_path/'override.docx';d.save(out);bp.checkpoint(work,'B0001',out,'显式15磅')
    assert bp.next_block(work)['completed']==1


@pytest.mark.parametrize('initial,final,shared',[(1,2,False),(2,1,False),(1,2,True),(2,1,True)])
def test_header_resize_updates_state_atomically(tmp_path,initial,final,shared):
    _,work=setup(tmp_path,headers=initial,ready=shared);old=bp.load(work)
    d=Document(old['current']);h=d.sections[0].header
    if final>initial:h.add_paragraph('新增页眉')
    else:h._element.remove(h.paragraphs[-1]._p)
    out=tmp_path/'header.docx';d.save(out)
    (bp.accept_shared_repair if shared else bp.accept_global)(work,out,'按模板修改页眉')
    s=bp.load(work);block=next(b for b in s['blocks'] if b['part']=='word/header1.xml')
    assert block['end']==final and block['status']=='pending'
    assert bp.next_block(work)['block']['id']=='B0001'


def test_global_validation_failure_does_not_save_broken_state(tmp_path,monkeypatch):
    _,work=setup(tmp_path,ready=False);before=(work/bp.STATE).read_bytes();s=bp.load(work)
    d=Document(s['current']);d.sections[0].header.paragraphs[0].text='新页眉';out=tmp_path/'out.docx';d.save(out)
    monkeypatch.setattr(bp,'validate_blocks',lambda *a,**k:(_ for _ in ()).throw(bp.BlockError('inject validation failure')))
    with pytest.raises(bp.BlockError):bp.accept_global(work,out,'测试事务')
    assert (work/bp.STATE).read_bytes()==before
    assert not (work/'checkpoints/r000001.docx').exists()


def test_write_failure_rolls_back_new_checkpoint_only(tmp_path,monkeypatch):
    _,work=setup(tmp_path);before=(work/bp.STATE).read_bytes();state=bp.load(work)
    d=Document(state['current']);d.paragraphs[0].runs[0].font.size=Pt(12);out=tmp_path/'p.docx';d.save(out)
    real=bp.write_json
    def fail(path,data):
        if Path(path)==work/bp.STATE:raise OSError('disk full')
        return real(path,data)
    monkeypatch.setattr(bp,'write_json',fail)
    with pytest.raises(OSError):bp.checkpoint(work,'B0001',out,'验证磁盘失败')
    assert (work/bp.STATE).read_bytes()==before
    assert not (work/f"checkpoints/r{state['revision']+1:06d}.docx").exists()


def test_correct_unchanged_block_has_local_receipt(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work);bp.checkpoint(work,'B0001',s['current'],'检查无需修改')
    receipt=bp.load(work)['blocks'][0]['local_audit']
    assert receipt['content_checked'] and receipt['field_structure_checked'] and receipt['template_checked']
    assert receipt['page_review']=='pending-final-layout'


def test_cross_paragraph_toc_is_not_split(tmp_path):
    src=tmp_path/'fields.docx';d=Document()
    p=d.add_paragraph();r=E.SubElement(p._p,bp.Q('r'));E.SubElement(r,bp.Q('fldChar'),{bp.Q('fldCharType'):'begin'})
    E.SubElement(E.SubElement(p._p,bp.Q('r')),bp.Q('instrText')).text=' TOC \\o "1-3" '
    E.SubElement(E.SubElement(p._p,bp.Q('r')),bp.Q('fldChar'),{bp.Q('fldCharType'):'separate'})
    d.add_paragraph('目录第一项');p=d.add_paragraph('目录第二项')
    E.SubElement(E.SubElement(p._p,bp.Q('r')),bp.Q('fldChar'),{bp.Q('fldCharType'):'end'})
    d.add_paragraph('后文');d.save(src)
    blocks=bp.make_blocks(src,target_chars=1,max_paragraphs=1)
    assert blocks[0]['start']==0 and blocks[0]['end']==3
    _,_,body=bp.package(src);ck.field_structure(list(body)[:3])


@pytest.mark.parametrize('kind',['open','orphan','locked'])
def test_bad_fields_stop_before_progress(tmp_path,kind):
    _,work=setup(tmp_path);state=bp.load(work)
    def change(body):
        r=E.SubElement(body[0],bp.Q('r'))
        attr={bp.Q('fldCharType'):'end' if kind=='orphan' else 'begin'}
        if kind=='locked':attr[bp.Q('fldLock')]='true'
        E.SubElement(r,bp.Q('fldChar'),attr)
    out=edit_body(state['current'],tmp_path/'bad.docx',change)
    with pytest.raises(bp.BlockError):bp.checkpoint(work,'B0001',out,'域错误')
    assert bp.load(work)['revision']==state['revision']


def test_operation_cache_reuses_readers_and_clears_after_operation(tmp_path,monkeypatch):
    _,work=setup(tmp_path);source=bp.load(work)['current'];counter=[];real=bp.zipfile.ZipFile
    def counted(file,*a,**k):
        if isinstance(file,(str,Path)) and str(file)==str(source):counter.append(str(file))
        return real(file,*a,**k)
    monkeypatch.setattr(bp.zipfile,'ZipFile',counted)
    with bp.package_session():
        for _ in range(7):bp.package(source)
    assert len(counter)==1
    bp.package(source);assert len(counter)==2


def test_checkpoint_plus_next_avoids_seven_package_reads(tmp_path,monkeypatch):
    _,work=setup(tmp_path);source=bp.load(work)['current'];counter=[];real=bp.zipfile.ZipFile
    def counted(file,*a,**k):
        if isinstance(file,(str,Path)) and Path(file)==Path(source):counter.append(str(file))
        return real(file,*a,**k)
    monkeypatch.setattr(bp.zipfile,'ZipFile',counted)
    with bp.package_session():
        bp.checkpoint(work,'B0001',source,'已核对');bp.next_block(work)
    assert len(counter)==1


@pytest.mark.parametrize('tag,val',[('kern','2'),('widowControl','0')])
def test_real_kerning_or_pagination_drift_is_not_ignored(tmp_path,tag,val):
    source=write_doc(tmp_path/'s.docx')
    def change(body):
        owner=body[0].find(bp.Q('r')) if tag=='kern' else body[0]
        pr=owner.find(bp.Q('rPr' if tag=='kern' else 'pPr'))
        if pr is None:pr=E.Element(bp.Q('rPr' if tag=='kern' else 'pPr'));owner.insert(0,pr)
        E.SubElement(pr,bp.Q(tag),{bp.Q('val'):val})
    out=edit_body(source,tmp_path/'saved.docx',change)
    assert office.compare(source,out)['status']=='failed'


def test_equivalent_explicit_defaults_do_not_create_save_loop(tmp_path):
    source=write_doc(tmp_path/'s.docx')
    def change(body):
        p=body[0];pp=E.Element(bp.Q('pPr'));p.insert(0,pp)
        E.SubElement(pp,bp.Q('widowControl'),{bp.Q('val'):'1'})
        E.SubElement(pp,bp.Q('ind'),{bp.Q('left'):'0'})
        r=p.find(bp.Q('r'));rp=E.Element(bp.Q('rPr'));r.insert(0,rp)
        E.SubElement(rp,bp.Q('kern'),{bp.Q('val'):'0'})
    out=edit_body(source,tmp_path/'saved.docx',change)
    assert office.compare(source,out)['status']=='passed'


def test_office_cleanup_failure_preserves_valid_output(tmp_path,monkeypatch):
    out=write_doc(tmp_path/'out.docx');warnings=[]
    monkeypatch.setattr(office.shutil,'rmtree',lambda p:(_ for _ in ()).throw(PermissionError('in use')))
    with office.office_workspace(tmp_path,warnings) as folder:
        assert Path(folder).is_dir();assert office.complete_output(out)['size']>0
    assert out.exists() and warnings[0]['code']=='office-cleanup-incomplete'


def test_cleanup_warning_does_not_hide_operation_failure(tmp_path,monkeypatch):
    monkeypatch.setattr(office.shutil,'rmtree',lambda p:(_ for _ in ()).throw(PermissionError('in use')))
    with pytest.raises(ValueError,match='operation failed'):
        with office.office_workspace(tmp_path,[]):raise ValueError('operation failed')


@pytest.mark.parametrize('data',[b'',b'not a zip',b'PK invalid'])
def test_invalid_output_never_passes_cleanup_path(tmp_path,data):
    p=tmp_path/'bad.docx';p.write_bytes(data)
    with pytest.raises(Exception):office.complete_output(p)


def test_probe_only_once_for_same_effective_setup(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work);calls=[]
    def save(src,dst,*a,**k):calls.append(1);shutil.copyfile(src,dst);return {'engine':'explicit test-double'}
    a=office.probe_once(work,s,s['current'],save);b=office.probe_once(work,s,s['current'],save)
    assert a['status']==b['status']=='passed' and len(calls)==1


def test_failed_probe_preserves_evidence_and_is_not_repeated(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work);calls=[]
    def save(src,dst,*a,**k):
        calls.append(1)
        def alter(body):
            p=body[0];pp=p.find(bp.Q('pPr'))
            if pp is None:pp=E.Element(bp.Q('pPr'));p.insert(0,pp)
            E.SubElement(pp,bp.Q('widowControl'),{bp.Q('val'):'0'})
        edit_body(src,dst,alter);return {'engine':'explicit test-double'}
    for _ in range(2):
        with pytest.raises(ValueError):office.probe_once(work,s,s['current'],save)
    assert len(calls)==1 and list((work/'office-probe').glob('*/saved.docx'))


def test_retry_fingerprint_ignores_rename_and_repackaging(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work);key=retry.key_for(s,SCRIPTS)
    out=rewrite(s['current'],tmp_path/'renamed-final-999.docx',lambda d:None)
    newer=copy.deepcopy(s);newer.update(current=str(out),current_sha256=bp.sha(out),revision=999)
    assert retry.key_for(newer,SCRIPTS)==key
    retry.record_failure(newer,key,[{'code':'repeat','path':'/one'}]);retry.record_failure(newer,key,[{'code':'repeat','path':'/two'}])
    with pytest.raises(bp.BlockError,match='连续两次'):retry.before_retry(newer,SCRIPTS)


def test_real_fix_releases_retry_lock(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work);key=retry.key_for(s,SCRIPTS)
    for _ in range(2):retry.record_failure(s,key,[{'code':'bad'}])
    d=Document(s['current']);d.paragraphs[0].runs[0].font.size=Pt(13);out=tmp_path/'different.docx';d.save(out);s['current']=str(out)
    assert retry.before_retry(s,SCRIPTS)!=key


def test_config_path_persists_on_first_explicit_use(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work);p=tmp_path/'vision.json';bp.write_json(p,{'command':['host-vision']})
    got=bw.persist_vision_config(work,s,p)
    assert got==str(p.resolve()) and bp.load(work)['settings']['vision_worker_config']==got
    assert bw.persist_vision_config(work,bp.load(work),None)==got


def test_config_cannot_silently_switch_provider(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work)
    for n in ('a','b'):bp.write_json(tmp_path/(n+'.json'),{'command':[n]})
    bw.persist_vision_config(work,s,tmp_path/'a.json')
    with pytest.raises(bp.BlockError):bw.persist_vision_config(work,bp.load(work),tmp_path/'b.json')


def vision_task(tmp_path, phase='initial'):
    im=tmp_path/'image.png';Image.new('RGB',(64,48),'white').save(im)
    page=tmp_path/'page.png';Image.new('RGB',(80,100),'white').save(page)
    ref=ve.prepare([im],tmp_path,{'source_sha256':'fixture','object_id':'F0003'},phase=phase,
        originals=[im] if phase!='initial' else (),pages=[{'name':'page.png',**ve.reference(page)}] if phase=='final' else ())
    with pytest.raises(ve.VisualError) as raised:ve.run(ref,None,tmp_path)
    assert raised.value.code=='visual-host-review-required'
    task=Path(raised.value.details['task']);data=ve.read(task);request=ve.read(ve.checked_file(data['request']))
    checks={k:('not_applicable' if k=='connections_correct' else 'pass') for k in ve.CHECKS}
    response={'_test_double':True,'request_id':request['request_id'],'image_type':'diagram','observation':'TEST DOUBLE; no model executed',
       'semantics_preserved':True,'page_match':True,'views':[{'view_id':v,'observation':'TEST DOUBLE', 'checks':dict(checks),'not_applicable_reasons':{'connections_correct':'TEST'}} for v in request['required_view_ids']],'findings':[]}
    log={'_test_double':True,'request_id':request['request_id'],'tool_calls':[{'tool':'explicit-test-double','reference':'TEST ONLY, not host evidence',
        'images':[{'id':i['id'],'sha256':i['sha256']} for i in request['images']]}]}
    return ref,task,response,log


def test_host_missing_worker_creates_pending_not_pass(tmp_path):
    ref,task,_,_=vision_task(tmp_path)
    assert not (task.parent/'inspection.json').exists()
    with pytest.raises(ve.VisualError) as raised:ve.run(ref,None,tmp_path)
    assert raised.value.details['task']==str(task)


@pytest.mark.parametrize('phase',['initial','repair','final'])
def test_host_protocol_completion_requires_real_input_coverage(tmp_path,phase):
    ref,task,response,log=vision_task(tmp_path,phase)
    result=hv.complete_host(task,response,log,allow_test_double=True)
    assert ve.validate_inspection(result['inspection'],allow_test_double=True)['verdict']=='pass'
    with pytest.raises(ve.VisualError,match='模拟'):ve.validate_inspection(result['inspection'])


@pytest.mark.parametrize('bad',['missing-call','missing-image','old-image','wrong-id','no-reference'])
def test_host_incomplete_logs_do_not_pass(tmp_path,bad):
    ref,task,response,log=vision_task(tmp_path)
    if bad=='missing-call':log['tool_calls']=[]
    if bad=='missing-image':log['tool_calls'][0]['images'].pop()
    if bad=='old-image':log['tool_calls'][0]['images'][0]['sha256']='old'
    if bad=='wrong-id':log['request_id']='old'
    if bad=='no-reference':log['tool_calls'][0]['reference']=''
    with pytest.raises(ve.VisualError):hv.complete_host(task,response,log,allow_test_double=True)
    assert not (task.parent/'inspection.json').exists()


def test_host_simulated_records_refused_by_production(tmp_path):
    _,task,response,log=vision_task(tmp_path)
    with pytest.raises(ve.VisualError,match='模拟'):hv.complete_host(task,response,log)


def test_host_failure_cannot_be_overwritten_by_pass(tmp_path):
    _,task,response,log=vision_task(tmp_path)
    response['views'][0]['checks']['text_inside_bounds']='fail'
    response['findings']=[{'view_id':response['views'][0]['view_id'],'check':'text_inside_bounds','description':'TEST ONLY','bbox':[0.1,0.1,0.2,0.2]}]
    result=hv.complete_host(task,response,log,allow_test_double=True);assert result['verdict']=='fail'
    response['findings']=[];response['views'][0]['checks']['text_inside_bounds']='pass'
    with pytest.raises(ve.VisualError,match='不得覆盖'):hv.complete_host(task,response,log,allow_test_double=True)


def test_repair_requires_original_pixels(tmp_path):
    im=tmp_path/'i.png';Image.new('RGB',(10,10)).save(im)
    with pytest.raises(ve.VisualError):ve.prepare([im],tmp_path,{},phase='repair')


def test_early_repaint_permission_refuses_stale_defect(tmp_path,monkeypatch):
    work=tmp_path;state={'source_sha256':'source','current_sha256':'current'};block={'id':'B0001'}
    bp.write_json(work/'image-plans/B0001.json',{'source_sha256':'source','current_sha256':'current','objects':[{'id':'F0003','redraw_allowed':True,'current_object_sha256':'new','defect_evidence':['old-proof']}]})
    monkeypatch.setattr(ve,'validate_inspection',lambda *a,**k:{'binding':{'source_sha256':'source','object_id':'F0003','object_sha256':'old'},'verdict':'fail','checks':{'text_inside_bounds':'fail'}})
    with pytest.raises(bp.BlockError,match='真实失败依据'):ir._permission(work,state,block,'F0003','new')


def test_early_repaint_permission_accepts_bound_real_defect_contract(tmp_path,monkeypatch):
    work=tmp_path;state={'source_sha256':'source','current_sha256':'current'};block={'id':'B0001'}
    bp.write_json(work/'image-plans/B0001.json',{'source_sha256':'source','current_sha256':'current','objects':[{'id':'F0003','redraw_allowed':True,'current_object_sha256':'bad','defect_evidence':['explicit-boundary-double']}]})
    monkeypatch.setattr(ve,'validate_inspection',lambda *a,**k:{'binding':{'source_sha256':'source','object_id':'F0003','object_sha256':'bad'},'verdict':'fail','checks':{'text_inside_bounds':'fail'}})
    ir._permission(work,state,block,'F0003','bad')


def test_no_image_plan_cannot_start_repaint(tmp_path):
    with pytest.raises(bp.BlockError,match='先执行'):ir._permission(tmp_path,{}, {'id':'B1'},'F0003','bad')


def test_style_ui_metadata_is_not_a_format_mutation(tmp_path):
    source=write_doc(tmp_path/'s.docx')
    def change(parts):
        root=bp.parse(parts['word/styles.xml']);style=root.find("w:style[@w:styleId='Normal']",bp.NS)
        E.SubElement(style,bp.Q('uiPriority'),{bp.Q('val'):'999'});parts['word/styles.xml']=xml(root)
    saved=rewrite(source,tmp_path/'saved.docx',change)
    assert eq.equivalent_styles(source,saved,'Normal')


@pytest.mark.parametrize('tag,val',[('kern','2'),('widowControl','0'),('sz','144')])
def test_style_effective_mutations_remain_blocking(tmp_path,tag,val):
    source=write_doc(tmp_path/'s.docx')
    def change(parts):
        root=bp.parse(parts['word/styles.xml']);s=root.find("w:style[@w:styleId='Normal']",bp.NS)
        pr=s.find(bp.Q('pPr' if tag=='widowControl' else 'rPr'))
        if pr is None:pr=E.SubElement(s,bp.Q('pPr' if tag=='widowControl' else 'rPr'))
        for n in list(pr.findall(bp.Q(tag))):pr.remove(n)
        E.SubElement(pr,bp.Q(tag),{bp.Q('val'):val});parts['word/styles.xml']=xml(root)
    saved=rewrite(source,tmp_path/'saved.docx',change)
    assert not eq.equivalent_styles(source,saved,'Normal')


def test_conformance_only_suppresses_proven_equivalent_style_diff(tmp_path,monkeypatch):
    source=write_doc(tmp_path/'s.docx')
    monkeypatch.setitem(sys.modules,'template_conformance',types.SimpleNamespace(check=lambda *a:{'status':'failed','issues':[
       {'severity':'error','code':'template-style-mutated','style_id':'Normal'}, {'severity':'error','code':'template-section-geometry-mismatch'}]}))
    import stable_template_audit as audit
    result=audit.check(source,source)
    assert result['status']=='failed' and len(result['issues'])==1
    assert result['issues'][0]['code']=='template-section-geometry-mismatch'


def test_office_quit_error_is_separate_warning_in_worker_source():
    text=(SCRIPTS/'office_com.ps1').read_text(encoding='utf-8-sig')
    assert 'cleanup_status' in text and 'operation_status' in text and 'output_validation_required' in text
    quit_line=next(line for line in text.splitlines() if "Quit: " in line)
    assert "status='failed'" not in quit_line


def test_user_ban_overrides_default_heiti_preservation(tmp_path):
    _,work=setup(tmp_path);s=bp.load(work)
    bp.write_json(s['settings']['text_rules'],{'body':{'preserve_fonts':['SimHei'],'forbidden_fonts':['黑体']}})
    d=Document(s['current']);d.paragraphs[0].runs[0].font.name='SimHei';out=tmp_path/'bad-font.docx';d.save(out)
    with pytest.raises(bp.BlockError,match='block-font'):bp.checkpoint(work,'B0001',out,'必须拒绝')


def test_explicit_font_replaces_inherited_theme_face():
    from template_format_contract import merge
    base=E.Element(bp.Q('rPr'));E.SubElement(base,bp.Q('rFonts'),{bp.Q('asciiTheme'):'minorAscii'})
    new=E.Element(bp.Q('rPr'));E.SubElement(new,bp.Q('rFonts'),{bp.Q('ascii'):'SimHei'})
    merge(base,new)
    assert base.find(bp.Q('rFonts')).get(bp.Q('asciiTheme')) is None


@pytest.mark.parametrize('rtl,left,start',[(False,'left','start'),(True,'right','start'),(False,'right','end'),(True,'left','end')])
def test_alignment_equivalence_respects_direction(rtl,left,start):
    from template_format_contract import value_key
    props=E.Element(bp.Q('pPr'));E.SubElement(props,bp.Q('bidi'),{bp.Q('val'):'1' if rtl else '0'});n=E.SubElement(props,bp.Q('jc'),{bp.Q('val'):start})
    assert value_key(props,'jc')==left


def test_explicit_complex_script_false_equals_absence():
    from template_format_contract import value_key
    a=E.Element(bp.Q('rPr'));b=E.Element(bp.Q('rPr'));E.SubElement(b,bp.Q('bCs'),{bp.Q('val'):'0'})
    assert value_key(a,'bCs')==value_key(b,'bCs')==False


def test_kern_numeric_spelling_is_normalized_not_ignored():
    from template_format_contract import value_key
    a=E.Element(bp.Q('rPr'));n=E.SubElement(a,bp.Q('kern'),{bp.Q('val'):'00'})
    assert value_key(a,'kern')=='0'
    n.set(bp.Q('val'),'2');assert value_key(a,'kern')=='2'


def test_auxiliary_inventory_uses_each_story_relationships(tmp_path,monkeypatch):
    """Real DOCX parts with an explicit small core-inventory boundary double."""
    import review_objects
    im=tmp_path/'logo.png';Image.new('RGB',(12,10)).save(im)
    source=tmp_path/'stories.docx';d=Document();d.add_picture(str(im));d.sections[0].header.paragraphs[0].add_run().add_picture(str(im));d.sections[0].footer.add_table(rows=1,cols=1,width=Pt(100));d.save(source)
    with zipfile.ZipFile(source) as z:files={n:z.read(n) for n in z.namelist()}
    class FakeDoc:
        def __init__(self):self.files=files;self.roots={};self.document=self.root(bp.DOC)
        def root(self,name):
            if name not in self.roots:self.roots[name]=bp.parse(self.files[name])
            return self.roots[name]
        def refresh(self):pass
    visited=[]
    def core(doc):
        # The real wrapper must switch BOTH story XML and relationship lookup.
        rels=doc.root(bp.RELS) if bp.RELS in doc.files else []
        mapping={r.get('Id'):r.get('Target') for r in rels}
        rows=[]
        for p in doc.document.iter(bp.Q('p')):
            refs=p.xpath('.//@r:embed',namespaces={'r':bp.R})
            if refs:
                assert all(r in mapping for r in refs)
                rows.append({'id':'F0001','kind':'figure','element':p,'caption':None,'hash':'fixture','paths':[mapping[refs[0]]]})
        for t in doc.document.iter(bp.Q('tbl')):rows.append({'id':'T0001','kind':'table','element':t,'caption':None,'hash':'fixture','paths':[]})
        visited.append(doc.document.tag);return rows
    monkeypatch.setitem(sys.modules,'numbering_policy',types.SimpleNamespace(inventory=core))
    got=review_objects.inventory(FakeDoc())
    assert {r['id'] for r in got}=={'F0001','S_header1_F0001','S_footer1_T0001'}
    assert {r['part'] for r in got}=={bp.DOC,'word/header1.xml','word/footer1.xml'}
    assert len(set(r['id'] for r in got))==3


def test_two_identical_final_failures_do_not_start_third_office_call(tmp_path,monkeypatch):
    _,work=setup(tmp_path)
    for b in list(bp.load(work)['blocks']):bp.checkpoint(work,b['id'],bp.load(work)['current'],'已核对')
    calls=[]
    def fail(*a):calls.append(1);raise RuntimeError('same actual Office failure')
    monkeypatch.setattr(bw,'_final_audit_impl',fail)
    for _ in range(2):
        with pytest.raises(RuntimeError):bw.final_audit(work,bp.load(work),types.SimpleNamespace(retry_final=True))
    with pytest.raises(bp.BlockError,match='连续两次'):bw.final_audit(work,bp.load(work),types.SimpleNamespace(retry_final=True))
    assert len(calls)==2


def test_table_audit_ignores_ui_metadata_but_keeps_real_properties(tmp_path,monkeypatch):
    import stable_table_audit as audit
    source=write_doc(tmp_path/'source.docx')
    class FakeDoc:
        def __init__(self,p):
            self.document=E.Element(bp.Q('body'));E.SubElement(self.document,bp.Q('tbl'))
            self.flat=E.Element(bp.Q('style'),{bp.Q('styleId'):'table'})
            E.SubElement(self.flat,bp.Q('name'),{bp.Q('val'):'UI renamed by Office'})
            rp=E.SubElement(self.flat,bp.Q('rPr'));E.SubElement(rp,bp.Q('kern'),{bp.Q('val'):'0'})
    mode={'bad':False}
    def apply(d,s):
        d.flat.find(bp.Q('name')).set(bp.Q('val'),'Template table display name')
        if mode['bad']:d.flat.find('w:rPr/w:kern',bp.NS).set(bp.Q('val'),'2')
        return []
    monkeypatch.setitem(sys.modules,'numbering_policy',types.SimpleNamespace(Doc=FakeDoc,Q=bp.Q,PolicyError=bp.BlockError,file_digest=bp.sha))
    monkeypatch.setitem(sys.modules,'template_table_style',types.SimpleNamespace(TemplateTables=lambda *a:types.SimpleNamespace(profile_hash='fixture'),
        project=lambda d:[{'table':1,'props':[],'style':[]}],theme_colors=lambda d:{},table_sid=lambda *a:'table',style_flat=lambda d,*a:copy.deepcopy(d.flat),apply=apply))
    assert audit.audit(source,source)['status']=='passed'
    mode['bad']=True;assert audit.audit(source,source)['status']=='failed'


def test_native_refresh_uses_real_output_check_even_on_cleanup_warning(tmp_path,monkeypatch):
    """Real saved ZIP plus explicit Office/field-boundary doubles, not native COM."""
    import importlib.util
    class Error(ValueError):
        def __init__(self,code,message,**details):super().__init__(message);self.code=code;self.details=details
    def native(engine,mode,*,source,output,**kwargs):
        shutil.copyfile(source,output)
        return {'status':'passed','engine':'Microsoft Word','fields_updated':True,'indexes_updated':True,'warnings':['Quit: simulated cleanup failure']}
    def preserve(src,candidate,out):shutil.copyfile(candidate,out);return {'status':'passed'}
    monkeypatch.setitem(sys.modules,'numbering_policy',types.SimpleNamespace(file_digest=bp.sha,parse=bp.parse,fields=lambda p:[],command=lambda f:'',Q=bp.Q,ERROR_TEXT=__import__('re').compile('Error!'),PolicyError=Error))
    monkeypatch.setitem(sys.modules,'office_backends',types.SimpleNamespace(PRIORITY=('word','wps','libreoffice'),OfficeError=Error,native_operation=native,find_office=lambda:None))
    monkeypatch.setitem(sys.modules,'caption_field_guard',types.SimpleNamespace(CaptionFieldError=Error,require_valid=lambda *a,**k:{'status':'passed'},preserve_native_instructions=preserve))
    monkeypatch.setitem(sys.modules,'runtime_preflight',types.SimpleNamespace(select_engine=lambda *a,**k:{'engine':'word'},PreflightError=Error))
    spec=importlib.util.spec_from_file_location('_field_refresh_test',SCRIPTS/'field_refresh.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
    src=write_doc(tmp_path/'s.docx');out=tmp_path/'saved.docx'
    result=m._refresh_once(src,out,'word',preflight={'engine':'word'})
    assert out.is_file() and result['output_sha256']==bp.sha(out) and result['warnings']
    def corrupt(*a,**k):Path(k['output']).write_bytes(b'not DOCX');return {'status':'passed'}
    m.native_operation=corrupt
    with pytest.raises(Exception):m._refresh_once(src,tmp_path/'bad.docx','word',preflight={'engine':'word'})
    assert not (tmp_path/'bad.docx').exists()


def test_inventory_pass_has_explicit_pending_visual_status():
    source=(SCRIPTS/'review_pipeline.py').read_text()
    assert "'visual_status':'pending' if actual else 'not_applicable'" in source


def test_unchanged_final_bad_image_is_not_sent_to_model_again(tmp_path,monkeypatch):
    """Exercise the real branch; image/document/Office discovery are explicit doubles."""
    import figure_inspection as fi
    import block_evidence as be
    source=tmp_path/'source.docx';source.write_bytes(b'explicit source boundary fixture')
    candidate=tmp_path/'candidate.docx';candidate.write_bytes(b'explicit candidate boundary fixture')
    image=tmp_path/'bad.png';Image.new('RGB',(20,20)).save(image)
    ledger=tmp_path/'ledger.json';ve.write(ledger,{'inspections':['explicit failed-call boundary fixture']})
    monkeypatch.setitem(sys.modules,'figure_review',types.SimpleNamespace(validate_initial=lambda *a:{'F0003':{'object_sha256':'bad','inspection':'initial-boundary'}}))
    monkeypatch.setitem(sys.modules,'render_docx',types.SimpleNamespace(validate_render=lambda *a,**k:None))
    monkeypatch.setattr(fi,'_doc',lambda *a:object())
    monkeypatch.setattr(fi,'_figures',lambda *a:[{'id':'F0003','hash':'bad','paths':['bad.png']}])
    monkeypatch.setattr(fi,'_extract',lambda *a:[image]);monkeypatch.setattr(fi,'_page_refs',lambda *a:[])
    monkeypatch.setattr(fi,'validate_inspection',lambda *a,**k:{'binding':{'object_id':'F0003'},'verdict':'fail','checks':{'text_inside_bounds':'fail'}})
    monkeypatch.setattr(fi,'_inspection_pixels',lambda *a:[ve.pixel_sha(ve.normalized_image(image))])
    monkeypatch.setattr(be,'reuse_source_content',lambda *a,**k:False)
    monkeypatch.setattr(fi,'run',lambda *a,**k:pytest.fail('unchanged defect must not trigger another model call'))
    visual={'candidate_sha256':ve.file_sha(candidate),'figures':{'objects':[{'id':'F0003'}]}}
    manifest={'candidate_sha256':ve.file_sha(candidate),'source_sha256':ve.file_sha(source),'render':{'data':{},'pages':[]},'image_discovery_ledger':ve.reference(ledger)}
    result=fi.inspect_final(source,candidate,{},visual,manifest,None,tmp_path/'inspect')
    assert result['overall_status']=='fail'
    assert result['figures']['objects'][0]['postcheck_issue']['code']=='late-image-defect-unchanged'
