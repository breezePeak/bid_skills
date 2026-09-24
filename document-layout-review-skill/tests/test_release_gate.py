import copy
import json
import unittest
from pathlib import Path
from unittest.mock import patch
from test_figure_gate import FigureGate
import numbering_policy as n
import figure_review as f
import finalize_review as g
import field_refresh as r
from review_pipeline import REQUIRED_GATES, write_json


class ReleaseGate(FigureGate):
    def setUp(self):
        super().setUp()
        for boundary in [
            patch('block_progress.validate_final_binding',return_value={}),
            patch('figure_inspection.verify_row',return_value={}),
            patch('figure_inspection.late_findings',return_value={'issue_ids':[],'authorizes_change':False}),
            patch.object(g,'validate_render',return_value={'status':'passed','engine':'explicit-render-test-double'}),
        ]:
            boundary.start();self.addCleanup(boundary.stop)

        # The release contract now requires an actual table template, not a
        # synthetic template containing headings only.
        from docx import Document
        d=Document(self.template);table=d.add_table(rows=2,cols=2);table.style='Table Grid'
        table.cell(0,0).text='设备';table.cell(0,1).text='数量'
        table.cell(1,0).text='示例';table.cell(1,1).text='1';d.save(self.template)

    def normalize(self,plan=None):
        result=super().normalize(plan)
        import template_table_style as table_style
        temporary=self.root/'template-tables.docx'
        table_style.repair(self.out,temporary,self.template);temporary.replace(self.out)
        return result

    def bundle(self):
        a,b,pages=self.data();initial=self.root/'initial.json';write_json(initial,a)
        style=self.root/'template-style.json';write_json(style,{})
        gates=[]
        for ident in sorted(REQUIRED_GATES):
            report=self.root/(ident+'.json');result={'status':'passed','issues':[]}
            # Synthetic gate metadata fixtures; no claim that Word executed here.
            if ident=='field-update':result={'status':'passed','engine':'Microsoft Word','fields_updated':True,'indexes_updated':True,'errors':[],'output_sha256':n.file_digest(self.out)}
            write_json(report,result);gates.append({'id':ident,'passed':True,'report':str(report),'report_sha256':n.file_digest(report)})
        m={'version':5,'status':'awaiting_visual_review','candidate':str(self.out),'candidate_sha256':n.file_digest(self.out),
           'source':str(self.source),'source_sha256':n.file_digest(self.source),'template':str(self.template),'template_sha256':n.file_digest(self.template),
           'template_style_json':str(style),'template_style_sha256':n.file_digest(style),'initial_visual_review':str(initial),'initial_visual_review_sha256':n.file_digest(initial),
           'gates':gates,'render':{'passed':True,'pages':pages}}
        v={'overall_status':'pass','candidate_sha256':n.file_digest(self.out),'pages':[{'name':p['name'],'sha256':p['sha256'],'status':'pass','observations':'Synthetic test page review.'} for p in pages],'figures':b,'table_reviews':[]}
        from content_integrity import make_plan
        for key,filename,data in [
            ('content_plan','content-plan.json',make_plan(self.source)),
            ('text_rules','text-rules.json',{}),
            ('runtime_preflight','runtime-preflight.json',{'status':'passed'}),
        ]:
            path=self.root/filename;write_json(path,data);m[key]=str(path)
            m['text_rules_file_sha256' if key=='text_rules' else key+'_sha256']=n.file_digest(path)
        discovery=self.root/'image-discoveries.json'
        write_json(discovery,{'version':1,'source_sha256':m['source_sha256'],'inspections':[]})
        m['image_discovery_ledger']={'path':str(discovery),'sha256':n.file_digest(discovery)}
        v['table_objects']=[{'id':o['id'],'object_sha256':o['hash'],'status':'pass','observations':'Explicit synthetic table observation fixture.',
                            'pages':[{'name':p['name'],'sha256':p['sha256']} for p in pages]}
                           for o in n.inventory(n.Doc(self.out)) if o['kind']=='table']
        return m,v
    def test_complete_contract_passes(self):
        m,v=self.bundle();self.assertEqual(g.validate_release(m,v)['status'],'passed')
    def test_missing_gate_cannot_vacuously_pass(self):
        m,v=self.bundle();m['gates']=[]
        with self.assertRaises(n.PolicyError) as c:g.validate_release(m,v)
        self.assertEqual(c.exception.code,'required-gate-missing')
    def test_missing_image_gate_rejected(self):
        m,v=self.bundle();m['gates']=[x for x in m['gates'] if x['id']!='image-inventory']
        with self.assertRaises(n.PolicyError):g.validate_release(m,v)
    def test_duplicate_gate_rejected(self):
        m,v=self.bundle();m['gates'][-1]=m['gates'][0]
        with self.assertRaises(n.PolicyError):g.validate_release(m,v)
    def test_empty_pages_rejected(self):
        m,v=self.bundle();m['render']['pages']=[];v['pages']=[]
        with self.assertRaises(n.PolicyError):g.validate_release(m,v)
    def test_duplicate_page_review_rejected(self):
        m,v=self.bundle();v['pages']*=2
        with self.assertRaises(n.PolicyError):g.validate_release(m,v)
    def test_missing_figure_entries_rejected(self):
        m,v=self.bundle();v['figures']['objects']=[]
        with self.assertRaises(n.PolicyError):g.validate_release(m,v)
    def test_mutated_audit_json_rejected(self):
        m,v=self.bundle();Path(m['gates'][0]['report']).write_text('{}')
        with self.assertRaises(n.PolicyError):g.validate_release(m,v)
    def test_changed_template_rejected(self):
        m,v=self.bundle();m['template_sha256']='old'
        with self.assertRaises(n.PolicyError):g.validate_release(m,v)
    def test_no_observation_rejected(self):
        m,v=self.bundle();v['pages'][0]['observations']=''
        with self.assertRaises(n.PolicyError):g.validate_release(m,v)
    def test_failed_release_preserves_existing_deliverable(self):
        m,v=self.bundle();v['overall_status']='fail';final=self.root/'final.docx';final.write_bytes(b'old')
        with self.assertRaises(n.PolicyError):g.finalize(m,v,final)
        self.assertEqual(final.read_bytes(),b'old')
    def test_release_copies_exact_reviewed_bytes(self):
        m,v=self.bundle();final=self.root/'final.docx';result=g.finalize(m,v,final)
        self.assertEqual(n.file_digest(final),m['candidate_sha256'])
    def test_updatefields_flag_is_not_engine_evidence(self):
        m,v=self.bundle()
        with self.assertRaises(n.PolicyError):r.validate_report(self.out,{'status':'passed','updateFields':True,'output_sha256':n.file_digest(self.out)})
    def test_field_refresh_after_modification_is_stale(self):
        m,v=self.bundle()
        with self.assertRaises(n.PolicyError):r.validate_report(self.out,{'status':'passed','engine':'Microsoft Word','fields_updated':True,'indexes_updated':True,'errors':[],'output_sha256':'old'})
    def test_table_review_unresolved_is_blocked(self):
        m,v=self.bundle();gate=next(x for x in m['gates'] if x['id']=='table-layout')
        write_json(gate['report'],{'issues':[{'severity':'review','code':'merge-candidate','table':1}]});gate['report_sha256']=n.file_digest(Path(gate['report']))
        with self.assertRaises(n.PolicyError) as c:g.validate_release(m,v)
        self.assertEqual(c.exception.code,'table-review-missing')

    def test_release_cannot_overwrite_original_source(self):
        m,v=self.bundle();before=self.source.read_bytes()
        with self.assertRaises(n.PolicyError):g.finalize(m,v,self.source)
        self.assertEqual(self.source.read_bytes(),before)

    def test_real_template_recheck_rejects_fake_pass_report(self):
        m,v=self.bundle()
        from test_numbering_policy import rewrite
        def change(files):
            root=n.parse(files['word/document.xml']);pr=root.find('.//w:tc/w:tcPr',n.NS)
            pr.append(n.node('shd',val='clear',fill='FF0000'));files['word/document.xml']=n.dump(root)
        rewrite(self.out,change)
        h=n.file_digest(self.out);m['candidate_sha256']=h;v['candidate_sha256']=h;v['figures']['candidate_sha256']=h
        field_gate=next(x for x in m['gates'] if x['id']=='field-update')
        data=json.loads(Path(field_gate['report']).read_text());data['output_sha256']=h
        write_json(field_gate['report'],data);field_gate['report_sha256']=n.file_digest(Path(field_gate['report']))
        with self.assertRaises(n.PolicyError) as ctx:g.validate_release(m,v)
        self.assertEqual(ctx.exception.code,'final-table-template-failed')


def load_tests(loader, tests, pattern):
    return unittest.TestSuite(ReleaseGate(name) for name in ReleaseGate.__dict__ if name.startswith('test_'))

# DLR_BLOCK_RELEASE_TESTS
