#!/usr/bin/env python3
"""Optional REAL Word integration. Does not run normalization after insert/delete.

Run on Windows with Microsoft Word installed:
  python tests/word_f9_integration.py --template assets/default-template.docx --out-dir word-test
A Linux skip is reported as skipped, never passed.
"""
from __future__ import annotations
import argparse
import copy
import json
import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches
from PIL import Image

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import numbering_policy as n
from field_refresh import refresh


def check(path,template,expected):
    actual=n.audit_document(path,template)
    if actual['status']!='passed':raise RuntimeError(json.dumps(actual,ensure_ascii=False))
    caps=[n.visible(o['caption']) for o in n.inventory(n.Doc(path)) if o['kind']=='figure']
    if caps!=expected:raise AssertionError((caps,expected))
    return actual


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--template',type=Path,required=True);ap.add_argument('--out-dir',type=Path,required=True);a=ap.parse_args();a.out_dir.mkdir(parents=True,exist_ok=True)
    report={'status':'skipped','engine':'Microsoft Word','checks':[]}
    if os.name!='nt':
        report['reason']='Windows Word COM is not available on this platform.'
        (a.out_dir/'result.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report));return 0
    try:
        image=a.out_dir/'fixture.png';Image.new('RGB',(120,80),'white').save(image)
        source=a.out_dir/'source.docx';d=Document()
        for chapter in ('第一章 项目要求','第二章 实施方案'):
            d.add_heading(chapter,1);d.add_picture(str(image),width=Inches(1));d.add_paragraph('图1 原图',style='Caption')
        d.save(source);fixed=a.out_dir/'fixed.docx';n.normalize_document(source,fixed,a.template)
        first=a.out_dir/'first-refresh.docx';report['first_refresh']=refresh(fixed,first,'word')
        check(first,a.template,['图1 原图','图1 原图']);report['checks'].append('initial-refresh')
        # Add a new figure and an actual native caption field before chapter one's
        # original figure. No numbering repair/normalizer is invoked afterwards.
        d=Document(first);anchor=next(p._p for p in d.paragraphs if p._p.xpath('.//w:drawing'))
        figure=d.add_paragraph();figure.add_run().add_picture(str(image),width=Inches(1));anchor.addprevious(figure._p)
        old=n.inventory(n.Doc(first))[0]['caption'];caption=copy.deepcopy(old)
        # Change only the title of the copied caption; retain the real SEQ code.
        ts=caption.findall('.//w:t',n.NS);ts[-1].text=' 新图'
        figure._p.addnext(caption)
        inserted=a.out_dir/'inserted.docx';d.save(inserted)
        updated=a.out_dir/'inserted-refreshed.docx';report['insert_refresh']=refresh(inserted,updated,'word')
        check(updated,a.template,['图1 新图','图2 原图','图1 原图']);report['checks'].append('insert-first-caption-and-F9-without-repair')
        d=n.Doc(updated);first_obj=n.inventory(d)[0]
        first_obj['caption'].getparent().remove(first_obj['caption']);first_obj['element'].getparent().remove(first_obj['element']);d.changed.add('word/document.xml')
        deleted=a.out_dir/'deleted.docx';d.save(deleted);updated2=a.out_dir/'deleted-refreshed.docx';report['delete_refresh']=refresh(deleted,updated2,'word')
        check(updated2,a.template,['图1 原图','图1 原图']);report['checks'].append('delete-first-caption-and-F9-without-repair')
        # Renaming a chapter changes no style lookup string in any caption.
        d=Document(updated2)
        for p in d.paragraphs:
            if p.style.name.lower().startswith(('heading','标题')):
                for run in p.runs:
                    if '项目要求' in run.text:run.text=run.text.replace('项目要求','更新后的项目要求')
        renamed=a.out_dir/'renamed.docx';d.save(renamed);renamed2=a.out_dir/'renamed-refreshed.docx';report['rename_refresh']=refresh(renamed,renamed2,'word')
        check(renamed2,a.template,['图1 原图','图1 原图']);report['checks'].append('rename-heading-and-F9-without-STYLEREF')
        report['status']='passed'
    except Exception as exc:report['status']='failed';report['error']=str(exc)
    (a.out_dir/'result.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2));return 0 if report['status']=='passed' else 2


if __name__=='__main__':raise SystemExit(main())
