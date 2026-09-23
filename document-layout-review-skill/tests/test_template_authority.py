"""Regressions: normalize the input TO the template, then protect that result."""
from __future__ import annotations
import copy
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches
from PIL import Image
from lxml import etree as E

SCRIPTS=Path(__file__).resolve().parents[1]/'scripts';sys.path.insert(0,str(SCRIPTS))
import numbering_policy as n
import template_table_style as t
import layout_invariant_guard as guard
from test_numbering_policy import rewrite,add_native_numbering


def xchange(path,fn,part='word/document.xml'):
    def transform(files):
        root=n.parse(files[part]);fn(root);files[part]=n.dump(root)
    rewrite(path,transform)


def prop(parent,kind,tag,**attrs):
    pp=parent.find(n.Q(kind))
    if pp is None:pp=n.node(kind);parent.insert(0,pp)
    old=pp.find(n.Q(tag))
    if old is not None:pp.remove(old)
    pp.append(n.node(tag,**attrs))


class TemplateAuthority(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.template=self.root/'template.docx';self.source=self.root/'source.docx';self.out=self.root/'out.docx'
        self.make(self.template);self.make(self.source)

    def tearDown(self):self.temp.cleanup()

    def make(self,path,cols=2,rows=3):
        d=Document();d.add_heading('项目方案',1);d.add_paragraph('表1 设备配置',style='Caption')
        table=d.add_table(rows=rows,cols=cols);table.style='Table Grid'
        for i,row in enumerate(table.rows):
            for j,cell in enumerate(row.cells):cell.text=(f'字段{j+1}' if i==0 else f'内容{i}-{j+1}')
        d.add_paragraph('不得修改的正文。');d.save(path)

    def gray(self,path,color='D9E1F2'):
        xchange(path,lambda root:[prop(c,'tcPr','shd',val='clear',fill=color) for c in root.find('.//w:tr',n.NS).findall(n.Q('tc'))])

    def repair(self,profile=None):return t.repair(self.source,self.out,self.template,profile)

    def doc(self,path=None):return n.Doc(path or self.out)

    def first_cell(self,path=None):return self.doc(path).document.find('.//w:tc',n.NS)

    def test_gray_header_is_removed_when_template_has_no_fill(self):
        self.gray(self.source);r=self.repair();self.assertEqual(r['status'],'passed')
        self.assertIsNone(self.first_cell().find('w:tcPr/w:shd',n.NS))
        self.assertEqual(t.audit(self.out,self.template)['status'],'passed')

    def test_template_blue_header_is_added_to_plain_input(self):
        self.gray(self.template,'204060');self.repair()
        self.assertEqual(n.value(self.first_cell().find('w:tcPr/w:shd',n.NS),'fill'),'204060')

    def test_input_blue_does_not_override_template_gray(self):
        self.gray(self.template,'C0C0C0');self.gray(self.source,'0000FF');self.repair()
        self.assertEqual(n.value(self.first_cell().find('w:tcPr/w:shd',n.NS),'fill'),'C0C0C0')

    def test_style_based_source_fill_is_removed(self):
        d=Document(self.source);d.tables[0].style='Light Shading Accent 1';d.save(self.source);self.repair()
        self.assertTrue(t.table_sid(self.doc(),self.doc().document.find('.//w:tbl',n.NS)).startswith('DLRT_'))
        self.assertEqual(t.audit(self.out,self.template)['status'],'passed')

    def test_template_conditional_header_is_retained(self):
        d=Document(self.template);d.tables[0].style='Light Shading Accent 1';d.save(self.template);self.repair()
        doc=self.doc();sid=t.table_sid(doc,doc.document.find('.//w:tbl',n.NS))
        self.assertIsNotNone(doc.style_map[sid].find('w:tblStylePr[@w:type="firstRow"]',n.NS))

    def test_template_white_header_text_is_retained(self):
        def change(root):
            s=root.find('w:style[@w:styleId="TableGrid"]',n.NS)
            cond=n.node('tblStylePr',type='firstRow');rp=n.node('rPr');rp.append(n.node('color',val='FFFFFF'));cond.append(rp);s.append(cond)
        xchange(self.template,change,'word/styles.xml');self.repair()
        self.assertEqual(n.value(self.first_cell().find('w:p/w:r/w:rPr/w:color',n.NS)),'FFFFFF')

    def test_paragraph_and_run_shading_are_removed(self):
        def change(root):
            cell=root.find('.//w:tc',n.NS);p=cell.find(n.Q('p'));r=p.find(n.Q('r'))
            prop(p,'pPr','shd',fill='123456');prop(r,'rPr','shd',fill='ABCDEF');prop(r,'rPr','highlight',val='yellow')
        xchange(self.source,change);self.repair();c=self.first_cell()
        self.assertEqual(n.value(c.find('w:p/w:pPr/w:shd',n.NS)),'nil')
        self.assertEqual(n.value(c.find('w:p/w:r/w:rPr/w:highlight',n.NS)),'none')

    def test_inherited_character_color_is_overridden_not_globally_deleted(self):
        def change(root):
            s=n.node('style',type='character',styleId='RedText');rp=n.node('rPr');rp.append(n.node('color',val='FF0000'));s.append(rp);root.append(s)
        xchange(self.source,change,'word/styles.xml')
        xchange(self.source,lambda root:prop(root.find('.//w:tc/w:p/w:r',n.NS),'rPr','rStyle',val='RedText'))
        self.repair();self.assertEqual(n.value(self.first_cell().find('w:p/w:r/w:rPr/w:color',n.NS)),'auto')
        self.assertIn('RedText',self.doc().style_map)

    def test_row_exception_shading_does_not_survive(self):
        xchange(self.source,lambda root:prop(root.find('.//w:tr',n.NS),'tblPrEx','shd',fill='0000FF'))
        self.repair();self.assertIsNone(self.doc().document.find('.//w:tblPrEx/w:shd',n.NS))

    def test_row_exception_from_template_is_applied(self):
        xchange(self.template,lambda root:prop(root.find('.//w:tr',n.NS),'tblPrEx','shd',fill='EEEEEE'))
        self.repair();self.assertEqual(n.value(self.doc().document.find('.//w:tblPrEx/w:shd',n.NS),'fill'),'EEEEEE')

    def test_borders_come_from_template(self):
        def change(root):
            c=root.find('.//w:tc',n.NS);cp=c.find(n.Q('tcPr'));b=n.node('tcBorders');b.append(n.node('top',val='double',color='FF0000'));cp.append(b)
        xchange(self.source,change);self.repair()
        self.assertIsNone(self.first_cell().find('w:tcPr/w:tcBorders',n.NS))
        borders=self.doc().document.find('.//w:tblPr/w:tblBorders',n.NS)
        self.assertEqual(n.value(borders.find(n.Q('top'))),'single')

    def test_no_border_is_a_real_template_state(self):
        d=Document(self.template);d.tables[0].style='Normal Table';d.save(self.template);self.repair()
        self.assertEqual(n.value(self.doc().document.find('.//w:tblPr/w:tblBorders/w:top',n.NS)),'nil')

    def test_widths_and_content_are_not_changed_by_appearance_repair(self):
        self.gray(self.source);before=guard.snapshot_docx(self.source);self.repair();after=guard.snapshot_docx(self.out)
        self.assertEqual(guard.changed_invariants(before,after,'template-table-style'),[])
        self.assertEqual(before['table_nonappearance'],after['table_nonappearance'])
        self.assertEqual(before['text_content'],after['text_content'])

    def test_merged_cells_are_not_restructured(self):
        d=Document(self.source);d.tables[0].cell(1,0).merge(d.tables[0].cell(2,0));d.save(self.source)
        before=guard.snapshot_docx(self.source);self.repair();self.assertEqual(before['table_structure'],guard.snapshot_docx(self.out)['table_structure'])

    def test_different_column_count_does_not_copy_template_grid(self):
        self.make(self.source,cols=5);self.gray(self.source);self.repair()
        self.assertEqual(len(self.doc().document.find('.//w:tblGrid',n.NS)),5)

    def test_no_table_document_does_not_require_sample(self):
        d=Document();d.add_paragraph('普通文本');d.save(self.source);d.save(self.template)
        self.assertEqual(self.repair()['change_count'],0)

    def test_missing_template_table_is_not_silent_preservation(self):
        d=Document();d.add_paragraph('没有表格示例');d.save(self.template)
        with self.assertRaises(n.PolicyError):self.repair()
        self.assertFalse(self.out.exists())

    def test_ambiguous_template_styles_are_not_selected_by_frequency(self):
        d=Document(self.template);a=d.add_table(rows=2,cols=2);a.style='Light Shading Accent 1'
        for row in a.rows:
            for c in row.cells:c.text='不同字段'
        d.save(self.template);self.make(self.source,cols=3)
        with self.assertRaises(n.PolicyError) as ctx:self.repair()
        self.assertEqual(ctx.exception.code,'template-table-style-ambiguous')

    def test_explicit_template_style_resolves_multiple_examples(self):
        d=Document(self.template);d.add_table(rows=2,cols=2).style='Light Shading Accent 1';d.save(self.template)
        profile=self.root/'profile.json';profile.write_text(json.dumps({'table':{'style_id':'TableGrid'}}))
        self.make(self.source,cols=4);self.gray(self.source);self.assertEqual(self.repair(profile)['status'],'passed')

    def test_exact_template_headers_choose_corresponding_style(self):
        d=Document(self.template);other=d.add_table(rows=2,cols=3);other.style='Light Shading Accent 1'
        for i,row in enumerate(other.rows):
            for j,c in enumerate(row.cells):c.text=f'其他{j}'
        d.save(self.template);self.repair();self.assertEqual(t.audit(self.out,self.template)['status'],'passed')

    def test_repeated_run_is_idempotent(self):
        self.gray(self.source);self.repair();again=self.root/'again.docx'
        r=t.repair(self.out,again,self.template);self.assertEqual(r['change_count'],0)
        self.assertEqual(t.project(self.doc()),t.project(n.Doc(again)))

    def test_later_width_guard_preserves_normalized_not_source_appearance(self):
        self.gray(self.source);before=guard.snapshot_docx(self.source);self.repair();normalized=guard.snapshot_docx(self.out)
        self.assertIn('table_appearance',guard.changed_invariants(before,normalized,'table-layout'))
        xchange(self.out,lambda root:root.find('.//w:gridCol',n.NS).set(n.Q('w'),'1200'))
        self.assertEqual(guard.changed_invariants(normalized,guard.snapshot_docx(self.out),'table-layout'),[])

    def test_later_recolor_is_rejected(self):
        self.repair();normalized=guard.snapshot_docx(self.out);self.gray(self.out)
        self.assertIn('table_appearance',guard.changed_invariants(normalized,guard.snapshot_docx(self.out),'table-layout'))
        self.assertEqual(t.audit(self.out,self.template)['status'],'failed')

    def test_style_definition_tampering_is_rejected(self):
        self.repair()
        def change(root):
            s=next(s for s in root.findall(n.Q('style')) if n.value(s,'styleId','').startswith('DLRT_'))
            prop(s,'tcPr','shd',fill='FF0000')
        xchange(self.out,change,'word/styles.xml');self.assertEqual(t.audit(self.out,self.template)['status'],'failed')

    def test_image_media_theme_and_body_are_preserved(self):
        image=self.root/'img.png';Image.new('RGB',(50,30),'green').save(image)
        d=Document(self.source);d.add_picture(str(image),width=Inches(1));d.save(self.source)
        before=self.doc(self.source);self.repair();after=self.doc()
        for part in before.files:
            if part not in {'word/styles.xml','word/document.xml'}:self.assertEqual(before.files[part],after.files[part],part)
        self.assertEqual(guard.snapshot_docx(self.source)['drawings'],guard.snapshot_docx(self.out)['drawings'])

    def test_template_theme_colors_are_resolved_without_changing_document_theme(self):
        d=Document(self.template);d.tables[0].style='Light Shading Accent 1';d.save(self.template)
        def recolor(root):
            for x in root.findall('.//a:accent1/a:srgbClr',n.NS):x.set('val','884422')
        xchange(self.template,recolor,'word/theme/theme1.xml')
        before=self.doc(self.source).files['word/theme/theme1.xml'];self.repair();after=self.doc()
        self.assertEqual(before,after.files['word/theme/theme1.xml'])
        sid=t.table_sid(after,after.document.find('.//w:tbl',n.NS))
        self.assertNotIn('themeFill',E.tostring(after.style_map[sid]).decode())

    def test_inplace_repair_is_refused(self):
        with self.assertRaises(n.PolicyError):t.repair(self.source,self.source,self.template)

    def test_failure_preserves_existing_deliverable(self):
        self.out.write_bytes(b'previous deliverable');d=Document();d.save(self.template)
        with self.assertRaises(n.PolicyError):self.repair()
        self.assertEqual(self.out.read_bytes(),b'previous deliverable')

    def test_template_sample_is_not_modified(self):
        before=self.template.read_bytes();self.gray(self.source);self.repair();self.assertEqual(before,self.template.read_bytes())

    def test_caption_below_table_moves_above_in_existing_numbering_workflow(self):
        add_native_numbering(self.template)
        d=Document(self.source);d.paragraphs[1]._p.getparent().remove(d.paragraphs[1]._p)
        d.add_paragraph('表1 设备配置',style='Caption');d.save(self.source)
        # Put the caption immediately after the table, not after unrelated prose.
        def move(root):
            table=root.find('.//w:tbl',n.NS);cap=[p for p in root.findall('.//w:p',n.NS) if n.visible(p).startswith('表1')][0]
            cap.getparent().remove(cap);table.addnext(cap)
        xchange(self.source,move)
        numbered=self.root/'numbered.docx';n.normalize_document(self.source,numbered,self.template)
        t.repair(numbered,self.out,self.template)
        doc=self.doc();tab=doc.document.find('.//w:tbl',n.NS)
        self.assertTrue(n.visible(tab.getprevious()).startswith('表1'))
        self.assertEqual(n.audit_document(self.out,self.template)['status'],'passed')

    def test_pipeline_has_required_template_gate(self):
        from review_pipeline import REQUIRED_GATES
        self.assertIn('table-template',REQUIRED_GATES)

    def test_pipeline_normalizes_template_before_width_repair(self):
        source=(SCRIPTS/'review_pipeline.py').read_text()
        self.assertLess(source.index("stage('04a-template-tables'"),source.index("stage('05-tables'"))
        self.assertIn("'template_style_enforce.py':'template-text-style'",source)

    def test_actual_width_script_keeps_template_colors(self):
        from table_layout_repair import repair_document_xml,write_docx
        self.gray(self.source);self.repair();before=guard.snapshot_docx(self.out)
        candidate=self.root/'widths.docx'
        xml,_=repair_document_xml(self.doc().files['word/document.xml'],self.template)
        write_docx(self.out,candidate,xml)
        self.assertEqual(guard.changed_invariants(before,guard.snapshot_docx(candidate),'table-layout'),[])
        self.assertEqual(t.audit(candidate,self.template)['status'],'passed')

    def test_template_scope_does_not_allow_editing_content(self):
        self.repair();before=guard.snapshot_docx(self.out)
        xchange(self.out,lambda root:setattr(root.find('.//w:tc/w:p/w:r/w:t',n.NS),'text','被改写的内容'))
        self.assertIn('text_content',guard.changed_invariants(before,guard.snapshot_docx(self.out),'template-table-style'))

    def test_tint_uses_word_retained_luminance(self):
        shd=n.node('shd',themeFill='accent1',themeFillTint='99')
        result=t.resolved_colors(shd,{'accent1':'4F81BD'})
        self.assertEqual(n.value(result,'fill'),'95B3D7')

    def test_missing_theme_color_blocks_without_guess(self):
        shd=n.node('shd',themeFill='accent1')
        with self.assertRaises(n.PolicyError):t.resolved_colors(shd,{})

    def test_stale_profile_is_not_used_with_new_template(self):
        profile=self.root/'profile.json';profile.write_text(json.dumps({'template':{'canonical_template_sha256':'sha256:old'},'table':{'style_id':'TableGrid'}}))
        with self.assertRaises(n.PolicyError) as ctx:self.repair(profile)
        self.assertEqual(ctx.exception.code,'template-profile-stale')


if __name__=='__main__':unittest.main()
