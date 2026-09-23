"""Focused regression tests; external Office availability is explicitly mocked.

Run: python -m unittest discover -s tests -p 'test_template_contract_regression.py' -v
"""
from __future__ import annotations
import copy
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as E

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import template_format_contract as fmt
import template_layout_contract as layout
import content_integrity as content
import runtime_preflight as env
import table_layout_repair as widths
import docx_layout_policy as policy


def property_node(name, **attrs):
    n = OxmlElement('w:' + name)
    for key, value in attrs.items():
        n.set(qn('w:' + key), str(value))
    return n


def set_property(owner, kind, name, **attrs):
    pr = owner.find(qn('w:' + kind))
    if pr is None:
        pr = OxmlElement('w:' + kind); owner.insert(0, pr)
    for old in list(pr.findall(qn('w:' + name))):
        pr.remove(old)
    n = property_node(name, **attrs); pr.append(n)
    return n


def package_replace(source, target, part, data):
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info, data if info.filename == part else src.read(info.filename))


def bytes_part(path, name):
    with zipfile.ZipFile(path) as z:
        return z.read(name)


class TemporaryTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def document(self, name='source.docx', text='工期30天，数量2套。'):
        d = Document()
        if text is not None:
            d.add_paragraph(text)
        p = self.root / name; d.save(p)
        return p


class FormatTests(TemporaryTest):
    def resolver(self, document):
        p = self.root / 'styles.docx'; document.save(p)
        return fmt.StyleResolver(bytes_part(p, 'word/styles.xml'))

    def test_missing_parent_is_blocked(self):
        xml = f'<w:styles xmlns:w="{fmt.W}"><w:style w:type="paragraph" w:styleId="Child"><w:basedOn w:val="Missing"/></w:style></w:styles>'.encode()
        with self.assertRaises(fmt.FormatContractError):
            fmt.StyleResolver(xml).chain('Child')

    def test_inheritance_cycle_is_blocked(self):
        xml = f'<w:styles xmlns:w="{fmt.W}"><w:style w:styleId="A"><w:basedOn w:val="B"/></w:style><w:style w:styleId="B"><w:basedOn w:val="A"/></w:style></w:styles>'.encode()
        with self.assertRaises(fmt.FormatContractError):
            fmt.StyleResolver(xml).chain('A')

    def test_parent_and_linked_style_are_copied(self):
        d = Document()
        base = d.styles.add_style('Template Parent', WD_STYLE_TYPE.PARAGRAPH)
        child = d.styles.add_style('Template Child', WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = base
        character = d.styles.add_style('Template Character', WD_STYLE_TYPE.CHARACTER)
        child.element.append(property_node('link', val=character.style_id))
        expected = self.resolver(d)
        target = self.resolver(Document())
        result = fmt.StyleResolver(fmt.copy_style_dependencies(fmt.serialize(target.root), fmt.serialize(expected.root), [child.style_id]))
        for sid in (child.style_id, base.style_id, character.style_id):
            self.assertEqual(fmt.signature(result.styles[sid]), fmt.signature(expected.styles[sid]))

    def test_next_style_self_reference_is_legal(self):
        d = Document()
        st = d.styles.add_style('Self Next', WD_STYLE_TYPE.PARAGRAPH)
        st.element.append(property_node('next', val=st.style_id))
        self.assertIn(st.style_id, self.resolver(d).dependency_ids([st.style_id]))

    def test_partial_property_inheritance(self):
        d = Document()
        base = d.styles.add_style('Parent', WD_STYLE_TYPE.PARAGRAPH)
        child = d.styles.add_style('Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
        set_property(base.element, 'pPr', 'ind', left=120, right=240)
        set_property(child.element, 'pPr', 'ind', left=360)
        ind = self.resolver(d).properties('pPr', child.style_id).find(fmt.Q('ind'))
        self.assertEqual(ind.get(fmt.Q('left')), '360')
        self.assertEqual(ind.get(fmt.Q('right')), '240')

    def test_style_toggle_and_direct_override(self):
        d = Document()
        base = d.styles.add_style('Parent', WD_STYLE_TYPE.PARAGRAPH)
        child = d.styles.add_style('Child', WD_STYLE_TYPE.PARAGRAPH); child.base_style = base
        set_property(base.element, 'rPr', 'caps', val='1')
        set_property(child.element, 'rPr', 'caps', val='1')
        resolver = self.resolver(d)
        self.assertFalse(fmt.value_key(resolver.properties('rPr', child.style_id), 'caps'))
        direct = E.Element(fmt.Q('rPr')); direct.append(property_node('caps', val='1'))
        self.assertTrue(fmt.value_key(resolver.properties('rPr', child.style_id, direct), 'caps'))

    def test_dirty_body_extra_formats_are_removed_and_audited(self):
        d = Document(); p = d.add_paragraph('合同工期30天'); run = p.runs[0]
        resolver = self.resolver(d)
        for name, attrs in [('color', {'val': 'FF0000'}), ('u', {'val': 'single'}),
                            ('highlight', {'val': 'yellow'}), ('strike', {'val': '1'}),
                            ('caps', {'val': '1'}), ('spacing', {'val': '25'})]:
            set_property(run._r, 'rPr', name, **attrs)
        self.assertTrue(fmt.audit_paragraph(p._p, 'Normal', resolver, resolver))
        fmt.repair_paragraph(p._p, 'Normal', resolver, resolver)
        self.assertEqual(fmt.audit_paragraph(p._p, 'Normal', resolver, resolver), [])
        self.assertEqual(p.text, '合同工期30天')

    def test_template_red_and_underline_are_retained(self):
        d = Document()
        set_property(d.styles['Normal'].element, 'rPr', 'color', val='C00000')
        set_property(d.styles['Normal'].element, 'rPr', 'u', val='single')
        p = d.add_paragraph('模板强调'); resolver = self.resolver(d)
        set_property(p.runs[0]._r, 'rPr', 'color', val='0000FF')
        fmt.repair_paragraph(p._p, 'Normal', resolver, resolver)
        self.assertEqual(fmt.audit_paragraph(p._p, 'Normal', resolver, resolver), [])
        self.assertEqual(p.runs[0]._r.find('w:rPr/w:color', fmt.NS).get(fmt.Q('val')), 'C00000')

    def test_body_heiti_exception_not_overwritten_by_extra_pass(self):
        d = Document(); p = d.add_paragraph('正文黑体'); resolver = self.resolver(d)
        set_property(p.runs[0]._r, 'rPr', 'rFonts', eastAsia='黑体')
        set_property(p.runs[0]._r, 'rPr', 'b', val='1')
        fmt.repair_paragraph(p._p, 'Normal', resolver, resolver)
        self.assertEqual(p.runs[0]._r.find('w:rPr/w:rFonts', fmt.NS).get(fmt.Q('eastAsia')), '黑体')
        self.assertEqual(p.runs[0]._r.find('w:rPr/w:b', fmt.NS).get(fmt.Q('val')), '1')

    def test_foreign_character_style_cannot_smuggle_color(self):
        base = Document(); expected = self.resolver(base)
        bad = Document(); s = bad.styles.add_style('Source Red', WD_STYLE_TYPE.CHARACTER)
        set_property(s.element, 'rPr', 'color', val='FF0000')
        p = bad.add_paragraph('正文'); p.runs[0].style = s; current = self.resolver(bad)
        self.assertTrue(fmt.audit_paragraph(p._p, 'Normal', expected, current))
        fmt.repair_paragraph(p._p, 'Normal', expected, current)
        self.assertIsNone(p.runs[0]._r.find('w:rPr/w:rStyle', fmt.NS))

    def test_run_splitting_does_not_change_business_text(self):
        d = Document(); p = d.add_paragraph(); p.add_run('工期'); p.add_run('30'); p.add_run('天')
        resolver = self.resolver(d)
        set_property(p.runs[1]._r, 'rPr', 'color', val='FF0000')
        fmt.repair_paragraph(p._p, 'Normal', resolver, resolver)
        self.assertEqual(p.text, '工期30天')

    def test_audit_detects_unmapped_story_paragraph(self):
        template = self.document('template.docx')
        d = Document(); st = d.styles.add_style('Foreign Body', WD_STYLE_TYPE.PARAGRAPH)
        d.add_paragraph('未映射', style=st); target = self.root / 'target.docx'; d.save(target)
        self.assertEqual(fmt.audit_extra_formats(template, target)['status'], 'failed')


class TableTests(TemporaryTest):
    def table_doc(self, name, align='top', table_align='left', repeat=False, margin=180):
        d = Document(); table = d.add_table(rows=2, cols=2)
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                cell.text = [['名称', '说明'], ['服务', '按要求完成']][i][j]
                set_property(cell._tc, 'tcPr', 'vAlign', val=align)
        set_property(table._tbl, 'tblPr', 'jc', val=table_align)
        mar = OxmlElement('w:tblCellMar'); mar.append(property_node('left', w=margin, type='dxa'))
        pr = table._tbl.tblPr
        for n in list(pr.findall(qn('w:tblCellMar'))): pr.remove(n)
        pr.append(mar)
        if repeat:
            set_property(table.rows[0]._tr, 'trPr', 'tblHeader', val='1')
        p = self.root / name; d.save(p); return p

    def test_template_top_alignment_is_applied_not_centered(self):
        template = self.table_doc('template.docx', 'top')
        target = self.table_doc('target.docx', 'center')
        updated, report = policy.repair_document_xml(bytes_part(target, 'word/document.xml'), template=template)
        out = self.root / 'out.docx'; policy.write_docx(target, out, updated)
        self.assertGreater(report['change_count'], 0)
        self.assertEqual(layout.audit(template, out)['status'], 'passed')
        for node in E.fromstring(updated).xpath('.//w:vAlign', namespaces=fmt.NS):
            self.assertEqual(node.get(fmt.Q('val')), 'top')

    def test_template_bottom_alignment_is_supported(self):
        template = self.table_doc('template.docx', 'bottom')
        target = self.table_doc('target.docx', 'center')
        updated, _ = policy.repair_document_xml(bytes_part(target, 'word/document.xml'), template=template)
        out = self.root / 'out.docx'; policy.write_docx(target, out, updated)
        self.assertEqual(layout.audit(template, out)['status'], 'passed')

    def test_template_alignment_margins_and_header_survive_width_repair(self):
        template = self.table_doc('template.docx', 'top', 'right', True, 220)
        source = self.table_doc('source.docx', 'center', 'center', False, 90)
        normalized, _ = policy.repair_document_xml(bytes_part(source, 'word/document.xml'), template=template)
        reflowed, _ = widths.repair_document_xml(normalized, template)
        out = self.root / 'out.docx'; policy.write_docx(source, out, reflowed)
        self.assertEqual(layout.audit(template, out)['status'], 'passed')
        self.assertEqual(content.audit(source, out)['status'], 'passed')

    def test_no_repeat_header_in_template_removes_old_repeat(self):
        template = self.table_doc('template.docx', repeat=False)
        target = self.table_doc('target.docx', repeat=True)
        updated, _ = policy.repair_document_xml(bytes_part(target, 'word/document.xml'), template=template)
        out = self.root / 'out.docx'; policy.write_docx(target, out, updated)
        self.assertEqual(layout.audit(template, out)['status'], 'passed')
        self.assertFalse(E.fromstring(updated).xpath('.//w:tblHeader', namespaces=fmt.NS))

    def test_layout_audit_does_not_accept_center_if_template_is_top(self):
        template = self.table_doc('template.docx', 'top')
        target = self.table_doc('target.docx', 'center')
        self.assertEqual(layout.audit(template, target)['status'], 'failed')

    def test_two_ambiguous_table_templates_block_instead_of_guess(self):
        template = self.table_doc('template.docx', 'top')
        d = Document(template)
        t = d.add_table(rows=2, cols=2)
        for row in t.rows:
            for c in row.cells:
                c.text = '其他'; set_property(c._tc, 'tcPr', 'vAlign', val='bottom')
        d.save(template)
        source = self.table_doc('source.docx')
        d = Document(source); d.tables[0].cell(0, 0).text = '新表头'; d.save(source)
        with self.assertRaises(fmt.FormatContractError):
            layout.audit(template, source)

    def test_top_level_table_uses_its_actual_section_width(self):
        xml = f'''<w:document xmlns:w="{fmt.W}"><w:body><w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="1"/></w:tblGrid><w:tr><w:tc><w:p><w:r><w:t>数据</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:sectPr><w:pgSz w:w="14000" w:h="16000"/><w:pgMar w:left="1000" w:right="1000"/></w:sectPr></w:body></w:document>'''.encode()
        _, report = widths.repair_document_xml(xml)
        self.assertEqual(report['tables'][0]['content_width_twips'], 12000)

    def test_cli_applies_template_without_modifying_source(self):
        template = self.table_doc('template.docx', 'top')
        source = self.table_doc('source.docx', 'center')
        old = source.read_bytes(); out = self.root / 'cli.docx'
        cp = subprocess.run([sys.executable, str(SCRIPTS/'docx_layout_policy.py'), str(source),
                             '--template', str(template), '--output', str(out)], capture_output=True, text=True)
        self.assertEqual(cp.returncode, 0, cp.stderr + cp.stdout)
        self.assertEqual(source.read_bytes(), old)
        self.assertEqual(layout.audit(template, out)['status'], 'passed')


class ContentTests(TemporaryTest):
    def test_business_number_change_is_blocked(self):
        source = self.document('source.docx', '工期30天')
        target = self.document('target.docx', '工期3天')
        self.assertEqual(content.audit(source, target)['status'], 'failed')

    def test_business_word_deletion_is_blocked(self):
        source = self.document('source.docx', '不得泄露资料')
        target = self.document('target.docx', '得泄露资料')
        self.assertEqual(content.audit(source, target)['status'], 'failed')

    def test_whitespace_and_punctuation_cleanup_are_allowed(self):
        source = self.document('source.docx', '软件 1 套,交付日期2026 年 10 月 31 日。')
        target = self.document('target.docx', '软件1套，交付日期2026年10月31日。')
        self.assertEqual(content.audit(source, target)['status'], 'passed')

    def test_english_word_boundary_cannot_be_removed(self):
        source = self.document('source.docx', 'API key')
        target = self.document('target.docx', 'APIkey')
        self.assertEqual(content.audit(source, target)['status'], 'failed')

    def test_latin_run_boundaries_are_not_business_changes(self):
        source = self.document('source.docx', 'API key')
        d = Document(); p = d.add_paragraph(); p.add_run('API '); p.add_run('key')
        target = self.root/'target.docx'; d.save(target)
        self.assertEqual(content.audit(source, target)['status'], 'passed')

    def test_decimal_and_negative_sign_are_protected(self):
        for old, new in [('阈值-0.5', '阈值0.5'), ('精度0.05米', '精度0.5米')]:
            source = self.document('source.docx', old); target = self.document('target.docx', new)
            self.assertEqual(content.audit(source, target)['status'], 'failed')

    def test_confirmed_heading_prefix_can_become_automatic(self):
        d = Document(); d.add_heading('1 技术方案', 1); source = self.root/'source.docx'; d.save(source)
        d = Document(); d.add_heading('技术方案', 1); target = self.root/'target.docx'; d.save(target)
        self.assertEqual(content.audit(source, target)['status'], 'passed')

    def test_year_at_heading_start_is_business_content(self):
        d = Document(); d.add_heading('2026年度实施计划', 1); source = self.root/'source.docx'; d.save(source)
        d = Document(); d.add_heading('年度实施计划', 1); target = self.root/'target.docx'; d.save(target)
        self.assertEqual(content.audit(source, target)['status'], 'failed')

    def test_page_field_cache_change_is_allowed(self):
        source = self.document('source.docx')
        d = Document(source); p = d.sections[0].footer.paragraphs[0]
        f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), ' PAGE ')
        r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text='1'; r.append(t); f.append(r); p._p.append(f)
        d.save(source)
        target=self.root/'target.docx'
        xml=bytes_part(source, 'word/footer1.xml').replace(b'>1<', b'>2<')
        package_replace(source,target,'word/footer1.xml',xml)
        self.assertEqual(content.audit(source,target)['status'],'passed')

    def test_header_business_text_is_protected(self):
        source = self.document('source.docx')
        d=Document(source);d.sections[0].header.paragraphs[0].text='采购项目A';d.save(source)
        d.sections[0].header.paragraphs[0].text='采购项目B';target=self.root/'target.docx';d.save(target)
        self.assertEqual(content.audit(source,target)['status'],'failed')

    def test_known_directional_reference_repair_keeps_business_text(self):
        source=self.document('source.docx','实施路径见图1，工期30天。')
        target=self.document('target.docx','实施路径见下图，工期30天。')
        self.assertEqual(content.audit(source,target)['status'],'passed')

    def test_stale_content_plan_is_rejected(self):
        source=self.document('source.docx');target=self.document('target.docx')
        with self.assertRaises(content.ContentIntegrityError):
            content.audit(source,target,{'version':1,'source_sha256':'sha256:stale'})

    def test_plain_body_numbered_item_is_not_treated_as_heading(self):
        source=self.document('source.docx','30 天完成采购')
        target=self.document('target.docx','天完成采购')
        self.assertEqual(content.audit(source,target)['status'],'failed')

    def test_unapproved_caption_title_change_is_blocked(self):
        source=self.document('source.docx',None);d=Document(source)
        d.add_paragraph('图1 系统架构',style='Caption');d.save(source)
        target=self.root/'target.docx';d.paragraphs[0].text='图1 错误题名';d.save(target)
        self.assertEqual(content.audit(source,target)['status'],'failed')


class PreflightTests(TemporaryTest):
    @patch.object(env, 'probe_libreoffice', return_value={'available':True,'uno_python':'/example/uno'})
    @patch.object(env, 'probe_word', return_value={'available':False})
    def test_auto_selects_available_uno(self, word, libreoffice):
        self.assertEqual(env.select_engine()['engine'],'libreoffice')

    @patch.object(env, 'probe_libreoffice', return_value={'available':True,'uno_python':'/example/uno'})
    @patch.object(env, 'probe_word', return_value={'available':False})
    def test_explicit_word_never_falls_back(self, word, libreoffice):
        with self.assertRaises(env.PreflightError):env.select_engine('word')
        libreoffice.assert_not_called()

    @patch.object(env, 'probe_libreoffice', return_value={'available':True,'uno_python':'/example/uno'})
    @patch.object(env, 'probe_word', return_value={'available':True})
    def test_auto_prefers_word_when_usable(self, word, libreoffice):
        self.assertEqual(env.select_engine()['engine'],'word')
        libreoffice.assert_not_called()

    @patch.object(env, 'probe_libreoffice', return_value={'available':False})
    @patch.object(env, 'probe_word', return_value={'available':False})
    def test_no_engine_is_blocking(self, word, libreoffice):
        with self.assertRaises(env.PreflightError):env.select_engine()

    def test_bad_docx_is_blocked_before_work(self):
        bad=self.root/'bad.docx';bad.write_bytes(b'not a docx')
        with self.assertRaises(env.PreflightError):env.validate_docx(bad)

    def test_stale_profile_is_rejected(self):
        template=self.document('template.docx');profile=self.root/'profile.json'
        profile.write_text(json.dumps({'template':{'canonical_template_sha256':'sha256:stale'}}))
        with self.assertRaises(env.PreflightError):env.validate_profile(template,profile)

    @patch.object(env, 'find_office', return_value=None)
    def test_missing_renderer_is_blocked_without_modifying_documents(self, office):
        source=self.document('source.docx');template=self.document('template.docx');old=source.read_bytes()
        with self.assertRaises(env.PreflightError):env.check(source,template,self.root/'work')
        self.assertEqual(old,source.read_bytes())




class MoreContentTests(TemporaryTest):
    def test_spaced_year_in_heading_is_not_a_numbering_prefix(self):
        d=Document();d.add_heading('2026 年度技术方案',1);source=self.root/'source.docx';d.save(source)
        d=Document();d.add_heading('2025年度技术方案',1);target=self.root/'target.docx';d.save(target)
        self.assertEqual(content.audit(source,target)['status'],'failed')

    def test_actual_inventory_ids_are_accepted(self):
        source=self.document()
        p={'version':1,'source_sha256':content.digest(source),'objects':[
            {'id':'F0001','object_sha256':'sha256:test','title':'系统架构'},
            {'id':'T0001','object_sha256':'sha256:test2','title':'资源清单'}]}
        data=content.make_plan(source,p)
        self.assertEqual([r['kind'] for r in data['allowed_caption_insertions']],['figure','table'])

    def native_note(self,body='参数精度0.05米'):
        d=Document();p=d.add_paragraph('说明');r=p.add_run('1');r.font.superscript=True
        d.add_paragraph('1 '+body)
        source=self.root/'source.docx';d.save(source)
        root=fmt.parse(bytes_part(source,'word/document.xml'))
        paras=list(root.iter(fmt.Q('p')))
        marker=paras[0].findall(fmt.Q('r'))[1]
        for t in list(marker.findall(fmt.Q('t'))):marker.remove(t)
        marker.append(property_node('footnoteReference',id='1'))
        note=copy.deepcopy(paras[1]);note.find('.//'+fmt.Q('t')).text=' '+body
        paras[1].getparent().remove(paras[1])
        notes=E.Element(fmt.Q('footnotes'),nsmap={'w':fmt.W})
        item=E.SubElement(notes,fmt.Q('footnote'),{fmt.Q('id'):'1'});item.append(note)
        target=self.root/'target.docx'
        with zipfile.ZipFile(source) as zin,zipfile.ZipFile(target,'w') as zout:
            for name in zin.namelist():zout.writestr(name,fmt.serialize(root) if name=='word/document.xml' else zin.read(name))
            zout.writestr('word/footnotes.xml',fmt.serialize(notes))
        receipt={'footnote_changes':[{'kind':'footnote','reference_paragraph':1,'note_paragraph':2,'before':'1 '+body,'note_id':'1'}]}
        return source,target,receipt

    def test_registered_native_note_move_preserves_business_words(self):
        source,target,receipt=self.native_note()
        plan=content.register_footnote_changes(source,target,receipt,content.make_plan(source))
        self.assertEqual(content.audit(source,target,plan)['status'],'passed')

    def test_native_note_number_change_is_blocked(self):
        source,target,receipt=self.native_note()
        plan=content.register_footnote_changes(source,target,receipt,content.make_plan(source))
        altered=self.root/'altered.docx'
        package_replace(target,altered,'word/footnotes.xml',bytes_part(target,'word/footnotes.xml').replace(b'0.05',b'0.5'))
        self.assertEqual(content.audit(source,altered,plan)['status'],'failed')

    def test_note_receipt_does_not_authorize_new_prose(self):
        source,target,receipt=self.native_note()
        altered=self.root/'altered.docx'
        package_replace(target,altered,'word/footnotes.xml',bytes_part(target,'word/footnotes.xml').replace(b'0.05',b'0.5'))
        with self.assertRaises(content.ContentIntegrityError):
            content.register_footnote_changes(source,altered,receipt,content.make_plan(source))

    def test_unregistered_note_move_is_not_silently_allowed(self):
        source,target,_=self.native_note()
        self.assertEqual(content.audit(source,target)['status'],'failed')


class LiteralContentTests(TemporaryTest):
    def test_url_punctuation_is_not_normalized_away(self):
        source=self.document('source.docx','访问 https://example.test/资料（版本1）')
        target=self.document('target.docx','访问 https://example.test/资料(版本1)')
        self.assertEqual(content.audit(source,target)['status'],'failed')

    def test_inline_code_cannot_change_chinese_punctuation(self):
        source=self.document('source.docx','命令 `print("你好，世界")`')
        target=self.document('target.docx','命令 `print("你好,世界")`')
        self.assertEqual(content.audit(source,target)['status'],'failed')

    def test_math_business_number_is_protected(self):
        source=self.document('source.docx','技术公式')
        root=fmt.parse(bytes_part(source,'word/document.xml'));p=next(root.iter(fmt.Q('p')))
        math=E.SubElement(p,'{'+content.M+'}oMath');r=E.SubElement(math,'{'+content.M+'}r');t=E.SubElement(r,'{'+content.M+'}t');t.text='0.05'
        old=self.root/'math-old.docx';package_replace(source,old,'word/document.xml',fmt.serialize(root))
        t.text='0.5';new=self.root/'math-new.docx';package_replace(source,new,'word/document.xml',fmt.serialize(root))
        self.assertEqual(content.audit(old,new)['status'],'failed')


class MergeContentTests(TemporaryTest):
    def merged(self,second='项目A'):
        d=Document();t=d.add_table(rows=2,cols=2)
        t.cell(0,0).text='项目A';t.cell(1,0).text=second;t.cell(0,1).text='要求1';t.cell(1,1).text='要求2'
        source=self.root/'source.docx';d.save(source)
        root=fmt.parse(bytes_part(source,'word/document.xml'));table=next(root.iter(fmt.Q('tbl')))
        cells=[r.findall(fmt.Q('tc'))[0] for r in table.findall(fmt.Q('tr'))]
        for cell,value in zip(cells,('restart','continue')):
            cp=cell.find(fmt.Q('tcPr'));cp.append(property_node('vMerge',val=value))
        for text in cells[1].iter(fmt.Q('t')):text.text=''
        target=self.root/'target.docx';package_replace(source,target,'word/document.xml',fmt.serialize(root))
        return source,target

    def test_equal_group_label_can_be_vertically_merged_without_false_content_loss(self):
        source,target=self.merged()
        self.assertEqual(content.audit(source,target)['status'],'passed')

    def test_merging_different_labels_cannot_hide_content_loss(self):
        source,target=self.merged('项目B')
        self.assertEqual(content.audit(source,target)['status'],'failed')

    def test_orphan_merge_is_blocked(self):
        source,target=self.merged()
        root=fmt.parse(bytes_part(target,'word/document.xml'))
        first=root.find('.//w:vMerge',fmt.NS);first.set(fmt.Q('val'),'continue')
        bad=self.root/'bad.docx';package_replace(target,bad,'word/document.xml',fmt.serialize(root))
        with self.assertRaises(content.ContentIntegrityError):content.audit(source,bad)

if __name__ == '__main__':
    unittest.main()
