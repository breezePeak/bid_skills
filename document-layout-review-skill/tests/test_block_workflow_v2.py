"""Block v2 regression: real DOCX scopes plus explicitly mocked engine boundaries."""
import copy
import hashlib
import importlib.util
import io
import json
import subprocess
import sys
import types
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree as E
from PIL import Image

SCRIPTS=Path(__file__).resolve().parents[1]/'scripts'
sys.path.insert(0,str(SCRIPTS))
import block_progress as bp
import block_workflow as bw
import block_cache as bc
import block_evidence as be


def save_parts(path,parts):
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in parts.items():
            z.writestr(name,data)


def fixture(tmp_path,n=3):
    src=tmp_path/'original-input.docx'; d=Document()
    for i in range(n): d.add_paragraph(f'第{i+1}段：工期30天，金额50万元。')
    d.save(src); work=tmp_path/'work'
    bp.initialize(src,work,target_chars=1,max_paragraphs=1)
    bp.accept_global(work,bp.load(work)['current'],'已核对全局设置')
    return src,work


@pytest.mark.parametrize('kind',['table','figure'])
@pytest.mark.parametrize('before',[True,False])
@pytest.mark.parametrize('blanks',[0,1,2])
def test_caption_stays_with_object_across_blank_paragraphs(tmp_path,kind,before,blanks):
    d=Document(); caption='表1 配置' if kind=='table' else '图1 结构'
    image=tmp_path/'pixel.png'; Image.new('RGB',(24,24)).save(image)
    if before:
        d.add_paragraph(caption)
        for _ in range(blanks):d.add_paragraph('')
    if kind=='table':d.add_table(rows=4,cols=2)
    else:d.add_picture(str(image))
    if not before:
        for _ in range(blanks):d.add_paragraph('')
        d.add_paragraph(caption)
    d.add_paragraph('后续正文')
    src=tmp_path/'doc.docx';d.save(src)
    blocks=bp.make_blocks(src,target_chars=1,max_paragraphs=1)
    assert blocks[0]['end']==2+blanks
    assert blocks[1]['start']==2+blanks


def test_image_without_text_does_not_absorb_following_body(tmp_path):
    im=tmp_path/'x.png';Image.new('RGB',(20,20)).save(im)
    d=Document();d.add_picture(str(im));d.add_paragraph('不属于图片的正文')
    src=tmp_path/'source.docx';d.save(src)
    blocks=bp.make_blocks(src,target_chars=1,max_paragraphs=1)
    assert [(b['start'],b['end']) for b in blocks]==[(0,1),(1,2)]


def test_noop_checkpoints_do_not_duplicate_whole_docx(tmp_path):
    _,work=fixture(tmp_path,30)
    for block in list(bp.load(work)['blocks']):
        bp.checkpoint(work,block['id'],bp.load(work)['current'],'已程序核对，未修改')
    assert len(list((work/'checkpoints').glob('*.docx')))==1
    assert bp.next_block(work)['completed']==30


def test_all_header_footer_parts_are_in_inventory(tmp_path):
    d=Document();d.add_paragraph('正文');d.sections[0].header.paragraphs[0].text='页眉'
    d.sections[0].footer.paragraphs[0].text='页脚';src=tmp_path/'story.docx';d.save(src)
    blocks=bp.make_blocks(src)
    assert {'word/header1.xml','word/footer1.xml'}<={b['part'] for b in blocks}


def test_story_local_change_is_accepted_and_other_story_is_protected(tmp_path):
    d=Document();d.add_paragraph('正文');d.sections[0].header.paragraphs[0].text='页眉'
    d.sections[0].footer.paragraphs[0].text='页脚';src=tmp_path/'story.docx';d.save(src)
    blocks=bp.make_blocks(src);block=next(b for b in blocks if b['part']=='word/header1.xml')
    d.sections[0].header.paragraphs[0].runs[0].font.size=Pt(12);p=tmp_path/'header.docx';d.save(p)
    assert bp.local_scope(src,p,block)[0]==0
    d.sections[0].footer.paragraphs[0].runs[0].font.size=Pt(18);p2=tmp_path/'both.docx';d.save(p2)
    with pytest.raises(bp.BlockError,match='块外'):bp.local_scope(src,p2,block)


def test_global_setup_creates_story_coverage_without_restarting_body(tmp_path):
    src=tmp_path/'s.docx';d=Document();d.add_paragraph('正文');d.save(src)
    w=tmp_path/'w';bp.initialize(src,w)
    d.sections[0].header.paragraphs[0].text='';p=tmp_path/'setup.docx';d.save(p)
    bp.accept_global(w,p,'模板页眉设置')
    assert any(b['part']=='word/header1.xml' for b in bp.load(w)['blocks'])


def test_shared_late_repair_keeps_completed_content_and_invalidates_final(tmp_path):
    _,work=fixture(tmp_path)
    for b in list(bp.load(work)['blocks']):bp.checkpoint(work,b['id'],bp.load(work)['current'],'已检查')
    st=bp.load(work);st['final']={'candidate':'old'};bp.write_json(work/bp.STATE,st)
    d=Document(st['current']);d.sections[0].left_margin=Pt(65)
    p=tmp_path/'shared.docx';d.save(p)
    bp.accept_shared_repair(work,p,'页边距不符合模板，影响本节分页')
    st=bp.load(work)
    assert all(b['status']=='checked' for b in st['blocks'])
    assert st['final'] is None and st['previous_final']['candidate']=='old'


def test_shared_repair_cannot_change_business_text(tmp_path):
    _,work=fixture(tmp_path);d=Document(bp.load(work)['current'])
    d.paragraphs[0].text='删掉业务内容';p=tmp_path/'bad.docx';d.save(p)
    with pytest.raises(bp.BlockError):bp.accept_shared_repair(work,p,'不允许')


def test_os_lock_is_released_when_worker_exits_abnormally(tmp_path):
    _,work=fixture(tmp_path)
    code="import sys,os;sys.path.insert(0,sys.argv[1]);from block_progress import session_lock\nwith session_lock(sys.argv[2]):os._exit(7)"
    result=subprocess.run([sys.executable,'-c',code,str(SCRIPTS),str(work)],capture_output=True)
    assert result.returncode==7
    with bp.session_lock(work):pass


def test_initialization_recovers_matching_partial_copy(tmp_path):
    d=Document();d.add_paragraph('原文');src=tmp_path/'s.docx';d.save(src)
    w=tmp_path/'w';w.mkdir();(w/'original.docx').write_bytes(src.read_bytes())
    assert bp.initialize(src,w)['source_sha256']==bp.sha(src)


def test_initialization_never_overwrites_different_partial_source(tmp_path):
    d=Document();d.add_paragraph('原文');src=tmp_path/'s.docx';d.save(src)
    w=tmp_path/'w';w.mkdir();(w/'original.docx').write_bytes(b'different')
    with pytest.raises(bp.BlockError):bp.initialize(src,w)


@pytest.mark.parametrize('damage',['missing','duplicate','out_of_range','overlap'])
def test_corrupt_block_coverage_is_not_accepted(tmp_path,damage):
    _,work=fixture(tmp_path);st=bp.load(work)
    if damage=='missing':st['blocks'].pop()
    elif damage=='duplicate':st['blocks'][1]['id']=st['blocks'][0]['id']
    elif damage=='out_of_range':st['blocks'][1]['end']=500
    else:st['blocks'][1]['start']=0
    bp.write_json(work/bp.STATE,st)
    with pytest.raises(bp.BlockError):bp.load(work)


def test_template_and_profile_same_basename_do_not_overwrite_each_other(tmp_path,monkeypatch):
    a=tmp_path/'a';b=tmp_path/'b';a.mkdir();b.mkdir()
    template=a/'same';profile=b/'same';template.write_bytes(b'template');profile.write_bytes(b'profile')
    monkeypatch.setitem(sys.modules,'review_pipeline',types.SimpleNamespace(resolve_template=lambda a,r:(template,profile)))
    monkeypatch.setitem(sys.modules,'text_rules',types.SimpleNamespace(load_rules=lambda _: {},write_rules=lambda p,r:bp.write_json(p,r)))
    reports=tmp_path/'work'/'reports';reports.mkdir(parents=True)
    args=types.SimpleNamespace(text_rules=None,renderer=None,field_engine=None,uno_python=None,vision_worker_config=None)
    result=bw.settings_from_args(args,reports)
    assert Path(result['template']).read_bytes()==b'template'
    assert Path(result['template_style_json']).read_bytes()==b'profile'


@pytest.fixture
def cached_script(tmp_path):
    scripts=tmp_path/'scripts';scripts.mkdir();(scripts/'audit.py').write_text('# checker v1')
    source=tmp_path/'candidate.docx';source.write_bytes(b'candidate')
    out=tmp_path/'report.json';calls=[]
    def execute(name,*args):
        calls.append(name);bp.write_json(out,{'status':'passed'})
        return {'returncode':0,'ok':True,'stdout':'{}','stderr':''}
    args=[source,'--json-out',out]
    return scripts,source,out,calls,execute,args


def test_successful_deterministic_audit_is_reused(cached_script):
    s,src,out,calls,run,args=cached_script
    bc.cached_run('audit.py',args,s,run);second=bc.cached_run('audit.py',args,s,run)
    assert len(calls)==1 and second['reused_step']


@pytest.mark.parametrize('change',['input','report','checker','arguments'])
def test_audit_cache_invalidates_changed_dependency(cached_script,change):
    s,src,out,calls,run,args=cached_script
    bc.cached_run('audit.py',args,s,run)
    if change=='input':src.write_bytes(b'new candidate')
    elif change=='report':out.write_text('{"status":"forged"}')
    elif change=='checker':(s/'audit.py').write_text('# checker v2')
    else:args.extend(['--rule','strict'])
    bc.cached_run('audit.py',args,s,run)
    assert len(calls)==2


def test_failed_audit_is_not_reused_as_pass(cached_script):
    s,src,out,calls,_,args=cached_script
    def fail(name,*args):calls.append(name);return {'returncode':2,'ok':False,'stdout':'{}','stderr':''}
    bc.cached_run('audit.py',args,s,fail);bc.cached_run('audit.py',args,s,fail)
    assert len(calls)==2


def test_execution_only_reconfiguration_preserves_completed_blocks(tmp_path,capsys):
    _,work=fixture(tmp_path);st=bp.load(work)
    st['objects_initialized']=True;st['settings'].update(renderer='auto',field_engine='auto',uno_python=None)
    bp.write_json(work/bp.STATE,st)
    bp.checkpoint(work,st['blocks'][0]['id'],st['current'],'已核对')
    assert bw.main(['--work-dir',str(work),'--action','configure','--renderer','libreoffice'])==0
    result=json.loads(capsys.readouterr().out)
    assert result['completed']==1 and bp.load(work)['settings']['renderer']=='libreoffice'


def test_effective_rules_cannot_be_silently_switched_on_resume(tmp_path,capsys):
    _,work=fixture(tmp_path);st=bp.load(work);st['objects_initialized']=True
    st['settings'].update(renderer='word',field_engine='word',uno_python=None);bp.write_json(work/bp.STATE,st)
    assert bw.main(['--work-dir',str(work),'--renderer','wps'])==2
    assert json.loads(capsys.readouterr().out)['status']=='blocked'


def test_rule_change_invalidates_identical_page_reuse(tmp_path):
    page=tmp_path/'p.png';page.write_bytes(b'page')
    row={'name':'p.png','path':str(page),'sha256':bp.sha(page)}
    m=tmp_path/'m.json';v=tmp_path/'v.json'
    bp.write_json(m,{'candidate_sha256':'old','template_sha256':'old-template','render':{'pages':[row]}})
    bp.write_json(v,{'candidate_sha256':'old','pages':[{'name':'p.png','sha256':row['sha256'],'status':'pass','observations':'已查看'}]})
    pending={'pages':[{'name':'p.png','sha256':row['sha256'],'status':'pending'}],'figures':{'objects':[]}}
    got=bw.reuse_unchanged_reviews({'manifest':str(m),'review':str(v),'candidate_sha256':'old'},pending,{'template_sha256':'new-template','render':{'pages':[row]}})
    assert got['pages'][0]['status']=='pending'


@pytest.fixture
def composite(tmp_path,monkeypatch):
    source=tmp_path/'source.docx';source.write_bytes(b'original')
    current=tmp_path/'current.docx';current.write_bytes(b'current')
    page=tmp_path/'page.png';page.write_bytes(b'current-page')
    pref={'name':'page.png','path':str(page),'sha256':bp.sha(page)}
    proof={'verdict':'pass','binding':{'source_sha256':bp.sha(source)},'bundle':'bundle',
           'checks':{'text_inside_bounds':'pass'},'image_type':'diagram','observation':'真实调用样例输出'}
    class VE(ValueError):pass
    monkeypatch.setitem(sys.modules,'visual_evidence',types.SimpleNamespace(
        VisualError=VE,checked_file=lambda r:Path(r['path']),validate_inspection=lambda *a,**k:proof,
        validate_bundle=lambda _: {'assets':[]},file_sha=bp.sha,read=lambda _: {'inspections':[]}))
    pixel_calls=[]
    monkeypatch.setitem(sys.modules,'figure_inspection',types.SimpleNamespace(verify_row=None,_assert_pixels=lambda *a:pixel_calls.append(1)))
    obj={'id':'F0001','hash':'same-image','paths':['word/media/image.png']}
    old={'object_sha256':'same-image','inspection':{'path':'actual-receipt','sha256':'receipt'}}
    row={'pages':[{'name':'page.png','sha256':bp.sha(page)}]}
    pages=[pref];reviews=[{'name':'page.png','sha256':bp.sha(page),'status':'pass','observations':'当前页图与图题完整'}]
    return source,current,obj,row,object(),old,pages,reviews,proof,pixel_calls


def test_normal_image_reuses_intrinsic_proof_and_current_page_without_new_call(composite):
    *args,proof,calls=composite
    assert be.reuse_source_content(*args)
    row=args[3]
    assert row['inspection'] is None
    assert row['content_review']['method']=='unchanged-content-current-pages'
    assert calls==[1]
    assert be.reuse_source_content(*args,validate=True)


@pytest.mark.parametrize('change',['image','initial-fail','page-pending','page-hash','native','no-media','no-page'])
def test_composite_reuse_never_skips_changed_or_unverified_object(composite,change):
    *args,proof,calls=composite
    if change=='image':args[2]['hash']='changed'
    elif change=='initial-fail':proof['verdict']='fail'
    elif change=='page-pending':args[7][0]['status']='pending'
    elif change=='page-hash':args[7][0]['sha256']='old'
    elif change=='native':args[5]['rendered_view']={'something':True}
    elif change=='no-media':args[2]['paths']=[]
    else:args[3]['pages']=[]
    assert not be.reuse_source_content(*args)


def test_old_composite_cannot_be_relabelled_as_current_page(composite):
    *args,proof,calls=composite
    assert be.reuse_source_content(*args)
    args[3]['content_review']['candidate_sha256']='forged'
    with pytest.raises(ValueError):be.reuse_source_content(*args,validate=True)
