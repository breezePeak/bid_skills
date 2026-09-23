"""Numbering regression tests using valid DOCX packages, not renamed XML files."""
from __future__ import annotations
import copy
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree as E

SCRIPTS=Path(__file__).resolve().parents[1]/'scripts'
sys.path.insert(0,str(SCRIPTS))
from numbering_core import (audit,repair,Package,NumberingError,NS,Q,R,REL,CT,item,run,field,
                            text,dump,xml,val,read_fields,set_number,heading_values,
                            heading_snapshot,preserve_heading_numbering,replace_span)


def mutate(path,fn,part='word/document.xml'):
    with zipfile.ZipFile(path) as z: files={n:z.read(n) for n in z.namelist()}
    root=xml(files[part]);fn(root);files[part]=dump(root)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in files.items():z.writestr(n,data)


def read(path,part='word/document.xml'):
    with zipfile.ZipFile(path) as z:return xml(z.read(part))


def make(path,paragraphs):
    doc=Document()
    for name,kind in [('Footnote Text',WD_STYLE_TYPE.PARAGRAPH),('Footnote Reference',WD_STYLE_TYPE.CHARACTER)]:
        if name not in doc.styles:doc.styles.add_style(name,kind)
    for value,style in paragraphs:doc.add_paragraph(value,style=style)
    doc.save(path)


def add_native(path,*,custom=False,missing_inside=False,year=False):
    pkg=Package(path);p=pkg.paragraphs[0]
    r=run();ref=item('footnoteReference',id=1)
    if custom:ref.set(Q('customMarkFollows'),'1')
    r.append(ref)
    if custom:
        t=item('t');t.text='*';r.append(t)
    p.append(r)
    notes=pkg.get('word/footnotes.xml','footnotes')
    for ident,kind,tag in [(-1,'separator','separator'),(0,'continuationSeparator','continuationSeparator')]:
        note=item('footnote',id=ident,type=kind);para=item('p');rr=run();rr.append(item(tag));para.append(rr);note.append(para);notes.append(note)
    note=item('footnote',id=1);para=item('p')
    if custom:para.append(run('* '))
    elif not missing_inside:
        rr=run();rr.append(item('footnoteRef'));para.append(rr)
    para.append(run('2026年资料' if year else 'Note content.'))
    note.append(para);notes.append(note)
    pkg.touch('word/footnotes.xml');pkg.touch('word/document.xml');pkg.ensure_part('word/footnotes.xml','footnotes')
    temp=path.with_suffix('.tmp.docx');pkg.save(temp);temp.replace(path)


def superscript(path,p_index,marker='1'):
    def change(root):
        p=root.findall('.//w:body/w:p',NS)[p_index]
        props=item('rPr');props.append(item('vertAlign',val='superscript'))
        p.append(run(marker,props))
    mutate(path,change)


class NumberingTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.dir=Path(self.tmp.name);self.src=self.dir/'input.docx';self.out=self.dir/'output.docx'

    def convert(self,paragraphs):
        make(self.src,paragraphs);report=repair(self.src,self.out)
        self.assertEqual(report['status'],'passed',report)
        self.assertEqual(audit(self.out)['status'],'passed')
        Document(self.out)
        return report

    def test_01_manual_hierarchy(self):
        r=self.convert([('1 项目','Heading 1'),('1.1 范围','Heading 2'),('1.2 目标','Heading 2'),('2 实施','Heading 1'),('2.1 步骤','Heading 2')])
        self.assertEqual(r['change_count'],5)
        pkg=Package(self.out);records=heading_snapshot(pkg)
        self.assertEqual([x['label'] for x in heading_values(pkg,records).values()],['1','1.1','1.2','2','2.1'])
        self.assertEqual([text(p) for p in pkg.paragraphs],['项目','范围','目标','实施','步骤'])
        self.assertEqual(len({pkg.num(p)['num_id'] for p in pkg.paragraphs}),1)

    def test_02_chinese_heading_format(self):
        self.convert([('第一章 范围','Heading 1'),('一、概述','Heading 2'),('二、目标','Heading 2'),('第二章 实施','Heading 1'),('一、方法','Heading 2')])
        pkg=Package(self.out)
        self.assertEqual([v['label'] for v in heading_values(pkg,heading_snapshot(pkg)).values()],['第一章','一、','二、','第二章','一、'])

    def test_03_no_manual_number_adds_native(self):
        self.convert([('概述','Heading 1'),('细节','Heading 2')])
        self.assertEqual([text(p) for p in Package(self.out).paragraphs],['概述','细节'])

    def test_04_gap_blocks_without_modifying_source(self):
        make(self.src,[('1 概述','Heading 1'),('3 方法','Heading 1')]);before=self.src.read_bytes()
        with self.assertRaises(NumberingError):repair(self.src,self.out)
        self.assertEqual(before,self.src.read_bytes());self.assertFalse(self.out.exists())

    def test_05_style_inherited_auto_detected(self):
        self.convert([('1 概述','Heading 1'),('2 方法','Heading 1')])
        pkg=Package(self.out);nid=pkg.num(pkg.paragraphs[0])['num_id']
        def clear(root):
            for p in root.findall('.//w:body/w:p',NS):
                pp=p.find('w:pPr',NS);pp.remove(pp.find('w:numPr',NS))
        mutate(self.out,clear)
        def linked(root):
            s=root.find('w:style[@w:styleId="Heading1"]',NS)
            pp=s.find('w:pPr',NS);n=item('numPr');n.append(item('numId',val=nid));pp.append(n)
        mutate(self.out,linked,'word/styles.xml')
        self.assertEqual(audit(self.out)['status'],'passed')
        output2=self.dir/'again.docx';report=repair(self.out,output2)
        self.assertEqual(report['change_count'],0);self.assertEqual(self.out.read_bytes(),output2.read_bytes())

    def test_06_num_id_zero_not_automatic(self):
        make(self.src,[('1 概述','Heading 1')]);mutate(self.src,lambda root:set_number(root.find('w:body/w:p',NS),'0',0))
        self.assertEqual(audit(self.src)['status'],'failed')
        self.assertEqual(repair(self.src,self.out)['status'],'passed')

    def test_07_broken_num_reference(self):
        make(self.src,[('1 概述','Heading 1')]);mutate(self.src,lambda root:set_number(root.find('w:body/w:p',NS),'9876',0))
        self.assertEqual(audit(self.src)['status'],'failed');self.assertEqual(repair(self.src,self.out)['status'],'passed')

    def test_08_body_lists_cover_toc_unchanged(self):
        self.convert([('投标文件','Title'),('Contents','TOC Heading'),('1. ordinary list','List Paragraph'),('图1显示效果。','Normal')])
        self.assertEqual(self.src.read_bytes(),self.out.read_bytes())

    def test_09_captions_split_runs(self):
        make(self.src,[('图1 系统架构','Caption'),('图2 流程','Caption'),('表1 参数','Caption')])
        def split(root):
            p=root.find('w:body/w:p',NS)
            for r in p.findall('w:r',NS):p.remove(r)
            p.append(run('图'));p.append(run('1'));pr=item('rPr');pr.append(item('b'));p.append(run(' 系统架构',pr))
        mutate(self.src,split)
        self.assertEqual(repair(self.src,self.out)['status'],'passed')
        p=read(self.out).find('w:body/w:p',NS)
        self.assertEqual(text(p),'图1 系统架构');self.assertIsNotNone(p.find('w:r/w:rPr/w:b',NS))
        self.assertIn('SEQ Figure',read_fields(p)[0]['instr'])

    def test_10_chapter_captions_native_prefix(self):
        self.convert([('第1章 架构','Heading 1'),('图1-1 架构图','Caption'),('图1-2 详细图','Caption'),('第2章 实施','Heading 1'),('图2-1 流程图','Caption')])
        pkg=Package(self.out)
        for p in pkg.paragraphs:
            if pkg.is_caption(p):
                fields=read_fields(p);self.assertEqual(len(fields),2)
                self.assertIn('STYLEREF',fields[0]['instr']);self.assertIn('\\s 1',fields[1]['instr'])

    def test_11_figure_table_sequences_independent(self):
        self.convert([('图1 第一图','Caption'),('表1 第一表','Caption'),('图2 第二图','Caption'),('表2 第二表','Caption')])
        self.assertEqual([read_fields(p)[0]['instr'].strip().split()[1] for p in Package(self.out).paragraphs],['Figure','Table','Figure','Table'])

    def test_12_native_captions_idempotent(self):
        self.convert([('图1 题注','Caption')]);second=self.dir/'second.docx'
        report=repair(self.out,second);self.assertEqual(report['change_count'],0);self.assertEqual(self.out.read_bytes(),second.read_bytes())

    def test_13_complex_seq_field_recognized(self):
        make(self.src,[('图','Caption')])
        def edit(root):
            p=root.find('w:body/w:p',NS)
            for tag,value in [('fldChar','begin'),('instrText',' SE'),('instrText','Q Figure '),('fldChar','separate'),('t','1'),('fldChar','end'),('t',' 标题')]:
                r=run();n=item(tag)
                if tag=='fldChar':n.set(Q('fldCharType'),value)
                else:n.text=value
                r.append(n);p.append(r)
        mutate(self.src,edit);self.assertEqual(audit(self.src)['status'],'passed')

    def test_14_seq_elsewhere_does_not_mask_static_label(self):
        make(self.src,[('图1 标题 ','Caption')]);mutate(self.src,lambda root:root.find('w:body/w:p',NS).append(field('SEQ Figure','1')))
        self.assertEqual(audit(self.src)['status'],'failed')

    def test_15_locked_field_unlocked(self):
        self.convert([('图1 标题','Caption')]);self.src.write_bytes(self.out.read_bytes());self.out.unlink()
        mutate(self.src,lambda root:root.find('.//w:fldSimple',NS).set(Q('fldLock'),'true'))
        self.assertEqual(audit(self.src)['status'],'failed');self.assertEqual(repair(self.src,self.out)['status'],'passed')

    def test_16_caption_gap_blocks(self):
        make(self.src,[('图1 第一图','Caption'),('图3 第三图','Caption')])
        with self.assertRaises(NumberingError):repair(self.src,self.out)
        self.assertFalse(self.out.exists())

    def test_17_caption_bookmark_survives(self):
        make(self.src,[('图','Caption')])
        def edit(root):
            p=root.find('w:body/w:p',NS);p.append(item('bookmarkStart',id=6,name='FigNo'));p.append(run('1'));p.append(item('bookmarkEnd',id=6));p.append(run(' 标题'))
        mutate(self.src,edit);self.assertEqual(repair(self.src,self.out)['status'],'passed')
        p=read(self.out).find('w:body/w:p',NS)
        self.assertLess(p.index(p.find('w:bookmarkStart',NS)),p.index(p.find('w:fldSimple',NS)))
        self.assertLess(p.index(p.find('w:fldSimple',NS)),p.index(p.find('w:bookmarkEnd',NS)))

    def test_18_manual_footnote_converts(self):
        make(self.src,[('正文内容','Normal'),('1 脚注说明','Footnote Text')]);superscript(self.src,0)
        self.assertEqual(audit(self.src)['status'],'failed')
        r=repair(self.src,self.out);self.assertEqual(r['status'],'passed',r)
        root=read(self.out);self.assertEqual(len(root.findall('.//w:footnoteReference',NS)),1)
        notes=read(self.out,'word/footnotes.xml');self.assertIn('脚注说明',text(notes))
        self.assertNotIn('脚注说明',text(root));Document(self.out)
        self.assertEqual(audit(self.out)['status'],'passed')

    def test_19_two_notes_same_paragraph(self):
        make(self.src,[('正文','Normal'),('1 说明甲','Footnote Text'),('2 说明乙','Footnote Text')]);superscript(self.src,0,'1');superscript(self.src,0,'2')
        r=repair(self.src,self.out);self.assertEqual(r['status'],'passed',r)
        self.assertEqual(len(read(self.out).findall('.//w:footnoteReference',NS)),2)

    def test_20_orphan_manual_note_blocks(self):
        make(self.src,[('1 脚注说明','Footnote Text')])
        with self.assertRaises(NumberingError):repair(self.src,self.out)
        self.assertFalse(self.out.exists())

    def test_21_exponent_not_footnote(self):
        make(self.src,[('面积 m','Normal')]);superscript(self.src,0,'2')
        self.assertEqual(repair(self.src,self.out)['change_count'],0);self.assertEqual(self.src.read_bytes(),self.out.read_bytes())

    def test_22_native_footnote_unchanged(self):
        make(self.src,[('正文','Normal')]);add_native(self.src)
        self.assertEqual(audit(self.src)['status'],'passed');self.assertEqual(repair(self.src,self.out)['change_count'],0)
        self.assertEqual(self.src.read_bytes(),self.out.read_bytes())

    def test_23_missing_note_content_blocks(self):
        make(self.src,[('正文','Normal')]);mutate(self.src,lambda root:root.find('w:body/w:p/w:r',NS).append(item('footnoteReference',id=99)))
        self.assertEqual(audit(self.src)['status'],'failed')
        with self.assertRaises(NumberingError):repair(self.src,self.out)

    def test_24_missing_inside_note_marker_no_year_loss(self):
        make(self.src,[('正文','Normal')]);add_native(self.src,missing_inside=True,year=True)
        self.assertEqual(repair(self.src,self.out)['status'],'passed');notes=read(self.out,'word/footnotes.xml')
        self.assertIn('2026年资料',text(notes));self.assertEqual(len(notes.findall('.//w:footnoteRef',NS)),1)

    def test_25_custom_footnote_mark_conversion(self):
        make(self.src,[('正文','Normal')]);add_native(self.src,custom=True)
        self.assertEqual(audit(self.src)['status'],'failed');r=repair(self.src,self.out);self.assertEqual(r['status'],'passed',r)
        self.assertIsNone(read(self.out).find('.//w:footnoteReference',NS).get(Q('customMarkFollows')))
        self.assertNotIn('*',text(read(self.out)));self.assertNotIn('*',text(read(self.out,'word/footnotes.xml')))

    def test_26_duplicate_note_id_rejected(self):
        make(self.src,[('正文','Normal')]);add_native(self.src)
        def dup(root):root.append(copy.deepcopy(root.find('w:footnote[@w:id="1"]',NS)))
        mutate(self.src,dup,'word/footnotes.xml');self.assertEqual(audit(self.src)['status'],'failed')
        with self.assertRaises(NumberingError):repair(self.src,self.out)

    def test_27_explicit_manual_footnote_map(self):
        make(self.src,[('引用[1]。','Normal'),('[1] 具体说明','Normal')]);pkg=Package(self.src)
        plan={'source_sha256':'sha256:'+hashlib.sha256(self.src.read_bytes()).hexdigest(),
              'footnotes':[{'reference_paragraph':1,'reference_text':'引用[1]。','start':2,'end':5,'note_paragraph':2,'note_text':'[1] 具体说明'}]}
        path=self.dir/'plan.json';path.write_text(json.dumps(plan))
        self.assertEqual(repair(self.src,self.out,path)['status'],'passed');self.assertEqual(text(read(self.out)),'引用。')

    def test_28_stale_map_not_applied(self):
        make(self.src,[('1 标题','Normal')]);plan=self.dir/'plan.json';plan.write_text(json.dumps({'source_sha256':'bad','headings':[{'paragraph':1,'text':'1 标题','level':1}]}))
        with self.assertRaises(NumberingError):repair(self.src,self.out,plan)
        self.assertFalse(self.out.exists())

    def test_29_map_classifies_unstyled_heading(self):
        make(self.src,[('1 标题','Normal')]);plan=self.dir/'plan.json';plan.write_text(json.dumps({'source_sha256':'sha256:'+hashlib.sha256(self.src.read_bytes()).hexdigest(),'headings':[{'paragraph':1,'text':'1 标题','level':1}]}))
        self.assertEqual(repair(self.src,self.out,plan)['status'],'passed');self.assertIsNotNone(Package(self.out).num(Package(self.out).paragraphs[0]))

    def test_30_source_never_overwritten(self):
        make(self.src,[('1 标题','Heading 1')])
        with self.assertRaises(NumberingError):repair(self.src,self.src)

    def test_31_heading_source_bookmark_ref_blocks(self):
        make(self.src,[('1 标题','Heading 1'),('参见','Normal')])
        def change(root):
            ps=root.findall('w:body/w:p',NS);ps[0].insert(1,item('bookmarkStart',id=1,name='H1'));ps[0].append(item('bookmarkEnd',id=1));ps[1].append(field('REF H1','1 标题'))
        mutate(self.src,change)
        with self.assertRaises(NumberingError):repair(self.src,self.out)

    def test_32_assets_table_section_unchanged(self):
        doc=Document();doc.add_heading('1 标题',level=1);doc.add_table(2,2);doc.sections[0].header.paragraphs[0].text='原页眉';doc.save(self.src)
        with zipfile.ZipFile(self.src,'a') as z:z.writestr('word/media/sentinel.bin',b'never change')
        self.assertEqual(repair(self.src,self.out)['status'],'passed')
        a,b=read(self.src),read(self.out)
        self.assertEqual(E.tostring(a.find('.//w:tbl',NS)),E.tostring(b.find('.//w:tbl',NS)))
        self.assertEqual(E.tostring(a.find('.//w:sectPr',NS)),E.tostring(b.find('.//w:sectPr',NS)))
        with zipfile.ZipFile(self.src) as x,zipfile.ZipFile(self.out) as y:
            for part in ('word/media/sentinel.bin','word/header1.xml','word/styles.xml','word/theme/theme1.xml'):
                self.assertEqual(x.read(part),y.read(part))

    def test_33_cli_reports_failure_without_output(self):
        make(self.src,[('1 标题','Heading 1'),('3 标题','Heading 1')]);report=self.dir/'report.json'
        cp=subprocess.run([sys.executable,str(SCRIPTS/'numbering_repair.py'),str(self.src),'--out',str(self.out),'--json-out',str(report)],capture_output=True,text=True)
        self.assertEqual(cp.returncode,2);self.assertEqual(json.loads(report.read_text())['status'],'blocked');self.assertFalse(self.out.exists())

    def test_34_cli_audit_and_repair(self):
        make(self.src,[('1 标题','Heading 1'),('图1 题注','Caption')])
        cp=subprocess.run([sys.executable,str(SCRIPTS/'numbering_audit.py'),str(self.src)],capture_output=True,text=True);self.assertEqual(cp.returncode,2)
        cp=subprocess.run([sys.executable,str(SCRIPTS/'numbering_repair.py'),str(self.src),'--out',str(self.out)],capture_output=True,text=True);self.assertEqual(cp.returncode,0,cp.stdout+cp.stderr)
        cp=subprocess.run([sys.executable,str(SCRIPTS/'numbering_audit.py'),str(self.out)],capture_output=True,text=True);self.assertEqual(cp.returncode,0,cp.stdout+cp.stderr)

    def test_35_partial_auto_chapter_prefix_fixed(self):
        self.convert([('第1章 主题','Heading 1')]);self.src.write_bytes(self.out.read_bytes());self.out.unlink()
        def change(root):
            p=item('p');pp=item('pPr');pp.append(item('pStyle',val='Caption'));p.append(pp);p.append(run('图1-'));p.append(field('SEQ Figure','1'));p.append(run(' 标题'));root.find('w:body',NS).insert(1,p)
        mutate(self.src,change);self.assertEqual(audit(self.src)['status'],'failed')
        self.assertEqual(repair(self.src,self.out)['status'],'passed')
        self.assertEqual(len(read_fields(Package(self.out).paragraphs[1])),2)

    def test_36_duplicate_resets_are_not_a_valid_sequence(self):
        make(self.src,[('图','Caption'),('图','Caption')])
        def change(root):
            for p in root.findall('w:body/w:p',NS):p.append(field('SEQ Figure \\r 1','1'));p.append(run(' 标题'))
        mutate(self.src,change);self.assertEqual(audit(self.src)['status'],'failed')
        self.assertEqual(repair(self.src,self.out)['status'],'blocked')

    def test_37_preserve_inherited_number_before_template(self):
        self.test_05_style_inherited_auto_detected()
        with zipfile.ZipFile(self.out) as z:
            root=xml(z.read('word/document.xml'));count=preserve_heading_numbering(root,z.read('word/styles.xml'),z.read('word/numbering.xml'))
        self.assertEqual(count,2);self.assertEqual(len(root.findall('.//w:numPr',NS)),2)

    def test_38_chapter_reset_is_hidden_and_idempotent(self):
        self.convert([('第1章 范围','Heading 1'),('图1-1 架构','Caption'),('第2章 方法','Heading 1'),('图2-1 流程','Caption')])
        pkg=Package(self.out)
        for _,p,level,num in heading_snapshot(pkg):
            self.assertNotRegex(text(p),r'\d')
            fields=read_fields(p);self.assertEqual(len(fields),1)
            self.assertIn(r'\r 0 \h',fields[0]['instr'])
            self.assertIsNotNone(fields[0]['node'].find('w:r/w:rPr/w:vanish',NS))
        second=self.dir/'second.docx';result=repair(self.out,second)
        self.assertEqual(result['change_count'],0);self.assertEqual(second.read_bytes(),self.out.read_bytes())

    def test_39_chapter_ref_without_number_switch_rejected(self):
        self.convert([('1 Chapter','Heading 1'),('Figure 1-1 A','Caption')])
        def change(root):
            for f in root.findall('.//w:fldSimple',NS):
                if val(f,'instr','').strip().startswith('STYLEREF'):
                    f.set(Q('instr'),' STYLEREF "heading 1" ')
        mutate(self.out,change)
        self.assertEqual(audit(self.out)['status'],'failed')

    def test_40_template_application_preserves_inherited_numbering(self):
        self.test_05_style_inherited_auto_detected()
        template=self.dir/'template.docx';Document().save(template)
        from template_style_profile import extract_profile
        profile=self.dir/'profile.json';profile.write_text(json.dumps(extract_profile(template)))
        final=self.dir/'enforced.docx'
        cp=subprocess.run([sys.executable,str(SCRIPTS/'template_style_enforce.py'),str(template),str(profile),str(self.out),'--out',str(final)],capture_output=True,text=True)
        self.assertEqual(cp.returncode,0,cp.stderr+cp.stdout)
        self.assertEqual(audit(final)['status'],'passed')
        self.assertEqual(len(read(final).findall('.//w:numPr',NS)),2)

    def test_41_initial_inspection_runs_numbering_audit(self):
        import inspect_document
        from template_style_profile import extract_profile
        from text_rules import load_rules,rules_digest
        make(self.src,[('1 Chapter','Heading 1'),('Figure 1 A','Caption')])
        profile=self.dir/'profile.json';profile.write_text(json.dumps(extract_profile(self.src)))
        work=self.dir/'inspect';calls=[]
        def runner(name,*args,allow=(0,2,3,4)):
            args=list(map(str,args));calls.append(name)
            if name=='numbering_audit.py':
                cp=subprocess.run([sys.executable,str(SCRIPTS/name),*args],capture_output=True,text=True)
                return {'ok':cp.returncode in allow,'returncode':cp.returncode,'stdout':cp.stdout,'stderr':cp.stderr}
            if '--json-out' in args:
                data={'status':'passed','issues':[],'text_rules_sha256':rules_digest(load_rules())}
                Path(args[args.index('--json-out')+1]).write_text(json.dumps(data))
            return {'ok':True,'returncode':0,'stdout':'{}','stderr':''}
        from unittest.mock import patch
        import contextlib,io
        argv=['inspect',str(self.src),'--template',str(self.src),'--template-style-json',str(profile),'--work-dir',str(work)]
        with patch.object(sys,'argv',argv),patch.object(inspect_document,'run_script',runner),contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(inspect_document.main(),0)
        result=json.loads((work/'inspection.json').read_text())
        check=next(c for c in result['checks'] if '自动编号' in c['name'])
        self.assertEqual(check['result']['status'],'failed');self.assertGreaterEqual(result['hard_error_count'],2)

    def test_42_final_gate_rejects_numbering_broken_by_later_stage(self):
        import review_pipeline
        from template_style_profile import extract_profile
        from text_rules import load_rules,rules_digest
        from unittest.mock import patch
        import contextlib,io,shutil
        make(self.src,[('1 Chapter','Heading 1'),('Figure 1 A','Caption')])
        profile=self.dir/'profile.json';profile.write_text(json.dumps(extract_profile(self.src)))
        work=self.dir/'pipeline';calls=[]
        def runner(name,*args,allow=(0,)):
            args=list(map(str,args));calls.append(name)
            if name in {'numbering_repair.py','numbering_audit.py'}:
                cp=subprocess.run([sys.executable,str(SCRIPTS/name),*args],capture_output=True,text=True)
                return {'ok':cp.returncode in allow,'returncode':cp.returncode,'stdout':cp.stdout,'stderr':cp.stderr}
            if name=='render_docx.py':
                page=Path(args[args.index('--out-dir')+1])/'page.png';page.write_bytes(b'mock render')
                return {'ok':True,'returncode':0,'stdout':json.dumps({'pages':[str(page)]}),'stderr':''}
            if '--out' in args or '--output' in args:
                oi=args.index('--out') if '--out' in args else args.index('--output')
                src=args[2] if name=='template_style_enforce.py' else args[1] if name=='template_usage_repair.py' else args[0]
                out=Path(args[oi+1]);shutil.copy2(src,out)
                if name=='docx_layout_policy.py':
                    def break_it(root):
                        p=root.find('w:body/w:p',NS);pp=p.find('w:pPr',NS);pp.remove(pp.find('w:numPr',NS))
                    mutate(out,break_it)
            if '--json-out' in args:
                Path(args[args.index('--json-out')+1]).write_text(json.dumps({'status':'passed','issues':[],'issue_count':0,'text_rules_sha256':rules_digest(load_rules())}))
            return {'ok':True,'returncode':0,'stdout':'{}','stderr':''}
        argv=['pipeline',str(self.src),'--template',str(self.src),'--template-style-json',str(profile),'--work-dir',str(work)]
        with patch.object(sys,'argv',argv),patch.object(review_pipeline,'run_script',runner),contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(review_pipeline.main(),2)
        result=json.loads((work/'review-manifest.json').read_text())
        gate=next(g for g in result['gates'] if '自动编号' in g['name'])
        self.assertFalse(gate['passed']);self.assertEqual(result['status'],'failed')
        self.assertLess(calls.index('numbering_repair.py'),calls.index('template_style_enforce.py'))

    @unittest.skipUnless(__import__('shutil').which('soffice') and __import__('shutil').which('pdftotext'),'LibreOffice/pdftotext not installed')
    def test_43_real_field_refresh_after_inserting_caption(self):
        self.convert([('1 Chapter One','Heading 1'),('Figure 1-1 First','Caption'),('Figure 1-2 Second','Caption'),('2 Chapter Two','Heading 1'),('Figure 2-1 Third','Caption')])
        # Insert a copy with stale cached number: an actual field refresh must
        # advance the old captions, not merely keep their cached strings.
        def change(root):
            body=root.find('w:body',NS);original=body.findall('w:p',NS)[1]
            inserted=copy.deepcopy(original)
            for t in inserted.findall('.//w:t',NS):
                if t.text and 'First' in t.text:t.text=t.text.replace('First','Inserted')
            body.insert(body.index(original),inserted)
        mutate(self.out,change)
        cp=subprocess.run(['soffice','-env:UserInstallation='+ (self.dir/'lo-profile').as_uri(),'--headless','--convert-to','pdf','--outdir',str(self.dir),str(self.out)],capture_output=True,text=True,timeout=40)
        self.assertEqual(cp.returncode,0,cp.stdout+cp.stderr)
        pdf=self.out.with_suffix('.pdf');self.assertTrue(pdf.is_file())
        cp=subprocess.run(['pdftotext','-layout',str(pdf),'-'],capture_output=True,text=True,check=True)
        content=' '.join(cp.stdout.split())
        for expected in ['Figure 1-1 Inserted','Figure 1-2 First','Figure 1-3 Second','Figure 2-1 Third']:
            self.assertIn(expected,content)


    def test_44_split_caption_keeps_label_space(self):
        self.convert([('Figure 1 Label','Caption')])
        p=Package(self.out).paragraphs[0]
        first=p.find('w:r/w:t',NS)
        self.assertEqual(first.text,'Figure ')
        self.assertEqual(first.get('{http://www.w3.org/XML/1998/namespace}space'),'preserve')

if __name__=='__main__':unittest.main(verbosity=2)
