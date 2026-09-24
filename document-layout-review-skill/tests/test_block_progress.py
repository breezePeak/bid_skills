"""Real DOCX fixtures for scope, grouping and resume; no Office/vision claims."""
import copy
import importlib.util
import json
import sys
import types
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from lxml import etree as E
from PIL import Image

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import block_progress as bp
import block_workflow as bw
import block_evidence as be


def save_parts(path, parts):
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for k, v in parts.items():
            z.writestr(k, v)


def xml(root):
    return E.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)


@pytest.fixture
def case(tmp_path):
    src = tmp_path / 'source.docx'
    d = Document()
    for s in ['第一段的原有业务内容30天。', '第二段的原有业务内容50万元。', '第三段的原有业务内容。']:
        d.add_paragraph(s)
    d.save(src)
    work = tmp_path / 'work'
    bp.initialize(src, work, target_chars=1, max_paragraphs=1)
    bp.accept_global(work, bp.load(work)['current'], '已核对模板和共享页面规则，无需修改。')
    return src, work, tmp_path


def proposal(case, index=0):
    _, work, root = case
    d = Document(bp.load(work)['current'])
    d.paragraphs[index].runs[0].font.size = Pt(13)
    p = root / f'proposal-{index}.docx'
    d.save(p)
    return p


def test_group_whole_table_and_caption(tmp_path):
    d = Document()
    d.add_paragraph('表1 配置清单')
    table = d.add_table(rows=100, cols=2)
    table.cell(0, 0).text = '编号'
    table.cell(99, 1).text = '最后一行'
    d.add_paragraph('表格之后的正文')
    p = tmp_path / 'table.docx'; d.save(p)
    b = bp.make_blocks(p, target_chars=1, max_paragraphs=1)
    assert (b[0]['start'], b[0]['end'], b[0]['kind']) == (0, 2, 'object')
    assert b[1]['start'] == 2


def test_heading_attaches_body(tmp_path):
    d = Document(); d.add_heading('标题', 1); d.add_paragraph('正文'); d.add_paragraph('后文')
    p = tmp_path / 'h.docx'; d.save(p)
    b = bp.make_blocks(p, target_chars=1, max_paragraphs=1)
    assert (b[0]['start'], b[0]['end']) == (0, 2)


def test_image_attaches_caption(tmp_path):
    im = tmp_path / 'img.png'; Image.new('RGB', (20, 20)).save(im)
    d = Document(); d.add_picture(str(im)); d.add_paragraph('图1 系统结构'); d.add_paragraph('后文')
    p = tmp_path / 'i.docx'; d.save(p)
    b = bp.make_blocks(p, target_chars=1, max_paragraphs=1)
    assert (b[0]['start'], b[0]['end'], b[0]['kind']) == (0, 2, 'object')


def test_body_blocks_use_budget_not_page_numbers(case):
    _, work, _ = case
    blocks = bp.load(work)['blocks']
    assert [b['id'] for b in blocks] == ['B0001', 'B0002', 'B0003']
    assert all('page' not in b for b in blocks)


def test_checkpoint_advances_and_preserves_original(case):
    src, work, _ = case
    before = bp.sha(src)
    bp.checkpoint(work, 'B0001', proposal(case), '当前块字号已修复；业务内容不变。')
    out = bp.next_block(work)
    assert out['block']['id'] == 'B0002' and out['completed'] == 1
    assert bp.sha(src) == before == bp.load(work)['source_sha256']


def test_no_change_checkpoint_is_allowed(case):
    _, work, _ = case
    state = bp.load(work)
    bp.checkpoint(work, 'B0001', state['current'], '当前块检查合格，无修改。')
    assert bp.next_block(work)['completed'] == 1


def test_outside_body_edit_rejected(case):
    _, work, _ = case
    before = bp.load(work)
    with pytest.raises(bp.BlockError, match='超出'):
        bp.checkpoint(work, 'B0001', proposal(case, 1), '测试块外误改')
    assert bp.load(work)['current_sha256'] == before['current_sha256']
    assert bp.next_block(work)['completed'] == 0


def test_shared_style_mutation_rejected(case):
    _, work, root = case
    d = Document(bp.load(work)['current']); d.styles['Normal'].font.size = Pt(19)
    p = root / 'global.docx'; d.save(p)
    with pytest.raises(bp.BlockError, match='共享|块外'):
        bp.checkpoint(work, 'B0001', p, '不能局部改全局样式')


def test_section_change_rejected(case):
    _, work, root = case
    d = Document(bp.load(work)['current']); d.sections[0].left_margin = Pt(100)
    p = root / 'section.docx'; d.save(p)
    with pytest.raises(bp.BlockError):
        bp.checkpoint(work, 'B0001', p, '不能改页面')


def test_no_skipping_ahead(case):
    _, work, _ = case
    with pytest.raises(bp.BlockError, match='当前待修块'):
        bp.checkpoint(work, 'B0002', proposal(case, 1), '不允许跳过前一块')


def test_done_block_does_not_repeat(case):
    _, work, _ = case
    bp.checkpoint(work, 'B0001', proposal(case), '完成')
    with pytest.raises(bp.BlockError, match='已经完成'):
        bp.checkpoint(work, 'B0001', bp.load(work)['current'], '不应重复')


def test_reopen_only_one_block(case):
    _, work, _ = case
    for b in list(bp.load(work)['blocks']):
        bp.checkpoint(work, b['id'], bp.load(work)['current'], '已检查')
    bp.reopen(work, 'B0002', '该表后续发现局部问题')
    assert [b['status'] for b in bp.load(work)['blocks']] == ['checked', 'pending', 'checked']
    assert bp.next_block(work)['block']['id'] == 'B0002'


def test_paragraph_insertion_does_not_change_later_ids(case):
    _, work, root = case
    state = bp.load(work); parts, doc, body = bp.package(state['current'])
    p = E.Element(bp.Q('p')); r = E.SubElement(p, bp.Q('r')); E.SubElement(r, bp.Q('t')).text = ''  # Empty layout paragraph; new business text is not authorized.
    body.insert(1, p); parts[bp.DOC] = xml(doc)
    out = root / 'insert.docx'; save_parts(out, parts)
    bp.checkpoint(work, 'B0001', out, '当前块增加一段，其他对象保持原样。')
    blocks = bp.load(work)['blocks']
    assert blocks[1]['id'] == 'B0002' and blocks[1]['start'] == 2
    assert blocks[2]['start'] == 3


def test_direct_edit_invalidates_checkpoint(case):
    _, work, _ = case
    current = Path(bp.load(work)['current']); current.write_bytes(current.read_bytes() + b'changed')
    with pytest.raises(bp.BlockError, match='检查点之外'):
        bp.load(work)


def test_other_source_cannot_replace_baseline(case):
    _, work, _ = case
    other = proposal(case)
    with pytest.raises(bp.BlockError, match='另一份输入'):
        bp.initialize(other, work)


def test_global_only_once(case):
    _, work, _ = case
    with pytest.raises(bp.BlockError, match='已经结束'):
        bp.accept_global(work, bp.load(work)['current'], '重复全局设置')


def test_shared_image_cannot_change(tmp_path):
    im = tmp_path / 'img.png'; Image.new('RGB', (25, 20), (1, 2, 3)).save(im)
    d = Document(); d.add_picture(str(im)); d.add_paragraph('正文'); d.add_picture(str(im))
    src = tmp_path / 'img.docx'; d.save(src)
    blocks = bp.make_blocks(src, target_chars=1, max_paragraphs=1)
    parts, _, _ = bp.package(src)
    name = next(k for k in parts if k.startswith('word/media/'))
    parts[name] += b'different'
    out = tmp_path / 'bad.docx'; save_parts(out, parts)
    with pytest.raises(bp.BlockError, match='块外'):
        bp.local_scope(src, out, blocks[0])


def test_current_block_image_can_be_replaced(tmp_path):
    im = tmp_path / 'img.png'; Image.new('RGB', (25, 20)).save(im)
    d = Document(); d.add_picture(str(im)); d.add_paragraph('正文')
    src = tmp_path / 'img.docx'; d.save(src)
    parts, _, _ = bp.package(src)
    name = next(k for k in parts if k.startswith('word/media/'))
    parts[name] += b'different'
    out = tmp_path / 'replace.docx'; save_parts(out, parts)
    delta, changes = bp.local_scope(src, out, bp.make_blocks(src, target_chars=1)[0])
    assert delta == 0 and name in changes


def test_block_local_new_style_allowed(case):
    _, work, root = case
    state = bp.load(work); parts, doc, body = bp.package(state['current'])
    styles = bp.parse(parts['word/styles.xml'])
    style = E.SubElement(styles, bp.Q('style'), {bp.Q('styleId'): 'NewLocal', bp.Q('type'): 'paragraph'})
    E.SubElement(style, bp.Q('name'), {bp.Q('val'): 'Local'})
    ppr = E.Element(bp.Q('pPr')); E.SubElement(ppr, bp.Q('pStyle'), {bp.Q('val'): 'NewLocal'}); body[0].insert(0, ppr)
    parts['word/styles.xml'], parts[bp.DOC] = xml(styles), xml(doc)
    out = root / 'newstyle.docx'; save_parts(out, parts)
    bp.checkpoint(work, 'B0001', out, '仅补入当前块使用的样式定义。')
    assert bp.next_block(work)['completed'] == 1


def test_global_setup_can_change_style(tmp_path):
    src = tmp_path / 'a.docx'; d = Document(); d.add_paragraph('原内容'); d.save(src)
    w = tmp_path / 'w'; bp.initialize(src, w)
    d.styles['Normal'].font.size = Pt(12); out = tmp_path / 'b.docx'; d.save(out)
    bp.accept_global(w, out, '一次性安装已确认共享样式')
    assert bp.load(w)['global_ready']


def test_global_setup_cannot_rewrite_body(tmp_path):
    src = tmp_path / 'a.docx'; d = Document(); d.add_paragraph('原内容'); d.save(src)
    w = tmp_path / 'w'; bp.initialize(src, w)
    d.paragraphs[0].text = '其他内容'; out = tmp_path / 'b.docx'; d.save(out)
    with pytest.raises(bp.BlockError, match='不批量改正文'):
        bp.accept_global(w, out, '错误操作')


def test_unplanned_footnote_creation_rejected_locally(case):
    _, work, root = case
    state = bp.load(work); parts, doc, body = bp.package(state['current'])
    r = E.SubElement(body[0], bp.Q('r')); E.SubElement(r, bp.Q('footnoteReference'), {bp.Q('id'): '1'})
    notes = E.Element(bp.Q('footnotes'), nsmap={'w': bp.W})
    note = E.SubElement(notes, bp.Q('footnote'), {bp.Q('id'): '1'})
    p = E.SubElement(note, bp.Q('p')); r = E.SubElement(p, bp.Q('r')); E.SubElement(r, bp.Q('t')).text = '脚注说明'
    rels = bp.parse(parts[bp.RELS]); relns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    E.SubElement(rels, '{' + relns + '}Relationship', Id='rIdNotes', Type=bp.R + '/footnotes', Target='footnotes.xml')
    ct = bp.parse(parts['[Content_Types].xml']); ctns = 'http://schemas.openxmlformats.org/package/2006/content-types'
    E.SubElement(ct, '{' + ctns + '}Override', PartName='/word/footnotes.xml', ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    parts.update({bp.DOC: xml(doc), bp.RELS: xml(rels), '[Content_Types].xml': xml(ct), 'word/footnotes.xml': xml(notes)})
    out = root / 'foot.docx'; save_parts(out, parts)
    with pytest.raises(bp.BlockError, match='当前块检查失败'):
        bp.checkpoint(work, 'B0001', out, '把当前块脚注转为原生引用。')
    assert bp.next_block(work)['completed'] == 0


def test_release_binding_requires_all_blocks(case):
    _, work, _ = case
    ref = {'path': str(work / bp.STATE), 'sha256': bp.sha(work / bp.STATE)}
    with pytest.raises(bp.BlockError, match='未完成块'):
        bp.validate_final_binding({'block_session': ref})


def test_release_binding_rejects_stale_revision(case):
    _, work, _ = case
    for b in list(bp.load(work)['blocks']):
        bp.checkpoint(work, b['id'], bp.load(work)['current'], '已检查')
    state = bp.load(work)
    state['final'] = {'revision': state['revision'], 'candidate_sha256': state['current_sha256']}
    bp.write_json(work / bp.STATE, state)
    m = {'candidate_sha256': state['current_sha256'], 'block_session': {'path': str(work / bp.STATE), 'sha256': bp.sha(work / bp.STATE)}}
    assert bp.validate_final_binding(m)['revision'] == state['revision']
    bp.reopen(work, 'B0001', '新发现')
    with pytest.raises(bp.BlockError, match='进度已改变'):
        bp.validate_final_binding(m)


def test_lock_prevents_parallel_write(case):
    _, work, _ = case
    with bp.session_lock(work):
        with pytest.raises(bp.BlockError, match='另一个操作'):
            bp.checkpoint(work, 'B0001', bp.load(work)['current'], '并行写入')


def test_state_baseline_change_rejected(case):
    _, work, root = case
    t = root / 'rules.json'; t.write_text('{}')
    state = bp.load(work)
    state['settings']['baselines'] = {'rules': {'path': str(t), 'sha256': bp.sha(t)}}
    bp.write_json(work / bp.STATE, state); t.write_text('{"changed":true}')
    with pytest.raises(bp.BlockError, match='基准变了'):
        bp.load(work)


def test_ready_state_after_forward_pass(case):
    _, work, _ = case
    for b in list(bp.load(work)['blocks']):
        bp.checkpoint(work, b['id'], bp.load(work)['current'], '已检查')
    assert bp.next_block(work)['status'] == 'ready_for_final'


def test_single_pending_block_after_failure(case):
    _, work, _ = case
    bp.checkpoint(work, 'B0001', bp.load(work)['current'], '完成')
    with pytest.raises(bp.BlockError):
        bp.checkpoint(work, 'B0002', proposal(case, 2), '错误范围')
    assert bp.next_block(work)['block']['id'] == 'B0002'
    assert bp.next_block(work)['completed'] == 1


# Receipt tests below explicitly simulate only protocol outputs, not a vision model.
class FakeVisualError(Exception):
    pass


@pytest.fixture
def receipt_env(monkeypatch):
    bound = {'source_sha256': 'source', 'candidate_sha256': 'old', 'object_id': 'F1',
             'object_sha256': 'image', 'media_sha256': ['media'],
             'pages': [{'name': 'page-1.png', 'sha256': 'page'}], 'rendered_view': None}
    result = {'binding': bound, 'verdict': 'pass'}
    monkeypatch.setitem(sys.modules, 'visual_evidence', types.SimpleNamespace(
        validate_inspection=lambda *a, **k: result, VisualError=FakeVisualError))
    return result


def test_same_pixels_can_reuse_real_receipt(receipt_env):
    expected = copy.deepcopy(receipt_env['binding']); expected['candidate_sha256'] = 'new'
    row = {'inspection': 'stored', 'reuse': {'previous_candidate_sha256': 'old', 'reason': 'unchanged-image-and-page-hashes'}}
    assert be.validate_bound_inspection(row, expected, phase='final') is receipt_env


@pytest.mark.parametrize('changed', ['page', 'image', 'source', 'native', 'failure'])
def test_changed_evidence_cannot_reuse(receipt_env, changed):
    expected = copy.deepcopy(receipt_env['binding']); expected['candidate_sha256'] = 'new'
    if changed == 'page': expected['pages'][0]['sha256'] = 'different'
    if changed == 'image': expected['object_sha256'] = 'different'
    if changed == 'source': expected['source_sha256'] = 'different'
    if changed == 'native': expected['rendered_view'] = {'report': 'different'}
    if changed == 'failure': receipt_env['verdict'] = 'fail'
    row = {'inspection': 'stored', 'reuse': {'previous_candidate_sha256': 'old', 'reason': 'unchanged-image-and-page-hashes'}}
    with pytest.raises(FakeVisualError):
        be.validate_bound_inspection(row, expected, phase='final')


def test_initial_cannot_be_rebound(receipt_env):
    expected = copy.deepcopy(receipt_env['binding']); expected['candidate_sha256'] = 'new'
    row = {'inspection': 'stored', 'reuse': {'previous_candidate_sha256': 'old', 'reason': 'unchanged-image-and-page-hashes'}}
    with pytest.raises(FakeVisualError):
        be.validate_bound_inspection(row, expected, phase='initial')


def test_long_table_view_does_not_cut_repair_scope(tmp_path):
    d = Document(); d.add_paragraph('表1 明细'); table = d.add_table(rows=80, cols=2)
    for i, row in enumerate(table.rows):
        row.cells[0].text = str(i)
        row.cells[1].text = '内容' * 40
    src = tmp_path / 'long.docx'; d.save(src)
    w = tmp_path / 'work'; bp.initialize(src, w)
    a, b = bp.next_block(w), bp.next_block(w, view_offset=14)
    assert a['block']['id'] == b['block']['id']
    assert (a['block']['start'], a['block']['end']) == (b['block']['start'], b['block']['end'])
    assert len(a['block']['paragraphs']) == 14 and a['block']['view']['has_more']
    assert Path(a['block']['xml_path']).is_file()
    assert a['block']['view']['total_units'] == 81


def test_unchanged_pages_reuse_except_changed_neighbours(tmp_path):
    oldpages, newpages, oldrows, newrows = [], [], [], []
    for i in range(1, 6):
        old = tmp_path / f'old-{i}.png'; new = tmp_path / f'new-{i}.png'
        old.write_bytes(f'page{i}'.encode()); new.write_bytes(f'page{i}'.encode() if i != 3 else b'changed')
        name = f'page-{i}.png'
        oldpages.append({'name': name, 'path': str(old), 'sha256': bp.sha(old)})
        newpages.append({'name': name, 'path': str(new), 'sha256': bp.sha(new)})
        oldrows.append({'name': name, 'sha256': bp.sha(old), 'status': 'pass', 'observations': '实际观察示例'})
        newrows.append({'name': name, 'sha256': bp.sha(new), 'status': 'pending', 'observations': ''})
    v = tmp_path / 'old-review.json'; m = tmp_path / 'old-manifest.json'
    bp.write_json(v, {'candidate_sha256': 'old', 'pages': oldrows, 'figures': {'objects': []}})
    bp.write_json(m, {'candidate_sha256': 'old', 'render': {'pages': oldpages}})
    template = {'pages': newrows, 'figures': {'objects': []}}
    result = bw.reuse_unchanged_reviews({'review': str(v), 'manifest': str(m), 'candidate_sha256': 'old'}, template, {'render': {'pages': newpages}})
    assert [p['status'] for p in result['pages']] == ['pass', 'pending', 'pending', 'pending', 'pass']


def test_same_revision_final_does_not_run_refresh_again(case, monkeypatch):
    _, work, root = case
    for b in list(bp.load(work)['blocks']):
        bp.checkpoint(work, b['id'], bp.load(work)['current'], '已检查')
    state = bp.load(work)
    initial = root / 'initial.json'; initial.write_text('{}')
    manifest = root / 'manifest.json'
    bp.write_json(manifest, {'status': 'awaiting_visual_review', 'initial_visual_review_sha256': bp.sha(initial)})
    state['initial_review'] = str(initial)
    state['final'] = {'revision': state['revision'], 'candidate': state['current'],
                      'candidate_sha256': state['current_sha256'], 'manifest': str(manifest), 'review': 'unused'}
    args = types.SimpleNamespace(retry_final=False, object_plan=None, content_plan=None)
    # Cache hit must not launch any engine or load the audit subprocess chain.
    state['final']['fingerprint'] = bw.audit_fingerprint(state)
    bp.write_json(manifest, {'status':'awaiting_visual_review','candidate_sha256':state['current_sha256']})
    import block_cache
    monkeypatch.setattr(block_cache,'prepared_valid',lambda m:True)
    result, code = bw.final_audit(work, state, args)
    assert code == 0 and result['reused_prepared_audit']


def test_unfinished_blocks_cannot_enter_final(case):
    _, work, _ = case
    args = types.SimpleNamespace(retry_final=False, object_plan=None, content_plan=None)
    with pytest.raises(bp.BlockError, match='未完成块'):
        bw.final_audit(work, bp.load(work), args)


def test_resume_reads_existing_state_without_whole_repair(case, monkeypatch, capsys):
    _, work, root = case
    state = bp.load(work); state['objects_initialized'] = True; state['object_inventory_version'] = 2
    state['settings'].update(renderer='auto', field_engine='auto', uno_python=None)
    bp.write_json(work / bp.STATE, state)
    assert bw.main(['--work-dir', str(work)]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result['block']['id'] == 'B0001' and result['completed'] == 0


def test_footnote_outside_selected_range_rejected(case):
    _, work, root = case
    state = bp.load(work)
    parts, _, _ = bp.package(state['current'])
    notes = E.Element(bp.Q('footnotes'))
    n = E.SubElement(notes, bp.Q('footnote'), {bp.Q('id'): '10'})
    E.SubElement(E.SubElement(E.SubElement(n, bp.Q('p')), bp.Q('r')), bp.Q('t')).text = '非当前块脚注'
    oldparts = dict(parts); oldparts['word/footnotes.xml'] = xml(notes)
    before = root / 'notes-before.docx'; save_parts(before, oldparts)
    n.find('.//' + bp.Q('t')).text = '被误改'
    changed = dict(oldparts); changed['word/footnotes.xml'] = xml(notes)
    after = root / 'notes-after.docx'; save_parts(after, changed)
    with pytest.raises(bp.BlockError, match='脚注/尾注'):
        bp.local_scope(before, after, state['blocks'][0])
