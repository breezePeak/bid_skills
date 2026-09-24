#!/usr/bin/env python3
"""Real Office render -> bound page crop smoke test, not a vision-model test."""
import argparse
import json
import sys
import uuid
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'scripts'))
from render_docx import render
import visual_evidence as ve
import figure_inspection as fi


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--work-dir', type=Path, required=True)
    ap.add_argument('--renderer', choices=['auto', 'word', 'wps', 'libreoffice'], default='auto')
    args = ap.parse_args()
    root = args.work_dir.resolve() / ('crop-smoke-' + uuid.uuid4().hex)
    root.mkdir(parents=True)
    source = root / 'fixture.docx'
    doc = Document(); doc.add_paragraph('Visual evidence integration fixture')
    doc.add_picture(str(ROOT / 'tests/fixtures/visual/right-column-overflow.png'), width=Inches(6))
    doc.save(source)
    source_hash = ve.file_sha(source)
    rendered = render(source, root / 'render', renderer=args.renderer)
    page = Path(rendered['pages'][0]); pixels = ve.normalized_image(page)
    row = {'id': 'fixture-1', 'rendered_view': {'render_report': ve.reference(rendered['report_path']),
           'locations': [{'name': page.name, 'bbox': [0, 0, pixels.width, pixels.height]}]}}
    images, pages = fi._rendered_pixels(source, row)
    assets = []
    for index, image in enumerate(images):
        path = root / f'crop-{index}.png'; image.save(path); assets.append(path)
    # This tests the rendered-view branch, not extraction of native Shape ids.
    obj = {'id': 'fixture-1', 'hash': 'fixture-object', 'paths': []}
    inventory_document = SimpleNamespace(files={})
    expected = fi.binding(source, source, obj, inventory_document, pages=pages, rendered_view=row['rendered_view'])
    bundle_ref = ve.prepare(assets, root, expected, phase='initial', pages=pages)
    bundle = ve.validate_bundle(bundle_ref)
    fi._assert_pixels(inventory_document, obj, bundle, pages, source, row)
    invalid = {**row, 'rendered_view': {**row['rendered_view'], 'locations': [{'name': page.name, 'bbox': [-1, 0, 10, 10]}]}}
    try:
        fi._rendered_pixels(source, invalid)
    except ve.VisualError:
        pass
    else:
        raise AssertionError('Out-of-bounds crop was accepted.')
    changed = root / 'changed-copy.docx'; changed.write_bytes(source.read_bytes() + b'stale-fixture')
    try:
        fi._rendered_pixels(changed, row)
    except Exception:
        pass
    else:
        raise AssertionError('Stale source was accepted.')
    assert ve.file_sha(source) == source_hash
    result = {'status': 'passed', 'kind': 'real-render-and-crop-transport', 'renderer': rendered['engine'],
              'page_count': len(rendered['pages']), 'page': str(page), 'bundle': bundle_ref,
              'source_unchanged': True, 'bad_crop_rejected': True, 'stale_source_rejected': True,
              'note': 'Not a model vision test, not full document acceptance, not native Shape inventory regression.'}
    ve.write(root / 'smoke-results.json', result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
