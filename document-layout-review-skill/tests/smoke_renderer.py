#!/usr/bin/env python3
"""Real engine smoke check. This does not certify a user's document or field update."""
from __future__ import annotations
import argparse
import json
import platform
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from office_backends import digest
from render_docx import RenderError, render, validate_render


def make_document(path, page_count):
    from docx import Document
    doc = Document()
    for number in range(1, page_count + 1):
        if number > 1:
            doc.add_page_break()
        doc.add_heading('Renderer verification - page ' + str(number), 1)
        doc.add_paragraph('Real Office export. Source bytes must stay unchanged.')
        table = doc.add_table(rows=3, cols=3)
        table.style = 'Table Grid'
        values = [('Item', 'Requirement', 'Result'),
                  ('Priority', 'Word > WPS > LibreOffice', 'Recorded'),
                  ('Evidence', 'Fresh PDF and PNG pages', 'Required')]
        for row, texts in zip(table.rows, values):
            for cell, text in zip(row.cells, texts):
                cell.text = text
        doc.add_paragraph('End of test page ' + str(number))
    doc.save(path)


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--renderer', choices=['auto', 'word', 'wps', 'libreoffice'], default='auto')
    parser.add_argument('--work-dir', type=Path, required=True)
    args = parser.parse_args()
    args.work_dir.mkdir(parents=True, exist_ok=True)
    # Do not overwrite files from another test or user task.
    work = Path(tempfile.mkdtemp(prefix='renderer-smoke-', dir=args.work_dir)).resolve()
    source = work / 'sample with spaces.docx'
    result = {'status': 'failed', 'platform': platform.platform(),
              'renderer_requested': args.renderer, 'work_dir': str(work), 'checks': []}
    try:
        make_document(source, 3)
        before = digest(source)
        first = render(source, work / 'render', renderer=args.renderer)
        validate_render(source, first, requested=args.renderer, pages=first['page_records'])
        require(first['page_count'] == 3, 'Expected exactly three pages')
        require(digest(source) == before, 'Rendering changed source bytes')
        result['checks'].append({'test': 'real export, three pages, source unchanged', 'status': 'passed', 'render': first})
        make_document(source, 1)
        before = digest(source)
        second = render(source, work / 'render', renderer=args.renderer)
        validate_render(source, second, requested=args.renderer, pages=second['page_records'])
        require(second['page_count'] == 1, 'Old page 2/3 was reused')
        require(second['pdf'] != first['pdf'], 'Run directories were reused')
        require(digest(source) == before, 'Rendering changed source bytes')
        result['checks'].append({'test': 'rerender same input/output root, one page, fresh evidence', 'status': 'passed', 'render': second})
        try:
            validate_render(source, first, requested=args.renderer)
        except RenderError as exc:
            require(exc.code == 'render-stale', 'Unexpected old-report rejection: ' + exc.code)
        else:
            raise AssertionError('Old report was incorrectly accepted')
        result['checks'].append({'test': 'changed-source rejects old report', 'status': 'passed'})
        result.update(status='passed', actual_engines=[first['engine'], second['engine']],
                      visual_acceptance='not_performed', field_refresh='not_tested')
    except Exception as exc:
        result['error'] = exc.as_issue() if hasattr(exc, 'as_issue') else str(exc)
    output = work / 'smoke-results.json'
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'status': result['status'], 'report': str(output),
                      'actual_engines': result.get('actual_engines'), 'error': result.get('error')},
                     ensure_ascii=False, indent=2))
    return 0 if result['status'] == 'passed' else 2


if __name__ == '__main__':
    raise SystemExit(main())
