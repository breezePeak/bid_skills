"""Caption invariants and real Office roundtrips, not a whole-skill acceptance test.

Unit tests: python -m unittest discover -s tests -p 'test_caption_baseline_regression.py' -v
Actual Office tests: set DLR_RUN_OFFICE_TESTS=1 and DLR_FIELD_ENGINE=word/libreoffice.
Word mode requires Windows Word and the repository's refresh_fields_word.ps1.
"""
from __future__ import annotations
import copy
import importlib.util
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import patch
import zipfile
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.shared import Inches
from lxml import etree as E
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import caption_field_guard as guard
import template_caption_policy as captions

Q, NS = guard.Q, guard.NS


def node(tag, **attrs):
    n = OxmlElement('w:' + tag)
    for k, v in attrs.items():
        n.set(Q(k), str(v))
    return n


def add_field(p, instruction, cached='', *, complex_field=False, locked=False, hidden=False):
    if complex_field:
        begin = node('fldChar', fldCharType='begin')
        if locked:
            begin.set(Q('fldLock'), 'true')
        r = node('r'); r.append(begin); p._p.append(r)
        for fragment in (instruction[:4], instruction[4:11], instruction[11:]):
            r = node('r'); n = node('instrText'); n.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); n.text = fragment; r.append(n); p._p.append(r)
        r = node('r'); r.append(node('fldChar', fldCharType='separate')); p._p.append(r)
        r = node('r')
        if hidden:
            rp = node('rPr'); rp.append(node('vanish')); r.append(rp)
        t = node('t'); t.text = cached; r.append(t); p._p.append(r)
        r = node('r'); r.append(node('fldChar', fldCharType='end')); p._p.append(r)
        return begin
    f = node('fldSimple', instr=' ' + instruction + ' ')
    if locked:
        f.set(Q('fldLock'), 'true')
    r = node('r')
    if hidden:
        rp = node('rPr'); rp.append(node('vanish')); r.append(rp)
    t = node('t'); t.text = cached; r.append(t); f.append(r); p._p.append(f)
    return f


def install_heading_numbers(doc):
    nr = doc.part.numbering_part.element
    abstract = node('abstractNum', abstractNumId='900')
    abstract.append(node('multiLevelType', val='multilevel'))
    lv = node('lvl', ilvl='0')
    for tag, val in [('start','1'), ('numFmt','decimal'), ('pStyle','Heading1'), ('suff','space'), ('lvlText','%1')]:
        lv.append(node(tag, val=val))
    abstract.append(lv)
    first_num = next((i for i, c in enumerate(nr) if c.tag == Q('num')), len(nr))
    nr.insert(first_num, abstract)
    num = node('num', numId='900'); num.append(node('abstractNumId', val='900')); nr.append(num)
    pp = doc.styles['Heading 1'].element.get_or_add_pPr()
    np = node('numPr'); np.append(node('ilvl', val='0')); np.append(node('numId', val='900')); pp.append(np)


def make_fixture(folder: Path, *, complex_fields=False, resets=True) -> Path:
    image = folder / 'fixture.png'
    Image.new('RGB', (24, 12), 'white').save(image)
    d = Document(); install_heading_numbers(d)
    for chapter in range(1, 3):
        heading = d.add_heading('章节标题_' + str(chapter), level=1)
        if resets:
            for label in ('图', '表'):
                add_field(heading, 'SEQ ' + label + r' \r 0 \h', hidden=True, complex_field=complex_fields)
        for i in range(1, 3):
            p = d.add_paragraph(style='Caption'); p.add_run('表')
            add_field(p, r'SEQ 表 \* ARABIC \s 1', str(i), complex_field=complex_fields)
            p.add_run(' 表格内容_' + str(chapter) + '_' + str(i))
            t = d.add_table(rows=1, cols=2); t.cell(0,0).text='项目'; t.cell(0,1).text='内容'
            d.add_paragraph('普通说明文字。')
            p = d.add_paragraph(); p.add_run().add_picture(str(image), width=Inches(.5))
            p = d.add_paragraph(style='Caption'); p.add_run('图')
            add_field(p, r'SEQ 图 \* ARABIC \s 1', str(i), complex_field=complex_fields)
            p.add_run(' 图片内容_' + str(chapter) + '_' + str(i))
            d.add_paragraph('后续说明文字。')
    path = folder/'source.docx'; d.save(path); return path


def rewrite(source, target, change):
    with zipfile.ZipFile(source) as z:
        content = {i.filename:z.read(i.filename) for i in z.infolist()}
    roots = {k:guard.parse(content[k]) for k in ('word/document.xml','word/styles.xml','word/numbering.xml')}
    change(roots)
    for name, root in roots.items():
        content[name] = E.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in content.items(): z.writestr(name,data)
    return target


def caption_paragraphs(root):
    return [p for p in root.iter(Q('p')) if p.find('w:pPr/w:pStyle',NS) is not None
            and guard.value(p.find('w:pPr/w:pStyle',NS)) == 'Caption']


def rename_titles(roots):
    for t in roots['word/document.xml'].iter(Q('t')):
        if (t.text or '').startswith('章节标题'):
            t.text = (t.text or '').replace('章节标题','重命名后的项目实施方案',1)


def rename_style(roots):
    style = roots['word/styles.xml'].find('w:style[@w:styleId="Heading1"]', NS)
    style.find(Q('name')).set(Q('val'),'项目章节_已重命名')


def change_number_format(roots):
    for lv in roots['word/numbering.xml'].findall('w:abstractNum/w:lvl[@w:ilvl="0"]',NS):
        lv.find(Q('numFmt')).set(Q('val'),'chineseCounting')
        lv.find(Q('lvlText')).set(Q('val'),'第%1章')


def delete_first_objects(roots):
    body=roots['word/document.xml'].find(Q('body'))
    table=next(n for n in body if n.tag==Q('tbl'))
    cap=table.getprevious(); body.remove(cap); body.remove(table)
    picture=next(n for n in body if n.tag==Q('p') and n.xpath('.//w:drawing',namespaces=NS))
    cap=picture.getnext(); body.remove(cap); body.remove(picture)


def insert_first_objects(roots):
    body=roots['word/document.xml'].find(Q('body'))
    table=next(n for n in body if n.tag==Q('tbl'))
    cap=table.getprevious()
    cap.addprevious(copy.deepcopy(cap)); cap.addprevious(copy.deepcopy(table))
    picture=next(n for n in body if n.tag==Q('p') and n.xpath('.//w:drawing',namespaces=NS))
    copied=copy.deepcopy(picture)
    for e in copied.iter():
        if E.QName(e).localname=='docPr': e.set('id','1001')
    picture.addprevious(copied); picture.addprevious(copy.deepcopy(picture.getnext()))


class TemporaryCase(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.folder=Path(self.tmp.name)


class PlacementTests(TemporaryCase):
    def template(self, direction='before'):
        d=Document(); p=d.add_paragraph('表1 清单',style='Caption'); t=d.add_table(rows=1,cols=1)
        if direction=='after': t._tbl.addnext(p._p)
        path=self.folder/'template.docx'; d.save(path); return path

    def test_table_caption_is_above_even_when_template_is_wrong(self):
        self.assertEqual(captions.placement_rules(self.template('after')),guard.PLACEMENT)

    def test_fixed_figure_below_table_above(self):
        self.assertEqual(captions.placement_rules(self.template()),{'figure':'after','table':'before'})

    def test_template_without_examples_does_not_request_position_choice(self):
        d=Document(); p=self.folder/'empty.docx'; d.save(p)
        self.assertEqual(captions.placement_rules(p),guard.PLACEMENT)

    def test_old_reversed_profile_cannot_override_baseline(self):
        p=self.folder/'profile.json'; p.write_text(json.dumps({'caption_placement':{'figure':'before','table':'after'}}))
        self.assertEqual(captions.placement_rules(self.template(),p),guard.PLACEMENT)

    def test_deprecated_invalid_position_is_not_an_active_rule(self):
        p=self.folder/'profile.json'; p.write_text(json.dumps({'caption_placement':{'table':'guess'}}))
        self.assertEqual(captions.placement_rules(self.template(),p,{'table'}),{'table':'before'})

    def test_unknown_object_kind_is_rejected(self):
        with self.assertRaises(ValueError): captions.placement_rules(self.template(),required={'chart'})

    def test_core_position_failure_is_not_filtered(self):
        source=make_fixture(self.folder)
        fake=types.ModuleType('numbering_policy')
        fake.audit_document=lambda *a,**k:{'status':'failed','issues':[{'code':'caption-position','severity':'error'}]}
        with patch.dict(sys.modules,{'numbering_policy':fake}):
            result=captions.audit_document(source,self.template())
        self.assertEqual(result['status'],'failed')
        self.assertIn('caption-position',{x['code'] for x in result['issues']})

    def test_core_failed_status_cannot_be_laundered_as_pass(self):
        source=make_fixture(self.folder)
        fake=types.ModuleType('numbering_policy'); fake.audit_document=lambda *a,**k:{'status':'failed','issues':[]}
        with patch.dict(sys.modules,{'numbering_policy':fake}): result=captions.audit_document(source,self.template())
        self.assertEqual(result['status'],'failed')


class FieldContractTests(TemporaryCase):
    def source(self, **kwargs): return make_fixture(self.folder,**kwargs)

    def mutate(self,change,**kwargs):
        return rewrite(self.source(**kwargs),self.folder/'changed.docx',change)

    def set_instruction(self,instruction):
        def change(roots):
            caption_paragraphs(roots['word/document.xml'])[0].find(Q('fldSimple')).set(Q('instr'),instruction)
        return self.mutate(change)

    def test_native_simple_fields_pass(self):
        result=guard.audit(self.source()); self.assertEqual(result['status'],'passed',result['issues']); self.assertEqual(result['caption_count'],8)

    def test_native_complex_split_instruction_fields_pass(self):
        result=guard.audit(self.source(complex_fields=True)); self.assertEqual(result['status'],'passed',result['issues'])

    def test_fixed_fields_do_not_contain_heading_names(self):
        for row in guard.audit(self.source())['captions']:
            self.assertNotIn('章节标题',row['instruction']); self.assertNotIn('Heading',row['instruction'])

    def test_title_text_edit_preserves_field_contract(self):
        result=guard.audit(self.mutate(rename_titles)); self.assertEqual(result['status'],'passed',result['issues'])

    def test_heading_style_display_name_edit_preserves_field_contract(self):
        result=guard.audit(self.mutate(rename_style)); self.assertEqual(result['status'],'passed',result['issues'])

    def test_heading_number_format_edit_preserves_field_contract(self):
        result=guard.audit(self.mutate(change_number_format)); self.assertEqual(result['status'],'passed',result['issues'])

    def test_heading_name_cannot_be_a_seq_identifier(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ "章节标题_1" \* ARABIC \s 1'))['status'],'failed')

    def test_heading_name_cannot_be_reset_argument(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ 表 \* ARABIC \s "Heading 1"'))['status'],'failed')

    def test_extra_bookmark_argument_blocked_despite_good_cache(self):
        result=guard.audit(self.set_instruction(r'SEQ 表 DeletedHeadingBookmark \* ARABIC \s 1'),check_cache=False)
        self.assertIn('caption-seq-dependency',{i['code'] for i in result['issues']})

    def test_wrong_restart_level_is_blocked(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ 表 \* ARABIC \s 2'))['status'],'failed')

    def test_fixed_restart_on_each_caption_is_blocked(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ 表 \r 1'))['status'],'failed')

    def test_duplicate_restart_switch_is_blocked(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ 表 \s 1 \s 1'))['status'],'failed')

    def test_visible_hidden_seq_is_blocked(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ 表 \s 1 \h'))['status'],'failed')

    def test_wrong_number_format_is_blocked(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ 表 \s 1 \* ROMAN'))['status'],'failed')

    def test_case_and_switch_order_supported(self):
        result=guard.audit(self.set_instruction(r'seq 表 \s 1 \* Arabic \* MERGEFORMAT'))
        self.assertEqual(result['status'],'passed',result['issues'])

    def test_missing_switch_argument_is_blocked(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ 表 \s'))['status'],'failed')

    def test_bad_quote_is_blocked(self):
        self.assertEqual(guard.audit(self.set_instruction('SEQ 表 \\s 1 "'))['status'],'failed')

    def test_misplaced_backslash_is_blocked(self):
        self.assertEqual(guard.audit(self.set_instruction(r'SEQ 表 s\ 1'))['status'],'failed')

    def test_locked_field_cannot_mask_stale_cache(self):
        def change(roots): caption_paragraphs(roots['word/document.xml'])[0].find(Q('fldSimple')).set(Q('fldLock'),'true')
        self.assertEqual(guard.audit(self.mutate(change))['status'],'failed')

    def test_styleref_prefix_is_blocked_even_when_cache_looks_good(self):
        def change(roots):
            p=caption_paragraphs(roots['word/document.xml'])[0]
            f=node('fldSimple',instr='STYLEREF "Deleted Heading Style" \\n'); r=node('r'); t=node('t'); t.text='';r.append(t);f.append(r);p.insert(1,f)
        result=guard.audit(self.mutate(change))
        self.assertIn('caption-heading-dependency',{i['code'] for i in result['issues']})

    def test_ref_prefix_is_blocked(self):
        def change(roots):
            p=caption_paragraphs(roots['word/document.xml'])[0];p.insert(1,node('fldSimple',instr='REF DeletedHeading'))
        self.assertEqual(guard.audit(self.mutate(change))['status'],'failed')

    def test_unclosed_complex_field_is_blocked(self):
        def change(roots):
            p=caption_paragraphs(roots['word/document.xml'])[0]
            mark=p.find('.//w:fldChar[@w:fldCharType="end"]',NS);mark.getparent().remove(mark)
        self.assertEqual(guard.audit(self.mutate(change,complex_fields=True))['status'],'failed')

    def test_orphan_instruction_is_blocked(self):
        def change(roots):
            p=caption_paragraphs(roots['word/document.xml'])[0];r=node('r');i=node('instrText');i.text='STYLEREF Missing';r.append(i);p.append(r)
        self.assertEqual(guard.audit(self.mutate(change))['status'],'failed')

    def test_native_seq_required_not_literal_field_code(self):
        def change(roots):
            p=caption_paragraphs(roots['word/document.xml'])[0];f=p.find(Q('fldSimple'));p.remove(f)
            r=node('r');t=node('t');t.text='{ SEQ 表 \\s 1 }';r.append(t);p.insert(2,r)
        self.assertEqual(guard.audit(self.mutate(change))['status'],'failed')

    def test_chinese_style_error_result_is_blocked(self):
        def change(roots): caption_paragraphs(roots['word/document.xml'])[0].find('.//w:fldSimple/w:r/w:t',NS).text='错误！文档中没有指定样式的文字。'
        self.assertEqual(guard.audit(self.mutate(change))['status'],'failed')

    def test_wrong_sequence_after_edit_requires_actual_refresh(self):
        changed=self.mutate(delete_first_objects)
        self.assertEqual(guard.audit(changed,check_cache=False)['status'],'passed')
        self.assertEqual(guard.audit(changed)['status'],'failed')

    def test_unrelated_styleref_in_normal_text_is_not_removed(self):
        def change(roots):
            p=node('p'); f=node('fldSimple',instr='STYLEREF "Heading 1"');p.append(f)
            roots['word/document.xml'].find(Q('body')).insert(1,p)
        self.assertEqual(guard.audit(self.mutate(change))['status'],'passed')

    def test_hidden_reset_in_body_is_blocked(self):
        def change(roots):
            p=node('p');p.append(node('fldSimple',instr='SEQ 图 \\r 0 \\h'))
            roots['word/document.xml'].find(Q('body')).insert(2,p)
        self.assertEqual(guard.audit(self.mutate(change))['status'],'failed')

    def test_duplicate_reset_is_blocked(self):
        def change(roots):
            p=next(p for p in roots['word/document.xml'].iter(Q('p')) if p.find(Q('fldSimple')) is not None)
            p.append(copy.deepcopy(p.find(Q('fldSimple'))))
        self.assertEqual(guard.audit(self.mutate(change))['status'],'failed')



class NativePreservationTests(TemporaryCase):
    def export_with_lost_switches(self):
        source=make_fixture(self.folder)
        def change(roots):
            for p in roots['word/document.xml'].iter(Q('p')):
                for f in p.findall(Q('fldSimple')):
                    kind=guard.sequence_kind(guard.value(f,'instr',''))
                    if kind:
                        f.set(Q('instr'),'SEQ '+guard.LABELS[kind]+' \\* ARABIC')
        return source,rewrite(source,self.folder/'office.docx',change)

    def test_restore_switches_without_recalculating_any_cache(self):
        source,office=self.export_with_lost_switches();out=self.folder/'result.docx'
        with zipfile.ZipFile(office) as z:
            before=[r[3].cached for r in guard._native_entries(guard.parse(z.read('word/document.xml')))]
        result=guard.preserve_native_instructions(source,office,out)
        with zipfile.ZipFile(out) as z:
            after=[r[3].cached for r in guard._native_entries(guard.parse(z.read('word/document.xml')))]
        self.assertEqual(before,after)
        self.assertEqual(result['instruction_restore_count'],12)
        self.assertEqual(guard.audit(out)['status'],'passed')

    def test_wrong_engine_sequence_is_not_fixed_by_writing_expected_cache(self):
        source,office=self.export_with_lost_switches()
        def change(roots):
            caption_paragraphs(roots['word/document.xml'])[0].find('.//w:fldSimple/w:r/w:t',NS).text='9'
        bad=rewrite(office,self.folder/'bad.docx',change);out=self.folder/'result.docx'
        with self.assertRaises(guard.CaptionFieldError):guard.preserve_native_instructions(source,bad,out)
        self.assertFalse(out.exists())

    def test_changed_caption_title_breaks_anchor(self):
        source,office=self.export_with_lost_switches()
        def change(roots):
            p=caption_paragraphs(roots['word/document.xml'])[0];list(p.iter(Q('t')))[-1].text=' 未授权的改写'
        bad=rewrite(office,self.folder/'bad.docx',change)
        with self.assertRaises(guard.CaptionFieldError) as cm:guard.preserve_native_instructions(source,bad,self.folder/'out.docx')
        self.assertEqual(cm.exception.code,'caption-roundtrip-anchor')

    def test_removed_field_is_not_recreated_from_old_cache(self):
        source,office=self.export_with_lost_switches()
        def change(roots):
            p=caption_paragraphs(roots['word/document.xml'])[0];p.remove(p.find(Q('fldSimple')))
        bad=rewrite(office,self.folder/'bad.docx',change)
        with self.assertRaises(guard.CaptionFieldError) as cm:guard.preserve_native_instructions(source,bad,self.folder/'out.docx')
        self.assertEqual(cm.exception.code,'caption-roundtrip-count')

    def test_no_instruction_change_is_byte_identical_copy(self):
        source=make_fixture(self.folder);out=self.folder/'out.docx'
        result=guard.preserve_native_instructions(source,source,out)
        self.assertEqual(result['instruction_restore_count'],0)
        self.assertEqual(source.read_bytes(),out.read_bytes())

    def test_original_cannot_be_overwritten(self):
        source,office=self.export_with_lost_switches();before=source.read_bytes()
        with self.assertRaises(guard.CaptionFieldError):guard.preserve_native_instructions(source,office,source)
        self.assertEqual(source.read_bytes(),before)

    def test_lost_restart_output_cannot_pass_on_cache_only(self):
        _,office=self.export_with_lost_switches()
        self.assertEqual(guard.audit(office)['status'],'failed')


class RefreshWiringTests(TemporaryCase):
    def load_refresh(self):
        # Only unmodified numbering-core parsing helpers are replaced here.
        # These wiring tests do not mock a successful Office update.
        fake=types.ModuleType('numbering_policy')
        fake.PolicyError=guard.CaptionFieldError;fake.file_digest=guard.digest
        fake.parse=guard.parse;fake.Q=guard.Q;fake.ERROR_TEXT=guard.ERROR_TEXT
        fake.fields=lambda p:[{'instruction':f.instruction,'cached':f.cached} for f in guard.fields(p)[0]]
        fake.command=lambda f:guard.command(f['instruction'])
        context=patch.dict(sys.modules,{'numbering_policy':fake});context.start();self.addCleanup(context.stop)
        spec=importlib.util.spec_from_file_location('caption_test_refresh',SCRIPTS/'field_refresh.py')
        mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod);return mod

    def test_malformed_caption_blocks_office_call(self):
        source=make_fixture(self.folder)
        def change(roots):caption_paragraphs(roots['word/document.xml'])[0].find(Q('fldSimple')).set(Q('instr'),'SEQ 表 DeletedBookmark')
        bad=rewrite(source,self.folder/'bad.docx',change);refresh=self.load_refresh()
        with patch.object(refresh.subprocess,'run') as engine:
            with self.assertRaises(guard.CaptionFieldError):refresh.refresh(bad,self.folder/'out.docx','libreoffice',preflight={'engine':'libreoffice'})
            engine.assert_not_called()

    def test_claimed_pass_report_cannot_hide_stripped_native_switches(self):
        source=make_fixture(self.folder)
        def change(roots):caption_paragraphs(roots['word/document.xml'])[0].find(Q('fldSimple')).set(Q('instr'),'SEQ 表 \\* ARABIC')
        bad=rewrite(source,self.folder/'bad.docx',change);refresh=self.load_refresh()
        report={'status':'passed','engine':'LibreOffice UNO','fields_updated':True,'indexes_updated':True,'errors':[],'output_sha256':guard.digest(bad)}
        with self.assertRaises(guard.CaptionFieldError):refresh.validate_report(bad,report)


@unittest.skipUnless(os.environ.get('DLR_RUN_OFFICE_TESTS')=='1','actual Office tests require DLR_RUN_OFFICE_TESTS=1')
class ActualOfficeTests(TemporaryCase):
    def update(self,source,stem):
        output=self.folder/(stem+'.docx');report=self.folder/(stem+'.json')
        engine=os.environ.get('DLR_FIELD_ENGINE','word' if os.name=='nt' else 'libreoffice')
        if engine=='word':
            if os.name!='nt': self.fail('明确要求 Word，当前不是 Windows；未回退其他引擎。')
            shell=shutil.which('powershell') or shutil.which('pwsh')
            cmd=[shell,'-NoProfile','-ExecutionPolicy','Bypass','-File',str(SCRIPTS/'refresh_fields_word.ps1'),'-InputPath',str(source),'-OutputPath',str(output),'-ReportPath',str(report)]
        elif engine=='libreoffice':
            python=os.environ.get('DLR_UNO_PYTHON','/usr/bin/python3')
            cmd=[python,str(SCRIPTS/'refresh_fields_uno.py'),str(source),str(output),str(report),'--soffice',shutil.which('libreoffice') or 'soffice']
        else: self.fail('未知实际更新引擎')
        completed=subprocess.run(cmd,capture_output=True,text=True,timeout=120)
        data=json.loads(report.read_text(encoding='utf-8-sig')) if report.is_file() else {}
        self.assertEqual(completed.returncode,0,{'engine':engine,'stdout':completed.stdout,'stderr':completed.stderr,'report':data})
        self.assertEqual(data.get('status'),'passed',data)
        preserved=self.folder/(stem+'-native.docx')
        preservation=guard.preserve_native_instructions(source,output,preserved)
        output=preserved
        result=guard.require_valid(output)
        evidence=os.environ.get('DLR_EVIDENCE_DIR')
        if evidence:
            dest=Path(evidence);dest.mkdir(parents=True,exist_ok=True)
            (dest/(self._testMethodName+'-'+stem+'.json')).write_text(json.dumps({'office':data,'native_preservation':preservation,'caption_check':result},ensure_ascii=False,indent=2),encoding='utf-8')
        return output,result

    def roundtrip(self,change,*,complex_fields=False):
        source=make_fixture(self.folder,complex_fields=complex_fields)
        original,_=self.update(source,'initial-refresh')
        edited=rewrite(original,self.folder/'edited.docx',change)
        guard.require_valid(edited,check_cache=False)
        refreshed,first=self.update(edited,'after-edit')
        _,second=self.update(refreshed,'second-refresh')
        self.assertEqual([(x['chapter'],x['kind'],x['number']) for x in first['captions']],[(x['chapter'],x['kind'],x['number']) for x in second['captions']])

    def test_title_edit_then_two_real_refreshes(self): self.roundtrip(rename_titles)
    def test_style_rename_then_two_real_refreshes(self): self.roundtrip(rename_style)
    def test_heading_number_format_change_then_two_real_refreshes(self): self.roundtrip(change_number_format)
    def test_delete_first_figure_and_table_then_two_real_refreshes(self): self.roundtrip(delete_first_objects)
    def test_insert_first_figure_and_table_then_two_real_refreshes(self): self.roundtrip(insert_first_objects)
    def test_complex_fields_title_edit_then_two_real_refreshes(self): self.roundtrip(rename_titles,complex_fields=True)


if __name__=='__main__': unittest.main()
