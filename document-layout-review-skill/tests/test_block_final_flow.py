"""Final-stage orchestration tests. Office/vision are explicit test doubles."""
import io
import json
import sys
import types
from contextlib import redirect_stdout
from pathlib import Path

from docx import Document
import pytest

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import block_progress as bp
import block_workflow as bw
import block_cache as bc


def module(monkeypatch,name,**items):
    value=types.ModuleType(name);value.__dict__.update(items)
    monkeypatch.setitem(sys.modules,name,value)
    return value


def setup_final(tmp_path,monkeypatch):
    src=tmp_path/'source.docx';d=Document();d.add_paragraph('工期30天');d.save(src)
    work=tmp_path/'work';bp.initialize(src,work,{'renderer':'auto','field_engine':'auto','template':str(src),'template_style_json':'style','text_rules':'rules'})
    bp.accept_global(work,bp.load(work)['current'],'全局规则核对')
    bp.checkpoint(work,'B0001',bp.load(work)['current'],'当前块核对')
    state=bp.load(work);reports=work/'reports';reports.mkdir(exist_ok=True)
    for name in ('initial_review','content_plan'):
        path=reports/(name+'.json');bp.write_json(path,{'version':1,'source_sha256':state['source_sha256']});state[name]=str(path)
    state['discovery_ledger']=str(reports/'image-discoveries.json');bp.write_json(state['discovery_ledger'],{'inspections':[]})
    bp.write_json(work/bp.STATE,state)
    calls={'field':0,'audit':0};fail={'audit':False,'refresh':False}
    def refresh(source,out,*a,**k):
        calls['field']+=1
        if fail['refresh']:raise RuntimeError('Office boundary failed')
        Path(out).write_bytes(Path(source).read_bytes())
        return {'source_sha256':bp.sha(source),'output_sha256':bp.sha(out)}
    def valid(candidate,report):
        result=bp.read_json(report)
        assert result['output_sha256']==bp.sha(candidate)
        return result
    def audit():
        calls['audit']+=1
        root=Path(sys.argv[sys.argv.index('--work-dir')+1]);rep=root/'reports';rep.mkdir(exist_ok=True)
        current=Path(sys.argv[1]);candidate=root/'candidate.docx';candidate.write_bytes(current.read_bytes())
        gate=rep/'gate.json';bp.write_json(gate,{'status':'passed'})
        page=root/'page-1.png';page.write_bytes(b'explicit fake page for binding only')
        m={'status':'failed' if fail['audit'] else 'awaiting_visual_review',
           'source_sha256':state['source_sha256'],'source':state['source'],
           'candidate':str(candidate),'candidate_sha256':bp.sha(candidate),
           'template_sha256':'T','template_style_sha256':'P','text_rules_sha256':'R',
           'gates':[{'id':'fixture','passed':not fail['audit'],'report':str(gate),'report_sha256':bp.sha(gate)}],
           'render':{'passed':not fail['audit'],'data':{},'pages':[{'name':page.name,'path':str(page),'sha256':bp.sha(page)}]}}
        bp.write_json(root/'review-manifest.json',m)
        bp.write_json(rep/'final-visual-review.template.json',{'candidate_sha256':m['candidate_sha256'],'overall_status':'pending','pages':[{'name':page.name,'sha256':bp.sha(page),'status':'pending','observations':''}],'figures':{'objects':[]},'table_reviews':[]})
        print(json.dumps({'status':m['status']}))
        return 2 if fail['audit'] else 0
    module(monkeypatch,'figure_review',validate_initial=lambda *a:{})
    module(monkeypatch,'field_refresh',refresh=refresh,validate_report=valid)
    module(monkeypatch,'review_pipeline',_audit_only_main=audit)
    module(monkeypatch,'render_docx',validate_render=lambda *a,**k:None)
    module(monkeypatch,'figure_inspection',ledger=lambda s,p:{'path':str(p),'sha256':bp.sha(p)})
    module(monkeypatch,'numbering_policy',Doc=lambda p:p,inventory=lambda d:[])
    args=types.SimpleNamespace(retry_final=False)
    return work,calls,fail,args


def test_same_final_candidate_does_not_refresh_or_audit_again(tmp_path,monkeypatch):
    work,calls,fail,args=setup_final(tmp_path,monkeypatch)
    bw.final_audit(work,bp.load(work),args)
    result,code=bw.final_audit(work,bp.load(work),args)
    assert result['reused_prepared_audit'] and code==0
    assert calls=={'field':1,'audit':1}


def test_audit_retry_reuses_successful_field_update(tmp_path,monkeypatch):
    work,calls,fail,args=setup_final(tmp_path,monkeypatch);fail['audit']=True
    assert bw.final_audit(work,bp.load(work),args)[1]==2
    fail['audit']=False;args.retry_final=True
    assert bw.final_audit(work,bp.load(work),args)[1]==0
    assert calls=={'field':1,'audit':2}


def test_failed_office_step_never_becomes_cached_success(tmp_path,monkeypatch):
    work,calls,fail,args=setup_final(tmp_path,monkeypatch);fail['refresh']=True
    with pytest.raises(RuntimeError):bw.final_audit(work,bp.load(work),args)
    fail['refresh']=False
    assert bw.final_audit(work,bp.load(work),args)[1]==0
    assert calls=={'field':2,'audit':1}


def test_nested_audit_json_is_captured_not_mixed_with_outer_result(tmp_path,monkeypatch):
    work,calls,fail,args=setup_final(tmp_path,monkeypatch)
    output=io.StringIO()
    with redirect_stdout(output):bw.final_audit(work,bp.load(work),args)
    assert output.getvalue()==''
    log=Path(bp.load(work)['final_attempt']['path'])/'audit-stdout.txt'
    assert json.loads(log.read_text())['status']=='awaiting_visual_review'


def test_late_defect_ledger_survives_rounds(tmp_path,monkeypatch):
    work,calls,fail,args=setup_final(tmp_path,monkeypatch)
    ledger=bp.load(work)['discovery_ledger'];bp.write_json(ledger,{'inspections':['explicit failed-call fixture']})
    bw.final_audit(work,bp.load(work),args)
    current=bp.load(work);bp.reopen(work,'B0001','终检发现局部问题')
    bp.checkpoint(work,'B0001',current['current'],'局部验证后继续')
    bw.final_audit(work,bp.load(work),args)
    m=bp.read_json(bp.load(work)['final']['manifest'])
    assert m['image_discovery_ledger']['path']==ledger
    assert bp.read_json(ledger)['inspections']==['explicit failed-call fixture']


def test_tampered_gate_report_invalidates_preparation_cache(tmp_path,monkeypatch):
    work,calls,fail,args=setup_final(tmp_path,monkeypatch)
    bw.final_audit(work,bp.load(work),args)
    m=bp.read_json(bp.load(work)['final']['manifest'])
    Path(m['gates'][0]['report']).write_text('changed')
    result,code=bw.final_audit(work,bp.load(work),args)
    assert not result.get('reused_prepared_audit')
    assert calls=={'field':1,'audit':2}


def test_empty_or_incomplete_success_cache_is_rejected(tmp_path):
    d=tmp_path/'d';d.write_bytes(b'fixture')
    m={'status':'awaiting_visual_review','candidate':str(d),'candidate_sha256':bp.sha(d),'gates':[]}
    assert not bc.prepared_valid(m)


def test_unfinished_block_cannot_prepare_final(tmp_path,monkeypatch):
    work,calls,fail,args=setup_final(tmp_path,monkeypatch);bp.reopen(work,'B0001','待修')
    with pytest.raises(bp.BlockError):bw.final_audit(work,bp.load(work),args)
    assert calls=={'field':0,'audit':0}


def test_legacy_session_adds_missing_story_without_erasing_body_progress(tmp_path):
    source=tmp_path/'source.docx';d=Document();d.add_paragraph('正文');d.sections[0].header.paragraphs[0].text='页眉';d.save(source)
    work=tmp_path/'work';bp.initialize(source,work);s=bp.load(work)
    s.pop('coverage_version');s['blocks']=[b for b in s['blocks'] if b['part']==bp.DOC]
    s['blocks'][0]['status']='checked';s['global_ready']=True;bp.write_json(work/bp.STATE,s)
    resumed=bp.load(work)
    assert resumed['blocks'][0]['status']=='checked'
    assert resumed['blocks'][1]['part']=='word/header1.xml' and resumed['blocks'][1]['status']=='pending'


def test_partial_caption_plans_are_accumulated_across_blocks(tmp_path,monkeypatch):
    work,calls,fail,args=setup_final(tmp_path,monkeypatch)
    module(monkeypatch,'numbering_policy',Doc=lambda p:p,load_object_plan=lambda *a:None)
    module(monkeypatch,'content_integrity',make_plan=lambda source,p:{'allowed_caption_insertions':p['objects']},validate_movements=lambda *a:None)
    plan=tmp_path/'partial.json'
    for ident in ('F0001','T0001'):
        state=bp.load(work);bp.write_json(plan,{'version':1,'source_sha256':state['source_sha256'],'objects':[{'id':ident,'title':'Title '+ident,'object_sha256':'fixture'}]})
        bw.persist_plans(work,state,types.SimpleNamespace(object_plan=plan,content_plan=None))
    state=bp.load(work)
    assert {r['id'] for r in bp.read_json(state['object_plan'])['objects']}=={'F0001','T0001'}
    assert {r['id'] for r in bp.read_json(state['content_plan'])['allowed_caption_insertions']}=={'F0001','T0001'}


def test_local_known_bad_picture_cannot_advance_unchanged(tmp_path,monkeypatch):
    initial=tmp_path/'initial.json';bp.write_json(initial,{'objects':[{'id':'F0001','object_sha256':'original','checks':{'text_inside_bounds':'fail'}}]})
    module(monkeypatch,'numbering_policy',Doc=lambda p:p,inventory=lambda d:[{'id':'F0001','hash':'original'}])
    module(monkeypatch,'visual_evidence',INTRINSIC={'text_inside_bounds'},validate_inspection=lambda *a:None,validate_bundle=lambda *a:None,checked_file=lambda *a:None,normalized_image=lambda *a:None,pixel_sha=lambda *a:None)
    with pytest.raises(bp.BlockError,match='尚未实际修复'):
        bw.require_local_image_change({'initial_review':str(initial)},{'figures':['F0001']},tmp_path/'candidate.docx')
