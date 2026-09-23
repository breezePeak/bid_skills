import json
import unittest
from pathlib import Path
from docx import Document
from test_numbering_policy import Fixtures, rewrite, add_native_numbering
import numbering_policy as n


class TemplateContractExtra(Fixtures):
    # Avoid inheriting all base test methods a second time: suite loader below
    # only selects the cases defined in this class.
    def test_actual_approved_numbering_definition(self):
        raw=(Path(__file__).parent/'fixtures'/'approved-default-numbering.xml').read_bytes()
        def patch(data):
            numbering=n.parse(raw);styles=n.parse(data['word/styles.xml'])
            for lv in range(3):
                s=styles.find('w:style[@w:styleId="Heading'+str(lv+1)+'"]',n.NS)
                sid=str(lv+2);s.set(n.Q('styleId'),sid)
                p=s.find(n.Q('pPr'));np=n.set_child(p,'numPr');n.set_child(np,'numId',val='1')
            doc=n.parse(data['word/document.xml'])
            for p in doc.iter(n.Q('p')):
                st=p.find('w:pPr/w:pStyle',n.NS)
                if st is not None and n.value(st,'val','').startswith('Heading'):
                    st.set(n.Q('val'),str(int(n.value(st)[7:])+1))
            data['word/numbering.xml']=raw;data['word/styles.xml']=n.dump(styles);data['word/document.xml']=n.dump(doc)
        rewrite(self.template,patch)
        self.sample();self.normalize();doc=n.Doc(self.out)
        for p in doc.paragraphs:
            if doc.level(p)==0:
                definition=doc.num(p)['definition']
                self.assertEqual(n.value(definition.find(n.Q('numFmt'))),'decimal')
                self.assertEqual(n.value(definition.find(n.Q('lvlText'))),'%1')
                self.assertEqual(n.value(definition.find(n.Q('suff'))),'space')

    def test_style_level_numbering_for_future_headings(self):
        self.sample();self.normalize();d=n.Doc(self.out)
        p=next(p for p in d.paragraphs if d.level(p)==0)
        new=n.fake_p(d.style_id(p));new.append(n.run('新章节'))
        self.assertEqual(d.num(p)['id'],d.num(new)['id'])

    def test_manual_template_preserves_inner_spaces(self):
        d=Document();d.add_heading('第 1 章 模板写法',1);d.add_paragraph('图1 示例',style='Caption');d.save(self.template)
        self.sample();self.normalize();doc=n.Doc(self.out)
        p=next(p for p in doc.paragraphs if doc.level(p)==0)
        self.assertEqual(n.value(doc.num(p)['definition'].find(n.Q('lvlText'))),'第 %1 章')

    def test_image_filename_not_rewritten_as_reference(self):
        d=Document();d.add_heading('实施',1);d.add_paragraph('源文件为图1.1.png，存放在C:\\图1.1.png。');d.save(self.source)
        self.normalize();self.assertIn('图1.1.png',n.visible(n.Doc(self.out).paragraphs[1]))

    def test_empty_caption_title_filled_from_real_alt(self):
        d=Document();d.add_heading('实施',1);s=d.add_picture(str(self.image));s._inline.docPr.set('descr','架构图');d.add_paragraph('图1 ',style='Caption');d.save(self.source)
        self.normalize();self.assertEqual(n.visible(n.inventory(n.Doc(self.out))[0]['caption']),'图1 架构图')

    def test_unknown_reference_plan_not_silently_ignored(self):
        self.sample()
        with self.assertRaises(n.PolicyError):self.normalize({'source_sha256':n.file_digest(self.source),'objects':[],'references':[]})

    def test_template_run_size_checked_after_repair(self):
        self.sample();self.normalize();d=n.Doc(self.out);p=next(p for p in d.paragraphs if d.level(p)==0)
        st=d.style_map[d.style_id(p)];rp=st.find(n.Q('rPr'));n.set_child(rp,'sz',val='80');d.changed.add('word/styles.xml');bad=self.root/'bad.docx';d.save(bad)
        issues=n.audit_document(bad,self.template)['issues'];self.assertIn('heading-template-style',[x['code'] for x in issues])

    def test_complex_hidden_reset_after_word_serialization(self):
        self.sample();self.normalize();d=n.Doc(self.out)
        for p in d.paragraphs:
            if d.level(p)!=0:continue
            for f in list(n.fields(p)):
                simple=f['node'];at=p.index(simple);p.remove(simple)
                pieces=[]
                rr=n.node('r');rr.append(n.node('fldChar',fldCharType='begin'));pieces.append(rr)
                rr=n.node('r');it=n.node('instrText');it.text=f['instruction'];rr.append(it);pieces.append(rr)
                rr=n.node('r');rr.append(n.node('fldChar',fldCharType='separate'));pieces.append(rr)
                rp=n.node('rPr');rp.append(n.node('vanish'));pieces.append(n.run('0',rp))
                rr=n.node('r');rr.append(n.node('fldChar',fldCharType='end'));pieces.append(rr)
                for i,c in enumerate(pieces):p.insert(at+i,c)
        d.changed.add('word/document.xml');converted=self.root/'complex.docx';d.save(converted)
        result=self.root/'rerun.docx';n.normalize_document(converted,result,self.template)
        after=n.Doc(result)
        self.assertEqual([n.visible(p) for p in after.paragraphs if after.level(p)==0],['项目需求','实施方案'])
        for p in after.paragraphs:
            if after.level(p)==0:self.assertEqual(len(n.fields(p)),2)

    def test_heading_spaced_year_and_quantity_preserved(self):
        d=Document()
        for value in ('2024 年数据处理方案','2024.10.31 执行记录','500 万元设备预算'):
            d.add_heading(value,1)
        d.save(self.source);self.normalize()
        self.assertEqual([n.visible(p) for p in n.Doc(self.out).paragraphs],['2024 年数据处理方案','2024.10.31 执行记录','500 万元设备预算'])

    def test_chinese_word_native_caption_sequence_names(self):
        self.sample();self.normalize();doc=n.Doc(self.out)
        for obj in n.inventory(doc):
            self.assertEqual(n.fields(obj['caption'])[0]['instruction'].split()[1], '图' if obj['kind']=='figure' else '表')

    def test_template_output_cannot_overwrite_source_template(self):
        self.sample();before=self.template.read_bytes()
        with self.assertRaises(n.PolicyError):n.normalize_document(self.source,self.template,self.template)
        self.assertEqual(self.template.read_bytes(),before)

    def test_explicit_template_font_beats_inherited_theme_font(self):
        def patch(data):
            styles=n.parse(data['word/styles.xml'])
            for sid in ('Heading1','Caption'):
                st=styles.find('w:style[@w:styleId="'+sid+'"]',n.NS)
                rp=st.find(n.Q('rPr'));rf=n.set_child(rp,'rFonts')
                for key in list(rf.attrib):rf.attrib.pop(key)
                for slot in ('ascii','hAnsi','eastAsia'):rf.set(n.Q(slot),'Noto Serif CJK SC')
            data['word/styles.xml']=n.dump(styles)
        rewrite(self.template,patch);self.sample();self.normalize();doc=n.Doc(self.out)
        for p in doc.paragraphs:
            if doc.level(p)==0 or doc.style_name(p).lower()=='caption':
                rf=doc.effective(doc.style_id(p),'rPr').find(n.Q('rFonts'))
                for slot in ('ascii','hAnsi','eastAsia'):self.assertEqual(rf.get(n.Q(slot)),'Noto Serif CJK SC')

    def test_generated_style_next_precedes_format_properties(self):
        self.sample();self.normalize();doc=n.Doc(self.out)
        p=next(p for p in doc.paragraphs if doc.level(p)==0);st=doc.style_map[doc.style_id(p)]
        children=list(st)
        self.assertLess(children.index(st.find(n.Q('next'))),children.index(st.find(n.Q('pPr'))))


def load_tests(loader, tests, pattern):
    return unittest.TestSuite(TemplateContractExtra(name) for name in TemplateContractExtra.__dict__ if name.startswith('test_'))
