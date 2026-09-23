import copy
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from lxml import etree as E
from PIL import Image

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS))
import numbering_policy as n


def rewrite(path, transform):
    with zipfile.ZipFile(path) as z:
        data = {p: z.read(p) for p in z.namelist()}
    transform(data)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, content in data.items(): z.writestr(name, content)


def add_native_numbering(path, fmt='decimal', pattern='%1', suffix='space', levels=3):
    def patch(data):
        root = n.parse(data['word/numbering.xml'])
        aid = max(int(n.value(x, 'abstractNumId')) for x in root.findall(n.Q('abstractNum'))) + 1
        nid = max(int(n.value(x, 'numId')) for x in root.findall(n.Q('num'))) + 1
        a = n.node('abstractNum', abstractNumId=aid)
        a.append(n.node('multiLevelType', val='multilevel'))
        for lv in range(levels):
            l = n.node('lvl', ilvl=lv)
            for tag, value in [('start', 1), ('numFmt', fmt if lv == 0 else 'decimal'),
                               ('pStyle', 'Heading' + str(lv+1)), ('suff', suffix),
                               ('lvlText', pattern if lv == 0 else '.'.join('%'+str(i+1) for i in range(lv+1))), ('lvlJc', 'left')]:
                l.append(n.node(tag, val=value))
            a.append(l)
        root.insert(0, a)
        num = n.node('num', numId=nid); num.append(n.node('abstractNumId', val=aid)); root.append(num)
        styles = n.parse(data['word/styles.xml'])
        for lv in range(levels):
            s = styles.find('w:style[@w:styleId="Heading'+str(lv+1)+'"]', n.NS)
            ppr = s.find(n.Q('pPr'))
            if ppr is None: ppr = n.node('pPr'); s.append(ppr)
            old = ppr.find(n.Q('numPr'))
            if old is not None: ppr.remove(old)
            np = n.node('numPr'); np.append(n.node('ilvl', val=lv)); np.append(n.node('numId', val=nid)); ppr.append(np)
        data['word/numbering.xml'] = n.dump(root); data['word/styles.xml'] = n.dump(styles)
    rewrite(path, patch)


class Fixtures(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.template = self.root / 'template.docx'
        d = Document(); d.add_heading('项目要求', 1); d.add_heading('实施方法', 2); d.add_heading('检查步骤', 3)
        d.add_paragraph('图1 示例图', style='Caption'); d.save(self.template)
        add_native_numbering(self.template)
        self.image = self.root / 'asset.png'; Image.new('RGB', (100, 60), 'white').save(self.image)
        self.source = self.root / 'source.docx'; self.out = self.root / 'out.docx'

    def tearDown(self): self.temp.cleanup()

    def sample(self, captions=True, headings=('第一章 项目需求', '第二章 实施方案')):
        d = Document()
        for i, title in enumerate(headings, 1):
            d.add_heading(title, 1)
            d.add_paragraph(f'总体路线见图{i}.1。')
            d.add_picture(str(self.image), width=Inches(2))
            if captions: d.add_paragraph(f'图{i}.1 总体架构图', style='Caption')
            d.add_paragraph(f'设备清单见表{i}.1。')
            if captions: d.add_paragraph(f'表{i}.1 设备配置明细表', style='Caption')
            t = d.add_table(rows=2, cols=2)
            t.cell(0,0).text='设备'; t.cell(0,1).text='数量'; t.cell(1,0).text='工作站'; t.cell(1,1).text='2'
        d.save(self.source)
        return self.source

    def normalize(self, plan=None):
        return n.normalize_document(self.source, self.out, self.template, object_plan=plan)

    def test_manual_chinese_heading_uses_template_decimal(self):
        self.sample(); self.normalize(); doc=n.Doc(self.out)
        hs=[p for p in doc.paragraphs if doc.level(p)==0]
        self.assertEqual([n.visible(p) for p in hs], ['项目需求','实施方案'])
        self.assertTrue(all(n.value(doc.num(p)['definition'].find(n.Q('numFmt'))) == 'decimal' for p in hs))
        self.assertTrue(all(n.value(doc.num(p)['definition'].find(n.Q('lvlText'))) == '%1' for p in hs))

    def test_valid_but_wrong_automatic_heading_is_corrected(self):
        self.sample(headings=('项目需求','实施方案'))
        add_native_numbering(self.source, fmt='chineseCounting',pattern='第%1章')
        self.normalize(); doc=n.Doc(self.out)
        self.assertTrue(all(n.value(doc.num(p)['definition'].find(n.Q('lvlText'))) == '%1' for p in doc.paragraphs if doc.level(p)==0))

    def test_template_chinese_pattern_is_respected(self):
        add_native_numbering(self.template, fmt='chineseCounting', pattern='第%1章', suffix='nothing')
        self.sample(); self.normalize(); doc=n.Doc(self.out)
        for p in doc.paragraphs:
            if doc.level(p)==0:
                self.assertEqual(n.value(doc.num(p)['definition'].find(n.Q('lvlText'))), '第%1章')
                self.assertEqual(n.value(doc.num(p)['definition'].find(n.Q('suff'))), 'nothing')

    def test_template_number_suffix_is_preserved(self):
        add_native_numbering(self.template, pattern='%1、', suffix='tab')
        self.sample(); self.normalize(); doc=n.Doc(self.out)
        for p in doc.paragraphs:
            if doc.level(p)==0:
                self.assertEqual(n.value(doc.num(p)['definition'].find(n.Q('lvlText'))), '%1、')
                self.assertEqual(n.value(doc.num(p)['definition'].find(n.Q('suff'))), 'tab')

    def test_chapter_local_captions_no_prefix(self):
        self.sample(); self.normalize(); doc=n.Doc(self.out)
        texts=[n.visible(o['caption']) for o in n.inventory(doc)]
        self.assertEqual(texts, ['图1 总体架构图','表1 设备配置明细表']*2)
        for o in n.inventory(doc):
            fs=n.fields(o['caption']); self.assertEqual(len(fs),1)
            self.assertIn('\\s 1',fs[0]['instruction']); self.assertNotIn('\\r',fs[0]['instruction'])
            self.assertNotIn('STYLEREF',fs[0]['instruction'])

    def test_reset_fields_anchored_to_each_chapter(self):
        self.sample(); self.normalize(); doc=n.Doc(self.out)
        for p in doc.paragraphs:
            if doc.level(p)==0:
                fs=n.fields(p); self.assertEqual(len(fs),2)
                self.assertTrue(all('\\r 0 \\h' in f['instruction'] for f in fs))
                self.assertTrue(all(f['node'].find('w:r/w:rPr/w:vanish',n.NS) is not None for f in fs))

    def test_relative_references(self):
        self.sample(); self.normalize(); doc=n.Doc(self.out)
        alltext=[n.visible(p) for p in doc.paragraphs]
        self.assertIn('总体路线见下图。',alltext)
        self.assertIn('设备清单见下表。',alltext)
        self.assertFalse(any('见图1.1' in t or '见表2.1' in t for t in alltext))

    def test_backward_reference(self):
        d=Document(); d.add_heading('第一章 实施',1); d.add_picture(str(self.image)); d.add_paragraph('图1.1 路线图',style='Caption'); d.add_paragraph('如图1.1所示。'); d.save(self.source)
        self.normalize(); self.assertIn('如上图所示。',[n.visible(p) for p in n.Doc(self.out).paragraphs])

    def test_missing_captions_use_agent_semantic_plan(self):
        d=Document(); d.add_heading('第一章 实施',1); d.add_picture(str(self.image)); t=d.add_table(rows=1,cols=2); t.cell(0,0).text='对象';t.cell(0,1).text='指标'; d.save(self.source)
        inv=n.inventory(n.Doc(self.source)); plan={'version':1,'source_sha256':n.file_digest(self.source),'objects':[{'id':o['id'],'object_sha256':o['hash'],'title':'总体架构图' if o['kind']=='figure' else '指标明细表'} for o in inv]}
        self.normalize(plan); self.assertEqual([n.visible(o['caption']) for o in n.inventory(n.Doc(self.out))],['图1 总体架构图','表1 指标明细表'])

    def test_missing_caption_cannot_silently_pass(self):
        d=Document(); d.add_heading('实施',1); d.add_picture(str(self.image)); d.save(self.source)
        with self.assertRaises(n.PolicyError) as c: self.normalize()
        self.assertEqual(c.exception.code,'caption-title-required'); self.assertFalse(self.out.exists())

    def test_title_from_real_alt_text(self):
        d=Document(); d.add_heading('实施',1); shape=d.add_picture(str(self.image)); shape._inline.docPr.set('descr','总体架构图'); d.save(self.source)
        self.normalize(); self.assertEqual(n.visible(n.inventory(n.Doc(self.out))[0]['caption']),'图1 总体架构图')

    def test_stale_plan_rejected(self):
        self.sample(); inv=n.inventory(n.Doc(self.source))
        with self.assertRaises(n.PolicyError) as c: self.normalize({'source_sha256':'sha256:old','objects':[]})
        self.assertEqual(c.exception.code,'object-plan-stale')

    def test_missing_caption_audit_is_error(self):
        d=Document();d.add_heading('实施',1);d.add_picture(str(self.image));d.save(self.source)
        result=n.audit_document(self.source,self.template)
        self.assertIn('caption-missing',[i['code'] for i in result['issues']])

    def test_failed_repair_does_not_overwrite_previous_output(self):
        self.out.write_bytes(b'previous deliverable')
        d=Document();d.add_heading('实施',1);d.add_picture(str(self.image));d.save(self.source)
        with self.assertRaises(n.PolicyError): self.normalize()
        self.assertEqual(self.out.read_bytes(),b'previous deliverable')

    def test_template_undefined_numbering_not_inferred_from_broken_source(self):
        d=Document();d.add_heading('项目要求',1);d.add_paragraph('图1 示例',style='Caption');d.save(self.template)
        self.sample()
        with self.assertRaises(n.PolicyError) as c: self.normalize()
        self.assertEqual(c.exception.code,'template-numbering-missing')

    def test_existing_simple_styleref_error_removed(self):
        d=Document();d.add_heading('实施',1);d.add_picture(str(self.image));p=d.add_paragraph(style='Caption')
        p._p.append(n.run('图 '));p._p.append(n.seq_field('STYLEREF "不存在的标题" \\n','错误！未定义样式。'))
        p._p.append(n.run('.'));p._p.append(n.seq_field('SEQ Figure \\* ARABIC','1'));p._p.append(n.run(' 总体架构图'));d.save(self.source)
        self.normalize();cap=n.inventory(n.Doc(self.out))[0]['caption']
        self.assertEqual(n.visible(cap),'图1 总体架构图');self.assertNotIn('STYLEREF',E.tostring(cap).decode())

    def test_existing_complex_caption_fields_replaced_whole(self):
        d=Document();d.add_heading('实施',1);d.add_picture(str(self.image));p=d.add_paragraph(style='Caption');p._p.append(n.run('图 '))
        for instruction,cached in [('STYLEREF "bad" \\n','错误！未定义样式。'),('SEQ Figure \\* ARABIC','1')]:
            rr=n.node('r');rr.append(n.node('fldChar',fldCharType='begin'));p._p.append(rr)
            rr=n.node('r');it=n.node('instrText');it.text=instruction;rr.append(it);p._p.append(rr)
            rr=n.node('r');rr.append(n.node('fldChar',fldCharType='separate'));p._p.append(rr)
            p._p.append(n.run(cached));rr=n.node('r');rr.append(n.node('fldChar',fldCharType='end'));p._p.append(rr)
            if instruction.startswith('STYLEREF'):p._p.append(n.run('.'))
        p._p.append(n.run(' 架构图'));d.save(self.source)
        self.normalize();cap=n.inventory(n.Doc(self.out))[0]['caption'];self.assertEqual(n.visible(cap),'图1 架构图');self.assertEqual(len(n.fields(cap)),1)

    def test_unaffected_table_geometry_and_media_preserved(self):
        self.sample();self.normalize()
        from layout_invariant_guard import snapshot_docx
        a,b=snapshot_docx(self.source),snapshot_docx(self.out)
        for key in ('table_geometry','table_structure','drawings','sections','media'):
            self.assertEqual(a[key],b[key],key)

    def test_repeated_run_no_duplicate_caption_or_reset(self):
        self.sample();self.normalize();second=self.root/'second.docx'
        n.normalize_document(self.out,second,self.template)
        a,b=n.Doc(self.out),n.Doc(second)
        self.assertEqual([n.visible(p) for p in a.paragraphs],[n.visible(p) for p in b.paragraphs])
        self.assertEqual(len(a.numbering.findall(n.Q('abstractNum'))),len(b.numbering.findall(n.Q('abstractNum'))))
        self.assertEqual(sum(len(n.fields(p)) for p in a.paragraphs),sum(len(n.fields(p)) for p in b.paragraphs))

    def test_insert_first_caption_and_recompute(self):
        self.sample();self.normalize()
        d=Document(self.out);old=d.paragraphs[1]._p
        p=d.add_paragraph();p.add_run().add_picture(str(self.image));old.addprevious(p._p)
        cap=d.add_paragraph('图1 新增架构图',style='Caption');p._p.addnext(cap._p)
        changed=self.root/'insert.docx';d.save(changed);new=self.root/'insert-out.docx'
        n.normalize_document(changed,new,self.template)
        figs=[n.visible(o['caption']) for o in n.inventory(n.Doc(new)) if o['kind']=='figure']
        self.assertEqual(figs,['图1 新增架构图','图2 总体架构图','图1 总体架构图'])

    def test_delete_first_caption_keeps_chapter_reset(self):
        self.sample();self.normalize();d=n.Doc(self.out);first=n.inventory(d)[0]
        first['caption'].getparent().remove(first['caption']);first['element'].getparent().remove(first['element'])
        # Drop its now-invalid body sentence too; remaining references are relative.
        for p in list(d.paragraphs):
            if n.visible(p)=='总体路线见下图。':p.getparent().remove(p);break
        d.changed.add('word/document.xml');changed=self.root/'delete.docx';d.save(changed);new=self.root/'delete-out.docx'
        n.normalize_document(changed,new,self.template)
        self.assertEqual([n.visible(o['caption']) for o in n.inventory(n.Doc(new)) if o['kind']=='figure'],['图1 总体架构图'])

    def test_heading_text_year_not_treated_as_number(self):
        d=Document();d.add_heading('2024年数据处理方案',1);d.save(self.source)
        self.normalize();self.assertEqual(n.visible(n.Doc(self.out).paragraphs[0]),'2024年数据处理方案')

    def test_table_caption_moved_above_table(self):
        d=Document();d.add_heading('实施',1);t=d.add_table(rows=1,cols=1);t.cell(0,0).text='设备';d.add_paragraph('表1 配置表',style='Caption');d.save(self.source)
        self.normalize();obj=n.inventory(n.Doc(self.out))[0];self.assertIs(obj['element'].getprevious(),obj['caption'])

    def test_ambiguous_reference_blocks_instead_of_related_diagram(self):
        d=Document();d.add_heading('实施',1);d.add_paragraph('见图1。')
        for _ in range(2):d.add_picture(str(self.image));d.add_paragraph('图1 架构图',style='Caption')
        d.save(self.source)
        with self.assertRaises(n.PolicyError) as c:self.normalize()
        self.assertEqual(c.exception.code,'reference-target-ambiguous')

    def test_caption_title_run_emphasis_and_bookmark_preserved(self):
        d=Document();d.add_heading('实施',1);d.add_picture(str(self.image));p=d.add_paragraph(style='Caption')
        p._p.append(n.node('bookmarkStart',id='40',name='FigureAnchor'));p.add_run('图1.1 ')
        p.add_run('总体架构图').underline=True;p._p.append(n.node('bookmarkEnd',id='40'));d.save(self.source)
        self.normalize();cap=n.inventory(n.Doc(self.out))[0]['caption']
        self.assertEqual(len(cap.findall(n.Q('bookmarkStart'))),1);self.assertEqual(len(cap.findall(n.Q('bookmarkEnd'))),1)
        self.assertIsNotNone(cap.find('.//w:rPr/w:u',n.NS))

    def test_decoration_exemption_requires_reason(self):
        d=Document();d.add_heading('实施',1);d.add_picture(str(self.image));d.save(self.source);obj=n.inventory(n.Doc(self.source))[0]
        plan={'source_sha256':n.file_digest(self.source),'objects':[{'id':obj['id'],'object_sha256':obj['hash'],'exempt':True}]}
        with self.assertRaises(n.PolicyError) as c:self.normalize(plan)
        self.assertEqual(c.exception.code,'exemption-without-reason')

    def test_legitimate_decoration_does_not_receive_fake_caption(self):
        d=Document();d.add_heading('实施',1);d.add_picture(str(self.image));d.save(self.source);obj=n.inventory(n.Doc(self.source))[0]
        plan={'source_sha256':n.file_digest(self.source),'objects':[{'id':obj['id'],'object_sha256':obj['hash'],'exempt':True,'reason':'正文装饰性分隔图，没有业务信息'}]}
        self.normalize(plan);self.assertIsNone(n.inventory(n.Doc(self.out))[0]['caption'])


if __name__=='__main__':unittest.main()
